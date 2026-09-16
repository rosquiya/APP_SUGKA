#!/usr/bin/env python3
"""
SUGKA LAB – API Server
REST API Flask para la aplicacion de gestion educativa UGEL IBIR-IMAZA.
Puerto: 8000
"""
from flask import Flask, request, jsonify, send_file, send_from_directory, g
from flask_cors import CORS
from functools import wraps
import base64
import csv
import hashlib
import io
import sqlite3
import json
import os
import re
import secrets
import shutil
import subprocess
import tempfile
import unicodedata
import urllib.request
from datetime import datetime, timedelta
from pathlib import Path
from rapidfuzz import fuzz as rapidfuzz_fuzz, process as rapidfuzz_process

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024  # 20MB por request (protege /api/ocr/upload_advanced)

# Orígenes permitidos para CORS. En producción, define ALLOWED_ORIGINS en el
# entorno con el/los dominios reales del frontend (separados por coma), p.ej.
# "https://app-sugka.onrender.com". Sin definir, se mantiene abierto (*) para
# no romper el desarrollo local.
_allowed_origins = [o.strip() for o in os.getenv('ALLOWED_ORIGINS', '*').split(',') if o.strip()]
CORS(app, origins=_allowed_origins if _allowed_origins != ['*'] else '*')

@app.errorhandler(413)
def handle_file_too_large(_e):
    return jsonify({'error': 'El archivo supera el tamaño máximo permitido (20MB).'}), 413

# Rutas relativas al directorio raíz del proyecto
ROOT = Path(__file__).parent.parent
SEED_DB_PATH = ROOT / 'data' / 'database' / 'sugka_demo.db'
GENERIC_EMAIL_DOMAIN = 'ugel-imaza.edu.pe'

def load_local_env():
    env_path = ROOT / '.env'
    if not env_path.exists():
        return
    for line in env_path.read_text(encoding='utf-8').splitlines():
        line = line.strip()
        if not line or line.startswith('#') or '=' not in line:
            continue
        key, value = line.split('=', 1)
        key = key.strip()
        value = value.strip().strip('"').strip("'")
        if value and not os.environ.get(key):
            os.environ[key] = value

load_local_env()

# En local, la BD vive en el repo (data/database/sugka_demo.db). En Render,
# SUGKA_DB_PATH apunta al disco persistente montado (ver render.yaml) --
# necesario porque el filesystem del contenedor es efimero y sin esto la BD
# se perderia en cada redeploy. Si el disco esta vacio (primer deploy), se
# copia una vez la BD semilla del repo para no arrancar sin instituciones,
# especialistas ni usuarios.
DB_PATH = Path(os.getenv('SUGKA_DB_PATH') or SEED_DB_PATH)
if DB_PATH.resolve() != SEED_DB_PATH.resolve() and not DB_PATH.exists() and SEED_DB_PATH.exists():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(SEED_DB_PATH, DB_PATH)

def get_azure_config():
    load_local_env()
    return (
        os.getenv("AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT", ""),
        os.getenv("AZURE_DOCUMENT_INTELLIGENCE_KEY", ""),
    )

def get_gemini_key():
    load_local_env()
    key = os.getenv("GEMINI_API_KEY", "")
    if key:
        genai.configure(api_key=key)
    return key

# gemini-1.5-flash y luego gemini-2.5-flash fueron retirados por Google
# (los modelos dejan de responder generateContent con 404 NotFound cuando
# se retiran) -- usar un modelo vigente, configurable via GEMINI_MODEL por
# si Google retira este tambien mas adelante.
GEMINI_DEFAULT_MODEL = os.getenv('GEMINI_MODEL', 'gemini-3.6-flash')

def get_openai_key():
    load_local_env()
    return os.getenv("OPENAI_API_KEY", "")

# Variables concretas del subproyecto "informes_campo_v3" (ver
# config/variables_concretas.json en ese subproyecto). Cada variable es una
# bandera 1/None con evidencia textual obligatoria en span_<variable>; no se
# calculan puntajes de gravedad/urgencia/confianza -- mismo principio que se
# uso para no "adivinar" en SIMON.
CAMPO_CONCRETE_VARIABLES = [
    'presenta_violencia',
    'ausencia_docente',
    'falta_materiales',
    'infraestructura_deficiente',
    'riesgo_salud_seguridad',
    'inasistencia_estudiantes',
    'eib_no_practicado',
    'gestion_directiva_debil',
    'retroalimentacion_debil',
    'acceso_dificil',
    'alimentacion_qaliwarma_problema',
]

# Metadata documentada de cada variable concreta (ver
# docs/01_diccionario_features.md y config/variables_concretas.json del
# subproyecto informes_campo_v3). Se usa para: (1) instruir a la IA con el
# criterio textual exacto en vez de un criterio generico, y (2) mostrarle al
# especialista, junto a cada hallazgo detectado, por que se marco y que
# significa -- nunca se calculan puntajes de gravedad/urgencia/confianza.
# `severidad` es una decision editorial de priorizacion de la app (no algo
# que la IA infiera del texto).
CAMPO_VARIABLE_META = {
    'presenta_violencia': {
        'label': 'Violencia o conflicto reportado',
        'criterio': 'Frase con violencia, maltrato, denuncia, agresion, conflicto, bullying o acoso.',
        'importancia': 'Senal social/institucional que SIMON no observa; puede requerir derivacion a otra instancia.',
        'ejemplo': 'denuncia de maltrato',
        'severidad': 'alta',
    },
    'ausencia_docente': {
        'label': 'Ausencia docente',
        'criterio': 'Frase de no asistio, docente ausente, falta docente, sin docente, plaza vacante.',
        'importancia': 'Afecta directamente la continuidad del servicio educativo.',
        'ejemplo': 'presentes 2 de 4 docentes',
        'severidad': 'alta',
    },
    'falta_materiales': {
        'label': 'Falta de materiales educativos',
        'criterio': 'Frase de no cuenta con materiales, falta material educativo, sin textos/cuadernos.',
        'importancia': 'Condicion operativa que limita el aprendizaje aunque el docente tenga buen desempeno.',
        'ejemplo': 'no llegaron los cuadernos de trabajo',
        'severidad': 'media',
    },
    'infraestructura_deficiente': {
        'label': 'Infraestructura o servicios deficientes',
        'criterio': 'Frase de deterioro, inoperatividad, sin agua/luz, SS.HH., aulas/mobiliario en mal estado.',
        'importancia': 'Se puede cruzar con Censo (P53/P61) para distinguir alerta puntual de condicion estructural conocida.',
        'ejemplo': 'aula sin techo, mobiliario deteriorado',
        'severidad': 'media',
    },
    'riesgo_salud_seguridad': {
        'label': 'Riesgo de salud o seguridad',
        'criterio': 'Frase de riesgo, peligro, seguridad, botiquin, primeros auxilios, condicion sanitaria.',
        'importancia': 'Variable mas frecuente en la corrida historica (41.9%); prioridad inmediata de acompanamiento.',
        'ejemplo': 'no presenta plan de gestion de riesgo',
        'severidad': 'alta',
    },
    'inasistencia_estudiantes': {
        'label': 'Inasistencia de estudiantes',
        'criterio': 'Frase de inasistencia, estudiantes ausentes, abandono, desercion.',
        'importancia': 'Alerta temprana de riesgo de desercion, complementaria a la matricula del Censo.',
        'ejemplo': 'alumnos ausentes de forma reiterada',
        'severidad': 'media',
    },
    'eib_no_practicado': {
        'label': 'EIB no practicado',
        'criterio': 'Frase explicita de no uso/no realizacion de actividades EIB o lengua materna.',
        'importancia': 'Relevante en Imaza por el contexto intercultural de la zona; mide brecha de pertinencia cultural.',
        'ejemplo': 'no se usa lengua materna en el aula',
        'severidad': 'media',
    },
    'gestion_directiva_debil': {
        'label': 'Gestion directiva debil',
        'criterio': 'Frase sobre PEI, PAT, RI, PCI, RD, comites no conformados o documentos ausentes.',
        'importancia': 'Senala capacidad institucional de la IE, factor que sostiene o debilita cualquier intervencion.',
        'ejemplo': 'no esta conformado el comite de gestion',
        'severidad': 'media',
    },
    'retroalimentacion_debil': {
        'label': 'Retroalimentacion pedagogica debil',
        'criterio': 'Frase de retroalimentacion elemental, no realizada o no reflexiva.',
        'importancia': 'Punto de triangulacion directo con AR02.3 de SIMON; refuerza o contrasta la senal pedagogica.',
        'ejemplo': 'retroalimentacion no descriptiva',
        'severidad': 'media',
    },
    'acceso_dificil': {
        'label': 'Acceso dificil a la IE',
        'criterio': 'Frase de dificil acceso, rio, trocha, traslado, distancia, llegada tarde por procedencia.',
        'importancia': 'Explica limitaciones operativas de la UGEL para monitorear con la frecuencia deseada.',
        'ejemplo': 'se accede solo por rio, 4 horas de traslado',
        'severidad': 'baja',
    },
    'alimentacion_qaliwarma_problema': {
        'label': 'Problema con alimentacion escolar (Qali Warma)',
        'criterio': 'Frase que menciona Qali Warma, alimentos o desayuno con problema.',
        'importancia': 'Afecta condiciones basicas de permanencia escolar, relevante en zona de alta ruralidad.',
        'ejemplo': 'no llego el desayuno escolar',
        'severidad': 'media',
    },
}

# Como cada IE solo tiene dos "cubetas" operativas historicas (infraestructura
# y pedagogica), agrupamos las 11 variables concretas de campo en esas dos
# para alimentar los contadores existentes; el resto (violencia, gestion
# directiva) entra en la cubeta pedagogica/institucional por ser mas cercana
# a seguimiento pedagogico-institucional que a infraestructura fisica.
CAMPO_INFRA_BUCKET = {
    'infraestructura_deficiente', 'riesgo_salud_seguridad',
    'acceso_dificil', 'alimentacion_qaliwarma_problema',
}
CAMPO_PEDAGOGIC_BUCKET = {
    'ausencia_docente', 'retroalimentacion_debil', 'eib_no_practicado',
    'inasistencia_estudiantes', 'falta_materiales', 'gestion_directiva_debil',
    'presenta_violencia',
}

def get_db():
    """Conexion sqlite3 cacheada en el contexto de la request actual (flask.g).

    Se cierra automaticamente en close_db_connection() al terminar la
    request, incluso si el handler lanza una excepcion antes de llegar a su
    propio conn.close() -- eso evita fugas de conexiones bajo carga.
    """
    if 'db_conn' not in g:
        g.db_conn = sqlite3.connect(str(DB_PATH))
        g.db_conn.row_factory = sqlite3.Row
        g.db_conn.execute('PRAGMA foreign_keys = ON')
    return g.db_conn

@app.teardown_appcontext
def close_db_connection(_exception=None):
    conn = g.pop('db_conn', None)
    if conn is not None:
        conn.close()

def rows_to_list(rows):
    return [dict(r) for r in rows]

# Roles del sistema. 'supervisor' (jefe) ve todo y puede corregir el registro de
# cualquier especialista, pero la administracion del sistema (usuarios,
# instituciones, preguntas dinamicas) sigue siendo exclusiva de administrador.
ROLES = ('administrador', 'supervisor', 'especialista')
ROLES_QUE_EDITAN_TODO = ('administrador', 'supervisor')
# Estados del flujo de revision de un informe de campo.
REVISION_ESTADOS = ('ok', 'pendiente', 'en_revision', 'resuelto')
ROLE_LABELS = {
    'administrador': 'Administrador',
    'supervisor': 'Supervisor',
    'especialista': 'Especialista',
}

def normalize_role(value='especialista'):
    role = str(value or 'especialista').strip().lower()
    if role in ('admin', 'administrador'):
        return 'administrador'
    if role in ('supervisor', 'jefe'):
        return 'supervisor'
    return 'especialista'

def role_label(role):
    return ROLE_LABELS.get(normalize_role(role), 'Especialista')

PBKDF2_ITERATIONS = 260_000

def hash_password(password, salt=None):
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac(
        'sha256', password.encode('utf-8'), salt.encode('utf-8'), PBKDF2_ITERATIONS
    ).hex()
    return salt, digest

def verify_password(password, salt, digest):
    if not password or not salt or not digest:
        return False
    _, candidate = hash_password(password, salt)
    if secrets.compare_digest(candidate, digest):
        return True
    # Compatibilidad con hashes antiguos (SHA-256 de una sola vuelta, sin
    # KDF). Permite que cuentas creadas antes de este cambio sigan
    # funcionando; en su proximo cambio de contraseña quedan migradas a
    # PBKDF2 automaticamente porque hash_password() ya usa el nuevo esquema.
    legacy_digest = hashlib.sha256(f'{salt}:{password}'.encode('utf-8')).hexdigest()
    return secrets.compare_digest(legacy_digest, digest)

def _seed_password(env_var, label):
    """Contrasena de un usuario semilla.

    Se toma de la variable de entorno indicada (en Render se configura con
    generateValue: true, asi que cada despliegue tiene una contrasena unica
    y aleatoria). Si no esta definida (por ejemplo en un entorno local sin
    .env), se genera una aleatoria en memoria y se imprime UNA vez en el
    log del proceso para que quien lo ejecute la pueda copiar. Nunca queda
    escrita en el codigo ni en el repositorio.
    """
    value = os.getenv(env_var)
    if value:
        return value
    generated = secrets.token_urlsafe(12)
    print(
        f"[SEED] {env_var} no definido: se genero una contrasena aleatoria "
        f"para el usuario '{label}': {generated}",
        flush=True,
    )
    return generated

def slugify_username(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode('ascii')
    text = re.sub(r'[^a-zA-Z0-9]+', '.', text).strip('.').lower()
    return text or 'usuario'

def public_user(row):
    user = dict(row)
    user.pop('password_hash', None)
    user.pop('password_salt', None)
    rol = normalize_role(user.get('rol'))
    user['rol'] = rol
    user['rol_label'] = role_label(rol)
    user['is_admin'] = rol == 'administrador'
    user['is_supervisor'] = rol == 'supervisor'
    # Lo consume el frontend para decidir si muestra los botones de editar de
    # registros ajenos, sin tener que repetir la tabla de roles alla.
    user['puede_editar_todo'] = rol in ROLES_QUE_EDITAN_TODO
    return user

def generic_email_for_name(nombre, domain=GENERIC_EMAIL_DOMAIN):
    return f'{slugify_username(nombre)}@{domain}'

def unique_account_identity(conn, username, email, exclude_user_id=None):
    username = slugify_username(username)
    email = str(email or generic_email_for_name(username)).strip().lower()
    email_user, _, email_domain = email.partition('@')
    if not email_domain:
        email_domain = GENERIC_EMAIL_DOMAIN

    base_username = username
    base_email_user = slugify_username(email_user)
    suffix = 1
    while True:
        params = [username, email]
        exclude_sql = ''
        if exclude_user_id:
            exclude_sql = ' AND user_id != ?'
            params.append(exclude_user_id)
        exists = conn.execute(
            f'''
            SELECT 1 FROM app_user
            WHERE (LOWER(username) = LOWER(?) OR LOWER(COALESCE(email, '')) = LOWER(?))
            {exclude_sql}
            LIMIT 1
            ''',
            params,
        ).fetchone()
        if not exists:
            return username, email
        suffix += 1
        username = f'{base_username}.{suffix}'
        email = f'{base_email_user}.{suffix}@{email_domain}'

def create_app_user(conn, nombre, username, password, rol='especialista', especialista_id=None, activo=1, email=None, cargo=None):
    if not nombre or not username or not password:
        raise ValueError('Nombre, usuario y contraseña son obligatorios.')
    username, email = unique_account_identity(
        conn,
        username,
        email or generic_email_for_name(nombre),
    )
    salt, digest = hash_password(password)
    user_id = f'user_{secrets.token_hex(6)}'
    conn.execute('''
        INSERT INTO app_user (
            user_id, nombre, username, email, password_salt, password_hash,
            rol, especialista_id, cargo, activo, creado_en, actualizado_en
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
    ''', (
        user_id,
        str(nombre).strip(),
        username,
        email,
        salt,
        digest,
        normalize_role(rol),
        especialista_id or None,
        str(cargo).strip() if cargo else None,
        1 if activo else 0,
    ))
    return conn.execute('SELECT * FROM app_user WHERE user_id = ?', (user_id,)).fetchone()

def json_dumps(value):
    return json.dumps(value if value is not None else {}, ensure_ascii=False)

def json_loads(value, default=None):
    if value in (None, ''):
        return default if default is not None else {}
    try:
        return json.loads(value)
    except (TypeError, ValueError):
        return default if default is not None else {}

LEVEL_LABELS = ('NIVEL I', 'NIVEL II', 'NIVEL III', 'NIVEL IV')

def normalize_question_code(value):
    code = str(value or '').strip().upper()
    code = re.sub(r'\s+', '-', code)
    code = re.sub(r'[^A-Z0-9_-]+', '', code)
    return code

def normalize_levels_payload(value):
    data = value if isinstance(value, dict) else {}
    roman_by_index = {1: 'I', 2: 'II', 3: 'III', 4: 'IV'}
    normalized = {}
    for index, roman in roman_by_index.items():
        label = f'NIVEL {roman}'
        raw = (
            data.get(label)
            or data.get(label.lower())
            or data.get(roman)
            or data.get(roman.lower())
            or data.get(f'nivel_{index}')
            or data.get(str(index))
            or ''
        )
        normalized[label] = str(raw).strip()
    return normalized

def build_instrument_payload(conn, row, include_inactive=False):
    base = json_loads(row['estructura_json'], {})
    item = dict(base)
    item['instrumento_id'] = row['instrumento_id']
    item['codigo'] = row['codigo']
    item['formulario'] = item.get('formulario') or row['nombre']
    item['tipo'] = row['tipo']
    item['version'] = row['version']
    item['fuente'] = row['fuente']

    sections = conn.execute('''
        SELECT *
        FROM app_instrumento_seccion
        WHERE instrumento_id = ?
        ORDER BY orden, seccion_id
    ''', (row['instrumento_id'],)).fetchall()

    if not sections:
        return item

    section_payload = []
    active_filter = '' if include_inactive else 'AND activo = 1'
    for section in sections:
        if not include_inactive and not section['activo']:
            continue
        questions = conn.execute(f'''
            SELECT *
            FROM app_instrumento_pregunta
            WHERE instrumento_id = ? AND seccion_id = ? {active_filter}
            ORDER BY orden, pregunta_id
        ''', (row['instrumento_id'], section['seccion_id'])).fetchall()
        section_payload.append({
            'seccion_id': section['seccion_id'],
            'clave': section['clave'],
            'nombre': section['nombre'],
            'orden': section['orden'],
            'activo': section['activo'],
            'preguntas': [
                {
                    'pregunta_id': question['pregunta_id'],
                    'id': question['codigo'],
                    'codigo': question['codigo'],
                    'item': question['item'],
                    'tipo_respuesta': question['tipo_respuesta'],
                    'niveles': json_loads(question['niveles_json'], {}),
                    'opciones': json_loads(question['opciones_json'], {}),
                    'orden': question['orden'],
                    'activo': question['activo'],
                    'observaciones': '',
                }
                for question in questions
            ],
        })
    item['secciones'] = section_payload
    return item

def get_dynamic_instrument_row(conn):
    return conn.execute('''
        SELECT *
        FROM app_instrumento
        WHERE tipo = 'dinamica' AND activo = 1
        ORDER BY nombre
        LIMIT 1
    ''').fetchone()

def sync_instrument_structure_json(conn, instrumento_id):
    row = conn.execute(
        'SELECT * FROM app_instrumento WHERE instrumento_id = ?',
        (instrumento_id,),
    ).fetchone()
    if not row:
        return
    payload = build_instrument_payload(conn, row)
    compact = {
        'formulario': payload.get('formulario'),
        'tipo': payload.get('tipo'),
        'datos_identificacion': payload.get('datos_identificacion', {}),
        'secciones': [
            {
                'clave': section.get('clave'),
                'nombre': section.get('nombre'),
                'preguntas': [
                    {
                        'id': question.get('codigo') or question.get('id'),
                        'item': question.get('item'),
                        'niveles': question.get('niveles', {}),
                        'observaciones': '',
                    }
                    for question in section.get('preguntas', [])
                ],
            }
            for section in payload.get('secciones', [])
        ],
        'campos_finales': payload.get('campos_finales', []),
        'codigo': payload.get('codigo'),
        'version': payload.get('version'),
    }
    conn.execute('''
        UPDATE app_instrumento
        SET estructura_json = ?, actualizado_en = CURRENT_TIMESTAMP
        WHERE instrumento_id = ?
    ''', (json_dumps(compact), instrumento_id))

def dynamic_questions_payload(conn):
    instrument = get_dynamic_instrument_row(conn)
    if not instrument:
        return None
    payload = build_instrument_payload(conn, instrument, include_inactive=True)
    sections = [
        {
            'seccion_id': section.get('seccion_id'),
            'clave': section.get('clave'),
            'nombre': section.get('nombre'),
            'orden': section.get('orden'),
            'activo': section.get('activo'),
        }
        for section in payload.get('secciones', [])
    ]
    questions = []
    for section in payload.get('secciones', []):
        for question in section.get('preguntas', []):
            questions.append({
                **question,
                'seccion_id': section.get('seccion_id'),
                'seccion_clave': section.get('clave'),
                'seccion_nombre': section.get('nombre'),
            })
    return {
        'instrumento': {
            'instrumento_id': instrument['instrumento_id'],
            'codigo': instrument['codigo'],
            'formulario': payload.get('formulario'),
            'tipo': instrument['tipo'],
            'version': instrument['version'],
        },
        'secciones': sections,
        'preguntas': questions,
    }

def build_empty_ficha_pdf(instrumentos):
    from fpdf import FPDF

    pdf = FPDF('P', 'mm', 'A4')
    font_family = 'Arial'
    unicode_font = False
    font_dir = Path(os.environ.get('WINDIR', 'C:/Windows')) / 'Fonts'
    regular_font = font_dir / 'arial.ttf'
    bold_font = font_dir / 'arialbd.ttf'
    if regular_font.exists() and bold_font.exists():
        try:
            try:
                pdf.add_font('ArialUnicode', '', str(regular_font), uni=True)
                pdf.add_font('ArialUnicode', 'B', str(bold_font), uni=True)
            except TypeError:
                pdf.add_font('ArialUnicode', '', str(regular_font))
                pdf.add_font('ArialUnicode', 'B', str(bold_font))
            font_family = 'ArialUnicode'
            unicode_font = True
        except Exception:
            font_family = 'Arial'
            unicode_font = False

    def safe(value):
        text = str(value or '')
        if unicode_font:
            return text
        replacements = {
            '\u2013': '-',
            '\u2014': '-',
            '\u2018': "'",
            '\u2019': "'",
            '\u201c': '"',
            '\u201d': '"',
            '\u2022': '-',
        }
        for old, new in replacements.items():
            text = text.replace(old, new)
        return text.encode('latin-1', 'replace').decode('latin-1')

    def add_page_header():
        pdf.set_fill_color(11, 53, 85)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font(font_family, 'B', 14)
        pdf.cell(0, 10, safe('SUGKA LAB - Ficha completa vacia'), 0, 1, 'C', True)
        pdf.set_text_color(16, 34, 53)
        pdf.set_font(font_family, '', 8)
        pdf.multi_cell(0, 5, safe('Instrumento imprimible generado desde las preguntas activas registradas en la base de datos.'))
        pdf.ln(2)

    def check_page(space=24):
        if pdf.get_y() > (297 - 14 - space):
            pdf.add_page()
            add_page_header()

    def section_title(title):
        check_page(18)
        pdf.set_fill_color(231, 242, 247)
        pdf.set_text_color(21, 95, 131)
        pdf.set_font(font_family, 'B', 10)
        pdf.cell(0, 8, safe(title), 0, 1, 'L', True)
        pdf.set_text_color(16, 34, 53)
        pdf.ln(1)

    def field_line(label):
        check_page(9)
        pdf.set_font(font_family, 'B', 8)
        pdf.cell(52, 7, safe(label[:36]), 0, 0)
        pdf.set_font(font_family, '', 8)
        pdf.cell(0, 7, '_' * 82, 0, 1)

    def grouped_identification(datos):
        for group_key, fields in (datos or {}).items():
            group_title = str(group_key).replace('_', ' ').title()
            section_title(group_title)
            for field in fields or []:
                field_line(field.get('etiqueta') or field.get('campo') or '')
            pdf.ln(1)

    def render_question(question):
        check_page(42)
        code = question.get('codigo') or question.get('id') or ''
        pdf.set_x(pdf.l_margin)
        pdf.set_font(font_family, 'B', 8)
        pdf.multi_cell(0, 5, safe(f'{code}. {question.get("item", "")}'), 0, 'L')
        pdf.set_font(font_family, '', 7)
        col_width = (pdf.w - pdf.l_margin - pdf.r_margin) / 4
        pdf.set_x(pdf.l_margin)
        for label in LEVEL_LABELS:
            pdf.cell(col_width, 7, safe(f'[  ] {label}'), 1, 0, 'C')
        pdf.ln(8)
        niveles = question.get('niveles') or {}
        for label in LEVEL_LABELS:
            text = niveles.get(label, '')
            if text:
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(0, 4, safe(f'{label}: {text}'))
        pdf.set_font(font_family, '', 8)
        pdf.set_x(pdf.l_margin)
        pdf.cell(25, 6, safe('Observacion:'), 0, 0)
        y = pdf.get_y() + 4
        pdf.line(pdf.get_x(), y, pdf.w - pdf.r_margin, y)
        pdf.ln(7)
        y = pdf.get_y() + 4
        pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
        pdf.ln(7)
        pdf.ln(2)

    def render_final_fields(fields):
        if not fields:
            return
        section_title('Campos finales')
        for field in fields:
            label = field.get('etiqueta') or field.get('campo') or ''
            pdf.set_font(font_family, 'B', 8)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(0, 5, safe(label))
            pdf.set_font(font_family, '', 8)
            y = pdf.get_y() + 4
            pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
            pdf.ln(7)
            if field.get('tipo') == 'texto_libre':
                y = pdf.get_y() + 4
                pdf.line(pdf.l_margin, y, pdf.w - pdf.r_margin, y)
                pdf.ln(7)
            pdf.ln(1)

    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    add_page_header()

    pdf.set_font(font_family, '', 9)
    pdf.cell(0, 6, safe(f'Fecha de impresion: {datetime.now().strftime("%d/%m/%Y %H:%M")}'), 0, 1)
    pdf.ln(2)

    for index, instrumento in enumerate(instrumentos):
        if index:
            pdf.add_page()
            add_page_header()
        section_title(instrumento.get('formulario') or instrumento.get('codigo') or 'Instrumento')
        pdf.set_font(font_family, '', 8)
        pdf.cell(0, 5, safe(f'Tipo: {"SIMON / fijo" if instrumento.get("tipo") == "fija" else "Dinamico UGEL"}  |  Version: {instrumento.get("version", "2026")}'), 0, 1)
        pdf.ln(1)
        grouped_identification(instrumento.get('datos_identificacion'))
        for section in instrumento.get('secciones', []):
            section_title(section.get('nombre') or section.get('clave') or 'Seccion')
            for question in section.get('preguntas', []):
                render_question(question)
        render_final_fields(instrumento.get('campos_finales') or [])

    content = pdf.output(dest='S')
    if isinstance(content, str):
        content = content.encode('latin-1')
    else:
        content = bytes(content)
    buffer = io.BytesIO(content)
    buffer.seek(0)
    return buffer

FICHA_MONITOREO_COLUMNS = {
    'source': 'TEXT',
    'file_name': 'TEXT',
    'extraction_method': 'TEXT',
    'confidence': 'REAL',
    'raw_text': 'TEXT',
    'extracted_json': 'TEXT',
    'region': 'TEXT',
    'ugel': 'TEXT',
    'n_visita': 'TEXT',
    'codigo_modular': 'TEXT',
    'nombre_ie': 'TEXT',
    'nivel_modalidad': 'TEXT',
    'director': 'TEXT',
    'director_cel': 'TEXT',
    'director_email': 'TEXT',
    'director_situacion_laboral': 'TEXT',
    'docente': 'TEXT',
    'docente_dni': 'TEXT',
    'docente_cel': 'TEXT',
    'docente_email': 'TEXT',
    'docente_situacion_laboral': 'TEXT',
    'grado': 'TEXT',
    'seccion': 'TEXT',
    'nro_estudiantes': 'INTEGER',
    'area': 'TEXT',
    'competencia': 'TEXT',
    'titulo_sesion': 'TEXT',
    'monitor': 'TEXT',
    'monitor_dni': 'TEXT',
    'iged': 'TEXT',
    'monitor_email': 'TEXT',
    'fecha_ejecucion': 'TEXT',
    'a1': 'INTEGER',
    'a2': 'INTEGER',
    'a3': 'INTEGER',
    'b1': 'INTEGER',
    'b2': 'INTEGER',
    'b3': 'INTEGER',
    'b4': 'INTEGER',
    'b5': 'INTEGER',
    'a1_observacion': 'TEXT',
    'a2_observacion': 'TEXT',
    'a3_observacion': 'TEXT',
    'b1_observacion': 'TEXT',
    'b2_observacion': 'TEXT',
    'b3_observacion': 'TEXT',
    'b4_observacion': 'TEXT',
    'b5_observacion': 'TEXT',
    'promedio': 'REAL',
    'observaciones': 'TEXT',
    'observaciones_recomendaciones': 'TEXT',
    'compromisos': 'TEXT',
    'compromisos_monitoreado': 'TEXT',
    # Estandar de auditoria: nunca se borra fisico, se marca estado_fila=0.
    'estado_fila': 'INTEGER DEFAULT 1',
    'actualizado_en': 'DATETIME',
    'creado_por': 'TEXT',
    'modificado_por': 'TEXT',
}

# Campos que el panel de "Registros" puede editar directamente: todo salvo
# metadata de origen/extraccion y las columnas de auditoria que solo maneja
# el servidor (estado_fila/creado_por/modificado_por/actualizado_en).
FICHA_NON_EDITABLE_FIELDS = {
    'estado_fila', 'actualizado_en', 'creado_por', 'modificado_por',
    'source', 'file_name', 'extraction_method', 'confidence',
    'raw_text', 'extracted_json',
}
FICHA_EDITABLE_FIELDS = [f for f in FICHA_MONITOREO_COLUMNS if f not in FICHA_NON_EDITABLE_FIELDS]

def ensure_db_schema():
    conn = get_db()
    columns_sql = ',\n            '.join(
        f'{name} {column_type}' for name, column_type in FICHA_MONITOREO_COLUMNS.items()
    )
    conn.execute(f'''
        CREATE TABLE IF NOT EXISTS fichas_monitoreo (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            {columns_sql},
            fecha_sincronizacion DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    ''')

    existing = {
        row['name'] for row in conn.execute('PRAGMA table_info(fichas_monitoreo)').fetchall()
    }
    for name, column_type in FICHA_MONITOREO_COLUMNS.items():
        if name not in existing:
            conn.execute(f'ALTER TABLE fichas_monitoreo ADD COLUMN {name} {column_type}')

    # Corrige filas creadas antes del fix de save_ficha_to_db (ver ahi el
    # comentario): quedaron con estado_fila NULL en vez de 1 y por eso el
    # panel de Registros las mostraba como "eliminadas" sin estarlo.
    conn.execute('UPDATE fichas_monitoreo SET estado_fila = 1 WHERE estado_fila IS NULL')

    conn.execute('CREATE INDEX IF NOT EXISTS idx_fichas_monitoreo_codigo ON fichas_monitoreo(codigo_modular)')

    # dim_institucion existe desde antes de este archivo (creada por el import
    # censal); solo le sumamos las columnas de auditoria/soft-delete si faltan.
    dim_institucion_tables = {
        row['name'] for row in conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name='dim_institucion'"
        ).fetchall()
    }
    if dim_institucion_tables:
        dim_existing = {
            row['name'] for row in conn.execute('PRAGMA table_info(dim_institucion)').fetchall()
        }
        for name, column_type in {
            'estado_fila': 'INTEGER DEFAULT 1',
            'creado_en': 'DATETIME',
            'actualizado_en': 'DATETIME',
        }.items():
            if name not in dim_existing:
                conn.execute(f'ALTER TABLE dim_institucion ADD COLUMN {name} {column_type}')

    # dim_alerta_categoria tambien es previa a este archivo (pensada para un
    # sistema de deteccion por keyword_column que ya no es el principal). Las
    # fuentes 'campo' (hallazgos de informe_campo_hallazgo) y 'dinamica'
    # (preguntas dinamicas Nivel I/II) se agregaron despues en
    # rebuild_simon_operational_summary() sin sumar su categoria aqui -- con
    # PRAGMA foreign_keys=ON eso rompe el INSERT en app_alerta_priorizada.
    if conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND name='dim_alerta_categoria'"
    ).fetchone():
        conn.executemany(
            'INSERT OR IGNORE INTO dim_alerta_categoria (alerta_codigo, nombre, tipo, descripcion, peso_base) '
            'VALUES (?, ?, ?, ?, ?)',
            [
                ('informe_campo', 'Hallazgo de informe de campo', 'informe_campo',
                 'Variable concreta detectada en un informe de campo (ausencia docente, infraestructura, etc.), con evidencia textual.', 3),
                ('preguntas_dinamicas', 'Preguntas dinamicas Nivel I/II', 'preguntas_dinamicas',
                 'Respuesta en Nivel I o II del instrumento dinamico (UGEL) para una IE.', 2),
            ],
        )

    conn.execute('''
        CREATE TABLE IF NOT EXISTS app_user (
            user_id TEXT PRIMARY KEY,
            nombre TEXT NOT NULL,
            username TEXT NOT NULL UNIQUE,
            email TEXT,
            password_salt TEXT NOT NULL,
            password_hash TEXT NOT NULL,
            rol TEXT NOT NULL DEFAULT 'especialista',
            especialista_id TEXT,
            activo INTEGER NOT NULL DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            actualizado_en DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    ''')

    user_existing = {
        row['name'] for row in conn.execute('PRAGMA table_info(app_user)').fetchall()
    }
    user_columns = {
        'nombre': 'TEXT',
        'username': 'TEXT',
        'email': 'TEXT',
        'password_salt': 'TEXT',
        'password_hash': 'TEXT',
        'rol': "TEXT NOT NULL DEFAULT 'especialista'",
        'especialista_id': 'TEXT',
        'cargo': 'TEXT',
        'activo': 'INTEGER NOT NULL DEFAULT 1',
        'creado_en': 'DATETIME DEFAULT CURRENT_TIMESTAMP',
        'actualizado_en': 'DATETIME DEFAULT CURRENT_TIMESTAMP',
    }
    for name, column_type in user_columns.items():
        if name not in user_existing:
            conn.execute(f'ALTER TABLE app_user ADD COLUMN {name} {column_type}')

    conn.execute('CREATE UNIQUE INDEX IF NOT EXISTS idx_app_user_username ON app_user(username)')
    conn.execute(
        "CREATE UNIQUE INDEX IF NOT EXISTS idx_app_user_email ON app_user(email) "
        "WHERE email IS NOT NULL AND email != ''"
    )

    conn.execute('''
        CREATE TABLE IF NOT EXISTS app_session (
            token TEXT PRIMARY KEY,
            user_id TEXT NOT NULL,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            expira_en DATETIME NOT NULL,
            FOREIGN KEY (user_id) REFERENCES app_user(user_id)
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_app_session_user ON app_session(user_id)')

    if conn.execute('SELECT COUNT(*) FROM app_user').fetchone()[0] == 0:
        create_app_user(
            conn,
            nombre='Administrador SUGKA',
            username='admin',
            email='admin@sugka.local',
            password=_seed_password('SEED_ADMIN_PASSWORD', 'admin'),
            rol='administrador',
        )
        first_specialist = conn.execute(
            'SELECT especialista_id, nombre FROM dim_especialista ORDER BY nombre LIMIT 1'
        ).fetchone()
        create_app_user(
            conn,
            nombre=first_specialist['nombre'] if first_specialist else 'Especialista Demo',
            username='especialista',
            email=generic_email_for_name(first_specialist['nombre']) if first_specialist else 'especialista@ugel-imaza.edu.pe',
            password=_seed_password('SEED_ESPECIALISTA_PASSWORD', 'especialista'),
            rol='especialista',
            especialista_id=first_specialist['especialista_id'] if first_specialist else None,
        )

    conn.execute('''
        CREATE TABLE IF NOT EXISTS institucion_resumen (
            codigo_modular TEXT PRIMARY KEY,
            campo_documentos INTEGER DEFAULT 0,
            simon_fichas INTEGER DEFAULT 0,
            simon_indicadores INTEGER DEFAULT 0,
            simon_promedio_nivel REAL,
            total_alertas_campo INTEGER DEFAULT 0,
            alertas_infraestructura_campo INTEGER DEFAULT 0,
            alertas_pedagogicas_campo INTEGER DEFAULT 0,
            coverage_category TEXT,
            campo_sin_simon INTEGER DEFAULT 0,
            simon_sin_campo INTEGER DEFAULT 0,
            ambas_fuentes INTEGER DEFAULT 0,
            priority_score REAL DEFAULT 0,
            priority_reason TEXT,
            campo_owners TEXT,
            campo_topics TEXT,
            simon_monitores TEXT,
            simon_docentes TEXT,
            simon_fechas TEXT,
            risk_score_infra REAL DEFAULT 0,
            risk_flags TEXT,
            FOREIGN KEY (codigo_modular) REFERENCES dim_institucion(codigo_modular)
        )
    ''')
    resumen_existing = {
        row['name'] for row in conn.execute('PRAGMA table_info(institucion_resumen)').fetchall()
    }
    resumen_columns = {
        'risk_score_infra': 'REAL DEFAULT 0',
        'risk_flags': 'TEXT',
        'riesgo_intervencion_simon': 'TEXT',
        'simon_promedio_general_desempeno': 'REAL',
        'simon_porcentaje_items_bajos': 'REAL',
        # Componentes del score de priorizacion del informe v5 (seccion 3.3).
        # Se guardan por separado, no solo el total, para que la app pueda
        # mostrarle al especialista de que esta hecho el puntaje de cada IE.
        'riesgo_simon_0_100': 'REAL',
        'riesgo_infra_0_100': 'REAL',
        'riesgo_campo_0_100': 'REAL',
        'brecha_evidencia_0_100': 'REAL',
        'score_priorizacion_ie': 'REAL',
        'prioridad_v5': 'TEXT',
        'campo_alertas_altas': 'INTEGER DEFAULT 0',
        'campo_alertas_medias': 'INTEGER DEFAULT 0',
        'campo_alertas_bajas': 'INTEGER DEFAULT 0',
        'campo_vars_distintas': 'INTEGER DEFAULT 0',
    }
    for name, column_type in resumen_columns.items():
        if name not in resumen_existing:
            conn.execute(f'ALTER TABLE institucion_resumen ADD COLUMN {name} {column_type}')

    conn.execute('''
        CREATE TABLE IF NOT EXISTS infraestructura_censo_2025 (
            codigo_modular TEXT PRIMARY KEY,
            codigo_modular_original TEXT,
            codigo_local TEXT,
            nombre_iiee TEXT,
            nivel_modalidad TEXT,
            distrito TEXT,
            centro_poblado TEXT,
            alumnos_censo INTEGER DEFAULT 0,
            edificaciones INTEGER DEFAULT 0,
            edificaciones_en_uso INTEGER DEFAULT 0,
            edificaciones_riesgo INTEGER DEFAULT 0,
            aulas INTEGER DEFAULT 0,
            aulas_en_uso INTEGER DEFAULT 0,
            alumnos_por_aula_en_uso REAL,
            alertas_infraestructura_campo INTEGER DEFAULT 0,
            risk_score_infra REAL DEFAULT 0,
            risk_flags TEXT,
            fuente TEXT DEFAULT 'censo_educativo_2025',
            actualizado_en DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_infra_censo_risk ON infraestructura_censo_2025(risk_score_infra)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_infra_censo_local ON infraestructura_censo_2025(codigo_local)')

    # ── Informes de campo: modelo relacional auditable ────────────────────
    # archivo_subido -> informe_campo_documento (1 doc puede citar varias IE)
    #   -> informe_campo_visita (1 fila = 1 visita a 1 IE)
    #     -> informe_campo_hallazgo (cada hallazgo/parrafo que sostiene una
    #        bandera, para poder auditar exactamente que evidencia la genero)
    conn.execute('''
        CREATE TABLE IF NOT EXISTS archivo_subido (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            tipo TEXT NOT NULL,
            nombre_archivo TEXT,
            mimetype TEXT,
            tamano_bytes INTEGER DEFAULT 0,
            contenido BLOB,
            hash_sha1 TEXT,
            subido_por TEXT,
            estado_fila INTEGER DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (subido_por) REFERENCES app_user(user_id)
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_archivo_subido_hash ON archivo_subido(hash_sha1)')

    # Vincula un archivo_subido (imagen/PDF que el especialista subio al OCR)
    # con la ficha SIMON que se confirmo a partir de el. Tabla puente en vez
    # de una FK directa en fichas_monitoreo porque una ficha puede venir de
    # varias fotos/paginas, y porque fichas_monitoreo es una tabla externa
    # (creada por el importador original) que no debemos tocar con columnas
    # nuevas mas alla de lo ya migrado en FICHA_MONITOREO_COLUMNS.
    conn.execute('''
        CREATE TABLE IF NOT EXISTS ficha_archivo (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            ficha_id INTEGER NOT NULL,
            archivo_subido_id INTEGER NOT NULL,
            orden INTEGER DEFAULT 0,
            estado_fila INTEGER DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (ficha_id) REFERENCES fichas_monitoreo(id),
            FOREIGN KEY (archivo_subido_id) REFERENCES archivo_subido(id)
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_ficha_archivo_ficha ON ficha_archivo(ficha_id)')

    # ── Core de profesores/directores y asistencia a talleres ─────────────
    # profesor es el padron maestro de personal: antes de esta tabla el docente
    # solo existia como texto libre dentro de cada ficha (fichas_monitoreo.docente),
    # sin forma de saber si "Perez Lopez, Ana" y "PEREZ LOPEZ ANA" eran la misma
    # persona. El DNI es la identidad real, pero se guarda NULL (no '') cuando
    # falta, para que el indice unico parcial tolere varios registros sin DNI.
    conn.execute('''
        CREATE TABLE IF NOT EXISTS profesor (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            dni TEXT,
            nombre_completo TEXT NOT NULL,
            celular TEXT,
            correo TEXT,
            cargo TEXT NOT NULL DEFAULT 'docente',
            condicion_laboral TEXT,
            codigo_modular_ie TEXT,
            fuente TEXT,
            estado_fila INTEGER DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            actualizado_en DATETIME,
            creado_por TEXT,
            modificado_por TEXT,
            FOREIGN KEY (codigo_modular_ie) REFERENCES dim_institucion(codigo_modular),
            FOREIGN KEY (creado_por) REFERENCES app_user(user_id),
            FOREIGN KEY (modificado_por) REFERENCES app_user(user_id)
        )
    ''')
    conn.execute('''
        CREATE UNIQUE INDEX IF NOT EXISTS idx_profesor_dni ON profesor(dni)
        WHERE dni IS NOT NULL AND dni != ''
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_profesor_nombre ON profesor(nombre_completo)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_profesor_ie ON profesor(codigo_modular_ie)')

    conn.execute('''
        CREATE TABLE IF NOT EXISTS taller_capacitacion (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre_evento TEXT,
            compromiso_indicador TEXT,
            nivel_educativo TEXT,
            dre TEXT,
            ugel TEXT,
            codigo_modular_sede TEXT,
            nombre_sede_detectado TEXT,
            fecha_inicio TEXT,
            fecha_fin TEXT,
            metodo_extraccion TEXT,
            parser TEXT,
            fuente TEXT DEFAULT 'app_upload',
            estado_fila INTEGER DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            actualizado_en DATETIME,
            creado_por TEXT,
            modificado_por TEXT,
            FOREIGN KEY (codigo_modular_sede) REFERENCES dim_institucion(codigo_modular),
            FOREIGN KEY (creado_por) REFERENCES app_user(user_id),
            FOREIGN KEY (modificado_por) REFERENCES app_user(user_id)
        )
    ''')

    # profesor_id es NOT NULL a proposito: el OCR no persiste nada hasta que el
    # especialista revisa y confirma cada asistente (mismo patron que fichas e
    # informes de campo). Un asistente sin identidad confirmada simplemente no
    # se envia a guardar; nunca se inserta una fila huerfana que nadie volvera
    # a vincular. Las columnas *_detectado guardan el texto crudo del OCR para
    # auditoria, aunque el vinculo real sea profesor_id.
    conn.execute('''
        CREATE TABLE IF NOT EXISTS taller_asistencia (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            taller_id INTEGER NOT NULL,
            profesor_id INTEGER NOT NULL,
            nombre_detectado TEXT,
            dni_detectado TEXT,
            grado_a_cargo TEXT,
            area TEXT,
            celular_detectado TEXT,
            correo_detectado TEXT,
            codigo_modular_ie_detectado TEXT,
            nombre_ie_detectado TEXT,
            match_metodo TEXT,
            match_score REAL,
            metadata_json TEXT,
            estado_fila INTEGER DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            actualizado_en DATETIME,
            creado_por TEXT,
            modificado_por TEXT,
            FOREIGN KEY (taller_id) REFERENCES taller_capacitacion(id),
            FOREIGN KEY (profesor_id) REFERENCES profesor(id),
            FOREIGN KEY (codigo_modular_ie_detectado) REFERENCES dim_institucion(codigo_modular),
            FOREIGN KEY (creado_por) REFERENCES app_user(user_id),
            FOREIGN KEY (modificado_por) REFERENCES app_user(user_id)
        )
    ''')
    taller_existing = {
        row['name'] for row in conn.execute('PRAGMA table_info(taller_capacitacion)').fetchall()
    }
    for columna, tipo in (('especialista_responsable', 'TEXT'), ('especialista_user_id', 'TEXT')):
        if columna not in taller_existing:
            conn.execute(f'ALTER TABLE taller_capacitacion ADD COLUMN {columna} {tipo}')

    conn.execute('CREATE INDEX IF NOT EXISTS idx_taller_asistencia_taller ON taller_asistencia(taller_id)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_taller_asistencia_profesor ON taller_asistencia(profesor_id)')

    conn.execute('''
        CREATE TABLE IF NOT EXISTS taller_archivo (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            taller_id INTEGER NOT NULL,
            archivo_subido_id INTEGER NOT NULL,
            orden INTEGER DEFAULT 0,
            estado_fila INTEGER DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (taller_id) REFERENCES taller_capacitacion(id),
            FOREIGN KEY (archivo_subido_id) REFERENCES archivo_subido(id)
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_taller_archivo_taller ON taller_archivo(taller_id)')

    conn.execute('''
        CREATE TABLE IF NOT EXISTS informe_campo_documento (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            archivo_subido_id INTEGER,
            nombre_archivo TEXT,
            ruta_relativa TEXT,
            hash_sha1 TEXT,
            tipo_documento TEXT,
            subtipo_documento TEXT,
            especialista_detectado TEXT,
            fecha_visita_inicio TEXT,
            fecha_visita_fin TEXT,
            metodo_extraccion TEXT,
            requiere_ocr INTEGER DEFAULT 0,
            longitud_texto_documento INTEGER DEFAULT 0,
            texto_extraido TEXT,
            fuente TEXT DEFAULT 'app_upload',
            subido_por TEXT,
            estado_fila INTEGER DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            actualizado_en DATETIME,
            FOREIGN KEY (archivo_subido_id) REFERENCES archivo_subido(id),
            FOREIGN KEY (subido_por) REFERENCES app_user(user_id)
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_campo_doc_hash ON informe_campo_documento(hash_sha1)')

    campo_variable_columns = ',\n            '.join(
        f'{var} INTEGER,\n            span_{var} TEXT' for var in CAMPO_CONCRETE_VARIABLES
    )
    conn.execute(f'''
        CREATE TABLE IF NOT EXISTS informe_campo_visita (
            visita_campo_id TEXT PRIMARY KEY,
            documento_id INTEGER,
            codigo_modular TEXT,
            nombre_ie_detectado TEXT,
            nombre_ie_padron TEXT,
            nivel_padron TEXT,
            centro_poblado_padron TEXT,
            distrito_padron TEXT,
            codlocal_padron TEXT,
            fecha_visita TEXT,
            anio_visita TEXT,
            mes_visita TEXT,
            especialista_detectado TEXT,
            {campo_variable_columns},
            n_variables_observadas INTEGER DEFAULT 0,
            variables_observadas TEXT,
            metodo_match_ie TEXT,
            flag_match_dudoso INTEGER DEFAULT 0,
            requiere_revision INTEGER DEFAULT 0,
            motivo_revision TEXT,
            fuente TEXT DEFAULT 'app_upload',
            estado_fila INTEGER DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            actualizado_en DATETIME,
            FOREIGN KEY (documento_id) REFERENCES informe_campo_documento(id)
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_campo_visita_codigo ON informe_campo_visita(codigo_modular)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_campo_visita_documento ON informe_campo_visita(documento_id)')

    campo_visita_existing = {
        row['name'] for row in conn.execute('PRAGMA table_info(informe_campo_visita)').fetchall()
    }
    if 'modificado_por' not in campo_visita_existing:
        conn.execute('ALTER TABLE informe_campo_visita ADD COLUMN modificado_por TEXT')

    # Flujo de revision de un informe de campo. 'requiere_revision' ya existia
    # pero era un booleano que no disparaba nada; estas columnas le dan estado,
    # responsable y rastro de quien pidio y quien cerro la revision.
    #   revision_estado: 'ok' | 'pendiente' | 'en_revision' | 'resuelto'
    # Mientras esta en 'pendiente' o 'en_revision' la visita NO alimenta las
    # alertas (ver rebuild_simon_operational_summary): un dato en duda no debe
    # mover los indicadores hasta que alguien lo confirme.
    for columna, tipo in (
        ('revision_estado', "TEXT DEFAULT 'ok'"),
        ('revision_asignado_a', 'TEXT'),
        ('revision_asignado_nombre', 'TEXT'),
        ('revision_solicitada_por', 'TEXT'),
        ('revision_solicitada_en', 'DATETIME'),
        ('revision_resuelto_por', 'TEXT'),
        ('revision_resuelto_en', 'DATETIME'),
        ('revision_nota', 'TEXT'),
    ):
        if columna not in campo_visita_existing:
            conn.execute(f'ALTER TABLE informe_campo_visita ADD COLUMN {columna} {tipo}')
    # Coherencia con el booleano viejo: lo que ya estaba marcado queda pendiente.
    conn.execute(
        "UPDATE informe_campo_visita SET revision_estado = 'pendiente' "
        "WHERE requiere_revision = 1 AND COALESCE(revision_estado, 'ok') = 'ok'"
    )
    conn.execute("UPDATE informe_campo_visita SET revision_estado = 'ok' WHERE revision_estado IS NULL")

    conn.execute('''
        CREATE TABLE IF NOT EXISTS informe_campo_hallazgo (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            visita_campo_id TEXT,
            documento_id INTEGER,
            codigo_modular_hallazgo TEXT,
            variable_detectada TEXT,
            tema TEXT,
            descripcion TEXT,
            evidencia_textual TEXT,
            requiere_revision INTEGER DEFAULT 0,
            motivo_revision TEXT,
            estado_fila INTEGER DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (visita_campo_id) REFERENCES informe_campo_visita(visita_campo_id),
            FOREIGN KEY (documento_id) REFERENCES informe_campo_documento(id)
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_campo_hallazgo_visita ON informe_campo_hallazgo(visita_campo_id)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_campo_hallazgo_variable ON informe_campo_hallazgo(variable_detectada)')

    # Validacion de hallazgos por especialista. El pipeline de OCR marca una
    # variable a partir de un fragmento de texto, y esa deteccion puede estar
    # mal (el fragmento habla del tema pero no reporta el problema, o es una
    # recomendacion y no una alerta). Estas columnas guardan el veredicto
    # humano para que la evidencia siga siendo trazable: un hallazgo
    # descartado deja de contar, pero no se pierde ni el fragmento ni la razon.
    hallazgo_existing = {
        fila[1] for fila in conn.execute('PRAGMA table_info(informe_campo_hallazgo)')
    }
    for columna, tipo in (
        ('validacion_estado', "TEXT DEFAULT 'sin_validar'"),
        ('validacion_comentario', 'TEXT'),
        ('validacion_por', 'TEXT'),
        ('validacion_en', 'DATETIME'),
        ('validacion_fuente', 'TEXT'),
    ):
        if columna not in hallazgo_existing:
            conn.execute(f'ALTER TABLE informe_campo_hallazgo ADD COLUMN {columna} {tipo}')

    conn.execute('''
        CREATE TABLE IF NOT EXISTS app_instrumento (
            instrumento_id INTEGER PRIMARY KEY AUTOINCREMENT,
            codigo TEXT NOT NULL UNIQUE,
            nombre TEXT NOT NULL,
            tipo TEXT NOT NULL,
            version TEXT DEFAULT '2026',
            activo INTEGER NOT NULL DEFAULT 1,
            fuente TEXT,
            estructura_json TEXT NOT NULL,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            actualizado_en DATETIME DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS app_instrumento_seccion (
            seccion_id INTEGER PRIMARY KEY AUTOINCREMENT,
            instrumento_id INTEGER NOT NULL,
            clave TEXT NOT NULL,
            nombre TEXT NOT NULL,
            orden INTEGER NOT NULL DEFAULT 0,
            activo INTEGER NOT NULL DEFAULT 1,
            FOREIGN KEY (instrumento_id) REFERENCES app_instrumento(instrumento_id)
        )
    ''')
    seccion_existing = {
        row['name'] for row in conn.execute('PRAGMA table_info(app_instrumento_seccion)').fetchall()
    }
    if 'activo' not in seccion_existing:
        conn.execute('ALTER TABLE app_instrumento_seccion ADD COLUMN activo INTEGER NOT NULL DEFAULT 1')

    conn.execute('''
        CREATE TABLE IF NOT EXISTS app_instrumento_pregunta_comentario (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            pregunta_id INTEGER NOT NULL,
            autor_id TEXT,
            texto TEXT NOT NULL,
            estado_fila INTEGER DEFAULT 1,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (pregunta_id) REFERENCES app_instrumento_pregunta(pregunta_id),
            FOREIGN KEY (autor_id) REFERENCES app_user(user_id)
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_pregunta_comentario_pregunta ON app_instrumento_pregunta_comentario(pregunta_id)')

    conn.execute('''
        CREATE TABLE IF NOT EXISTS app_instrumento_pregunta (
            pregunta_id INTEGER PRIMARY KEY AUTOINCREMENT,
            instrumento_id INTEGER NOT NULL,
            seccion_id INTEGER,
            codigo TEXT NOT NULL,
            item TEXT NOT NULL,
            tipo_respuesta TEXT NOT NULL DEFAULT 'nivel',
            niveles_json TEXT,
            opciones_json TEXT,
            orden INTEGER NOT NULL DEFAULT 0,
            activo INTEGER NOT NULL DEFAULT 1,
            FOREIGN KEY (instrumento_id) REFERENCES app_instrumento(instrumento_id),
            FOREIGN KEY (seccion_id) REFERENCES app_instrumento_seccion(seccion_id)
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS ficha_respuesta_instrumento (
            respuesta_id INTEGER PRIMARY KEY AUTOINCREMENT,
            ficha_id INTEGER NOT NULL,
            instrumento_codigo TEXT NOT NULL,
            instrumento_tipo TEXT,
            seccion_clave TEXT,
            pregunta_codigo TEXT NOT NULL,
            nivel INTEGER,
            respuesta_texto TEXT,
            observacion TEXT,
            metadata_json TEXT,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (ficha_id) REFERENCES fichas_monitoreo(id)
        )
    ''')
    conn.execute('''
        CREATE TABLE IF NOT EXISTS ficha_campo_instrumento (
            campo_respuesta_id INTEGER PRIMARY KEY AUTOINCREMENT,
            ficha_id INTEGER NOT NULL,
            instrumento_codigo TEXT NOT NULL,
            campo TEXT NOT NULL,
            valor TEXT,
            metadata_json TEXT,
            creado_en DATETIME DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (ficha_id) REFERENCES fichas_monitoreo(id)
        )
    ''')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_ficha_respuesta_ficha ON ficha_respuesta_instrumento(ficha_id)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_ficha_respuesta_pregunta ON ficha_respuesta_instrumento(instrumento_codigo, pregunta_codigo)')
    conn.execute('CREATE INDEX IF NOT EXISTS idx_ficha_campo_ficha ON ficha_campo_instrumento(ficha_id)')

    specialist_rows = conn.execute(
        'SELECT especialista_id, nombre FROM dim_especialista ORDER BY nombre'
    ).fetchall()
    for esp in specialist_rows:
        existing_user = conn.execute(
            'SELECT * FROM app_user WHERE especialista_id = ?',
            (esp['especialista_id'],),
        ).fetchone()
        base_username = slugify_username(esp['nombre'])
        base_email = generic_email_for_name(esp['nombre'])
        if existing_user:
            username = existing_user['username'] or base_username
            email = existing_user['email'] or base_email
            if username == 'especialista':
                username = base_username
            username, email = unique_account_identity(conn, username, email, existing_user['user_id'])
            conn.execute('''
                UPDATE app_user
                SET nombre = ?, username = ?, email = ?, rol = 'especialista',
                    actualizado_en = CURRENT_TIMESTAMP
                WHERE user_id = ?
            ''', (esp['nombre'], username, email, existing_user['user_id']))
        else:
            create_app_user(
                conn,
                nombre=esp['nombre'],
                username=base_username,
                email=base_email,
                password=_seed_password('SEED_ESPECIALISTA_PASSWORD', base_username),
                rol='especialista',
                especialista_id=esp['especialista_id'],
            )

    conn.commit()
    conn.close()

# La migracion de esquema y el seed de usuarios se ejecutan UNA sola vez al
# cargar el modulo (ya sea con `python api_server.py` o al ser importado por
# gunicorn), no en cada request. Antes se llamaba a ensure_db_schema() al
# inicio de casi cada endpoint, lo que repetia ~25 sentencias DDL/DML y el
# hashing de contraseñas en cada peticion -- costoso y, bajo trafico
# concurrente, una fuente de errores "database is locked" en SQLite.
with app.app_context():
    ensure_db_schema()

def first_text(data, *keys):
    for key in keys:
        value = data.get(key)
        if value is None:
            continue
        value = str(value).strip()
        if value:
            return value
    return ''

def safe_int(value, default=0):
    if value is None or value == '':
        return default
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value)
    match = re.search(r'-?\d+', str(value))
    return int(match.group(0)) if match else default

def safe_float(value, default=0.0):
    if value is None or value == '':
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
            match = re.search(r'-?\d+(?:[.,]\d+)?', str(value))
            return float(match.group(0).replace(',', '.')) if match else default

def canonical_codigo_modular(conn, value):
    code = re.sub(r'\D+', '', str(value or '').strip())
    if not code:
        return ''
    candidates = []
    for candidate in (code, code.zfill(7), f'0{code}'):
        if candidate and candidate not in candidates:
            candidates.append(candidate)
    for candidate in candidates:
        exists = conn.execute(
            'SELECT 1 FROM dim_institucion WHERE codigo_modular = ? LIMIT 1',
            (candidate,),
        ).fetchone()
        if exists:
            return candidate
    return code

def resolve_ficha_codigo_modular(conn, codigo_value, nombre_ie_value):
    """Resuelve el codigo modular de una ficha SIMON contra dim_institucion,
    igual de estricto que resolve_codigo_modular_strict (informes de campo):
    solo confirma un codigo, nunca lo adivina. Primero prueba el codigo
    (normalizado), y si no calza intenta por nombre de IE -- exacto primero,
    y por contencion (LIKE) solo si el resultado es unico, para no
    quedarnos con un match ambiguo entre varias IE con nombre parecido.
    Devuelve (codigo_resuelto, nombre_iiee_bd); codigo_resuelto es '' si no
    se pudo confirmar."""
    code = re.sub(r'\D+', '', str(codigo_value or '').strip())
    if code:
        for candidate in (code, code.zfill(7), f'0{code}'):
            row = conn.execute(
                'SELECT codigo_modular, nombre_iiee FROM dim_institucion WHERE codigo_modular = ?',
                (candidate,),
            ).fetchone()
            if row:
                return row['codigo_modular'], row['nombre_iiee']

    nombre = str(nombre_ie_value or '').strip()
    if nombre:
        row = conn.execute(
            'SELECT codigo_modular, nombre_iiee FROM dim_institucion WHERE UPPER(nombre_iiee) = UPPER(?) LIMIT 1',
            (nombre,),
        ).fetchone()
        if not row:
            like_rows = conn.execute(
                'SELECT codigo_modular, nombre_iiee FROM dim_institucion WHERE UPPER(nombre_iiee) LIKE UPPER(?) LIMIT 2',
                (f'%{nombre}%',),
            ).fetchall()
            row = like_rows[0] if len(like_rows) == 1 else None
        if row:
            return row['codigo_modular'], row['nombre_iiee']

    return '', ''

def normalize_level(value):
    if value is None or value == '':
        return 0
    if isinstance(value, (int, float)):
        level = int(value)
        return level if 0 <= level <= 4 else 0

    text = str(value).strip().upper()
    roman_levels = {'IV': 4, 'III': 3, 'II': 2, 'I': 1}
    for roman, score in roman_levels.items():
        if re.search(rf'\b{roman}\b', text):
            return score

    match = re.search(r'[0-4]', text)
    if match:
        return int(match.group(0))
    return 0

def calculate_promedio(ficha):
    levels = [ficha.get(k, 0) for k in ('a1', 'a2', 'a3', 'b1', 'b2', 'b3', 'b4', 'b5')]
    valid = [level for level in levels if level > 0]
    return round(sum(valid) / len(valid), 2) if valid else 0.0

SIMON_CODE_TO_FIELD = {
    'A-01': 'a1',
    'A-02': 'a2',
    'A-03': 'a3',
    'B-01': 'b1',
    'B-02': 'b2',
    'B-03': 'b3',
    'B-04': 'b4',
    'B-05': 'b5',
}
SIMON_FIELD_TO_CODE = {field: code for code, field in SIMON_CODE_TO_FIELD.items()}
SIMON_LEVEL_FIELDS = tuple(SIMON_CODE_TO_FIELD.values())

def iter_instrument_payload(data):
    instruments = data.get('instrumentos') or data.get('instrument_responses') or []
    if isinstance(instruments, dict):
        instruments = [instruments]
    return [item for item in instruments if isinstance(item, dict)]

def apply_instrument_data_to_ficha(ficha, data):
    compromisos = []
    for instrumento in iter_instrument_payload(data):
        tipo = str(instrumento.get('tipo') or instrumento.get('instrumento_tipo') or '').lower()
        codigo = str(instrumento.get('codigo') or instrumento.get('instrumento_codigo') or '').lower()
        is_simon = tipo == 'fija' or 'simon' in codigo or 'regional' in codigo

        for respuesta in instrumento.get('respuestas') or []:
            if not isinstance(respuesta, dict):
                continue
            pregunta_codigo = str(respuesta.get('pregunta_codigo') or respuesta.get('id') or '').upper()
            field = SIMON_CODE_TO_FIELD.get(pregunta_codigo)
            if is_simon and field:
                ficha[field] = normalize_level(
                    respuesta.get('nivel') or respuesta.get('respuesta') or respuesta.get('valor')
                )
                obs = first_text(respuesta, 'observacion', 'observaciones')
                if obs:
                    ficha[f'{field}_observacion'] = obs

        campos = instrumento.get('campos_finales') or instrumento.get('campos') or {}
        if isinstance(campos, dict):
            obs = first_text(campos, 'observaciones_recomendaciones', 'observaciones')
            if obs:
                ficha['observaciones'] = obs
                ficha['observaciones_recomendaciones'] = obs
            for key, value in campos.items():
                if str(key).startswith('compromiso_monitoreado') and str(value).strip():
                    compromisos.append(str(value).strip())

    if compromisos:
        joined = '\n'.join(compromisos)
        ficha['compromisos'] = joined
        ficha['compromisos_monitoreado'] = joined

def normalize_ficha(data, source='manual'):
    data = data if isinstance(data, dict) else {}
    ficha = {
        'source': first_text(data, 'source') or source,
        'file_name': first_text(data, 'file_name', 'filename'),
        'extraction_method': first_text(data, 'extraction_method', 'method'),
        'confidence': safe_float(data.get('confidence'), 0.0),
        'raw_text': first_text(data, 'raw_text'),
        'extracted_json': first_text(data, 'extracted_json'),
        'region': first_text(data, 'region'),
        'ugel': first_text(data, 'ugel'),
        'n_visita': first_text(data, 'n_visita', 'numero_visita', 'nro_visita', 'visita'),
        'codigo_modular': first_text(data, 'codigo_modular', 'cod_modular', 'codModular'),
        'nombre_ie': first_text(data, 'nombre_ie', 'ie', 'institucion_educativa', 'nombre_iiee'),
        'nivel_modalidad': first_text(data, 'nivel_modalidad', 'nivel', 'modalidad'),
        'director': first_text(data, 'director', 'directora', 'director_nombre'),
        'director_cel': first_text(data, 'director_cel', 'director_celular', 'cel_director'),
        'director_email': first_text(data, 'director_email', 'email_director'),
        'director_situacion_laboral': first_text(data, 'director_situacion_laboral', 'situacion_laboral_director'),
        'docente': first_text(data, 'docente', 'docente_nombre', 'nombre_docente'),
        'docente_dni': first_text(data, 'docente_dni', 'dni_docente', 'dni'),
        'docente_cel': first_text(data, 'docente_cel', 'docente_celular', 'cel_docente', 'cel'),
        'docente_email': first_text(data, 'docente_email', 'email_docente', 'email'),
        'docente_situacion_laboral': first_text(data, 'docente_situacion_laboral', 'situacion_laboral_docente'),
        'grado': first_text(data, 'grado'),
        'seccion': first_text(data, 'seccion', 'sección'),
        'nro_estudiantes': safe_int(first_text(data, 'nro_estudiantes', 'numero_estudiantes', 'n_estudiantes')),
        'area': first_text(data, 'area', 'área'),
        'competencia': first_text(data, 'competencia'),
        'titulo_sesion': first_text(data, 'titulo_sesion', 'titulo_de_la_sesion', 'sesion'),
        'monitor': first_text(data, 'monitor', 'monitor_nombre', 'especialista', 'especialista_monitor'),
        'monitor_dni': first_text(data, 'monitor_dni', 'dni_monitor'),
        'iged': first_text(data, 'iged'),
        'monitor_email': first_text(data, 'monitor_email', 'email_monitor'),
        'fecha_ejecucion': first_text(data, 'fecha_ejecucion', 'fecha'),
    }

    for key in ('a1', 'a2', 'a3', 'b1', 'b2', 'b3', 'b4', 'b5'):
        ficha[key] = normalize_level(data.get(key))

    for key in (
        'a1_observacion', 'a2_observacion', 'a3_observacion',
        'b1_observacion', 'b2_observacion', 'b3_observacion',
        'b4_observacion', 'b5_observacion',
    ):
        ficha[key] = first_text(data, key, key.replace('_observacion', '_obs'))

    observaciones = first_text(data, 'observaciones', 'observaciones_recomendaciones')
    compromisos = first_text(data, 'compromisos', 'compromisos_monitoreado')
    ficha['observaciones'] = observaciones
    ficha['observaciones_recomendaciones'] = first_text(data, 'observaciones_recomendaciones') or observaciones
    ficha['compromisos'] = compromisos
    ficha['compromisos_monitoreado'] = first_text(data, 'compromisos_monitoreado') or compromisos
    apply_instrument_data_to_ficha(ficha, data)
    ficha['promedio'] = safe_float(data.get('promedio', data.get('promedio_nivel')), 0.0) or calculate_promedio(ficha)
    ficha['promedio_nivel'] = str(ficha['promedio'])

    if not ficha['extracted_json']:
        ficha['extracted_json'] = json.dumps(data, ensure_ascii=False)

    return ficha

def save_instrument_responses(conn, ficha_id, data):
    for instrumento in iter_instrument_payload(data):
        codigo = first_text(instrumento, 'codigo', 'instrumento_codigo') or 'instrumento'
        tipo = first_text(instrumento, 'tipo', 'instrumento_tipo')
        respuestas = instrumento.get('respuestas') or []
        for respuesta in respuestas:
            if not isinstance(respuesta, dict):
                continue
            pregunta_codigo = first_text(respuesta, 'pregunta_codigo', 'id')
            if not pregunta_codigo:
                continue
            conn.execute('''
                INSERT INTO ficha_respuesta_instrumento (
                    ficha_id, instrumento_codigo, instrumento_tipo, seccion_clave,
                    pregunta_codigo, nivel, respuesta_texto, observacion, metadata_json
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                ficha_id,
                codigo,
                tipo,
                first_text(respuesta, 'seccion_clave', 'seccion'),
                pregunta_codigo,
                normalize_level(respuesta.get('nivel')),
                first_text(respuesta, 'respuesta_texto', 'respuesta', 'valor'),
                first_text(respuesta, 'observacion', 'observaciones'),
                json_dumps(respuesta),
            ))

        campos = instrumento.get('campos_finales') or instrumento.get('campos') or {}
        if isinstance(campos, dict):
            for campo, valor in campos.items():
                if valor is None or str(valor).strip() == '':
                    continue
                conn.execute('''
                    INSERT INTO ficha_campo_instrumento (
                        ficha_id, instrumento_codigo, campo, valor, metadata_json
                    )
                    VALUES (?, ?, ?, ?, ?)
                ''', (
                    ficha_id,
                    codigo,
                    str(campo),
                    str(valor),
                    json_dumps({'campo': campo, 'valor': valor, 'tipo': tipo}),
                ))

def save_ficha_archivos(conn, ficha_id, archivos):
    """Persiste en BD las imagenes/PDF que el especialista subio para el OCR
    de esta ficha (base64 en 'file_b64'), una fila de archivo_subido por
    archivo mas su vinculo en ficha_archivo -- asi lo pedido explicitamente
    de no perder el documento original tras extraerle el texto."""
    if not archivos:
        return
    try:
        user_id = g.current_user['user_id'] if getattr(g, 'current_user', None) else None
    except RuntimeError:
        user_id = None
    for idx, item in enumerate(archivos):
        if not isinstance(item, dict) or not item.get('file_b64'):
            continue
        try:
            contenido = base64.b64decode(item['file_b64'])
        except Exception:
            continue
        cur = conn.execute('''
            INSERT INTO archivo_subido (tipo, nombre_archivo, mimetype, tamano_bytes, contenido, hash_sha1, subido_por)
            VALUES ('ficha_ocr', ?, ?, ?, ?, ?, ?)
        ''', (
            item.get('filename', ''), item.get('mimetype', ''), len(contenido),
            contenido, item.get('hash_sha1', ''), user_id,
        ))
        conn.execute('''
            INSERT INTO ficha_archivo (ficha_id, archivo_subido_id, orden) VALUES (?, ?, ?)
        ''', (ficha_id, cur.lastrowid, idx))

def save_ficha_to_db(conn, data):
    data = data if isinstance(data, dict) else {}
    ficha = normalize_ficha(data, source=first_text(data, 'source') or 'manual')
    if getattr(g, 'current_user', None):
        ficha['creado_por'] = g.current_user['user_id']
    # estado_fila esta en la lista explicita de columnas del INSERT (mas abajo),
    # asi que si no se fija aqui se inserta NULL en vez de aplicar el DEFAULT 1
    # de la columna (los defaults de SQL solo aplican cuando la columna se omite
    # del INSERT, no cuando se manda explicitamente NULL) -- eso hacia que toda
    # ficha nueva quedara marcada como "eliminada" para el panel de Registros.
    ficha['estado_fila'] = 1
    columns = list(FICHA_MONITOREO_COLUMNS.keys())
    placeholders = ', '.join('?' for _ in columns)
    cur = conn.execute(
        f'INSERT INTO fichas_monitoreo ({", ".join(columns)}) VALUES ({placeholders})',
        [ficha.get(column) for column in columns]
    )
    ficha['id'] = cur.lastrowid
    save_instrument_responses(conn, cur.lastrowid, data)
    save_ficha_archivos(conn, cur.lastrowid, data.get('archivos'))
    return ficha

# Los 8 indicadores pedagogicos del instrumento SIMON, agrupados en sus dos
# dimensiones (AR01 preparacion / AR02 ensenanza). No son variables inventadas
# por el proyecto: son las preguntas del instrumento de monitoreo aplicado por
# los especialistas, aqui en su version resumida para la ficha tecnica del dashboard.
SIMON_INDICATOR_DEFINITIONS = [
    {'codigo_reporte': 'AR01.1', 'codigo_instrumento': 'A-01', 'campo': 'a1',
     'dimension': 'Preparacion para el aprendizaje',
     'pregunta': 'La planificacion curricular evidencia conocimiento de estandares y alineacion con el contexto?'},
    {'codigo_reporte': 'AR01.2', 'codigo_instrumento': 'A-02', 'campo': 'a2',
     'dimension': 'Preparacion para el aprendizaje',
     'pregunta': 'Las situaciones de aprendizaje son desafiantes, factibles y coherentes con los propositos?'},
    {'codigo_reporte': 'AR01.3', 'codigo_instrumento': 'A-03', 'campo': 'a3',
     'dimension': 'Preparacion para el aprendizaje',
     'pregunta': 'Existe coherencia entre los criterios de evaluacion y los propositos de aprendizaje?'},
    {'codigo_reporte': 'AR02.1', 'codigo_instrumento': 'B-01', 'campo': 'b1',
     'dimension': 'Ensenanza para el aprendizaje',
     'pregunta': 'El docente promueve el interes de los estudiantes y el sentido de lo aprendido?'},
    {'codigo_reporte': 'AR02.2', 'codigo_instrumento': 'B-02', 'campo': 'b2',
     'dimension': 'Ensenanza para el aprendizaje',
     'pregunta': 'Las actividades estimulan la formulacion creativa, comprension de principios o relaciones conceptuales?'},
    {'codigo_reporte': 'AR02.3', 'codigo_instrumento': 'B-03', 'campo': 'b3',
     'dimension': 'Ensenanza para el aprendizaje',
     'pregunta': 'El docente monitorea avances/dificultades y brinda retroalimentacion formativa?'},
    {'codigo_reporte': 'AR02.4', 'codigo_instrumento': 'B-04', 'campo': 'b4',
     'dimension': 'Ensenanza para el aprendizaje',
     'pregunta': 'El docente se comunica con respeto, calidez, y atiende necesidades afectivas/fisicas?'},
    {'codigo_reporte': 'AR02.5', 'codigo_instrumento': 'B-05', 'campo': 'b5',
     'dimension': 'Ensenanza para el aprendizaje',
     'pregunta': 'El docente establece normas de convivencia claras y las hace cumplir formativamente?'},
]

# Regla documentada de riesgo_intervencion (ver simon_documented_risk arriba),
# repetida aqui en forma de tabla para mostrarla como ficha tecnica en el dashboard.
SIMON_RISK_RULE = [
    {'nivel': 'Alto',
     'condicion': 'Al menos un indicador en Nivel I, o promedio_general_desempeno < 2.25, o porcentaje_items_bajos >= 75%',
     'justificacion': 'Cualquiera de estas condiciones indica una visita con desempeno critico o mayoritariamente bajo.'},
    {'nivel': 'Medio',
     'condicion': 'No cumple Alto, y (promedio_general_desempeno < 2.75 o cantidad_items_bajos >= 3)',
     'justificacion': 'Desempeno no critico, pero con senales que ameritan seguimiento.'},
    {'nivel': 'Bajo',
     'condicion': 'No cumple condiciones de Alto ni Medio',
     'justificacion': 'Desempeno relativo mejor respecto a los umbrales definidos.'},
]

# Ejemplo ilustrativo (ficha real C-345) que muestra por que el umbral porcentual
# puede activar "Alto" incluso sin ningun indicador en Nivel I ni promedio critico.
SIMON_RISK_EXAMPLE = {
    'ficha': 'C-345',
    'items': [
        {'codigo_reporte': 'AR01.1', 'nivel': 2}, {'codigo_reporte': 'AR01.2', 'nivel': 2}, {'codigo_reporte': 'AR01.3', 'nivel': 2},
        {'codigo_reporte': 'AR02.1', 'nivel': 2}, {'codigo_reporte': 'AR02.2', 'nivel': 2}, {'codigo_reporte': 'AR02.3', 'nivel': 2},
        {'codigo_reporte': 'AR02.4', 'nivel': 3}, {'codigo_reporte': 'AR02.5', 'nivel': 3},
    ],
    'calculo': 'promedio_general_desempeno = (2x6 + 3x2) / 8 = 2.25 | nivel_minimo_obtenido = 2 (no hay Nivel I) | cantidad_items_bajos (Nivel I o II) = 6 -> porcentaje_items_bajos = 75%',
    'resultado': 'Alto',
    'explicacion': (
        'No hay Nivel I y el promedio no es estrictamente menor a 2.25 (es igual), pero porcentaje_items_bajos >= 75% '
        'si se cumple exactamente: esta ficha se clasifica como Alto aunque ningun indicador individual llego al nivel mas bajo.'
    ),
}

# ─── SCORE DE PRIORIZACION (informe final Sugka Lab v5, seccion 3.3) ──────────
# Regla ponderada explicable, 0 a 100, que une las tres fuentes por
# codigo_modular. Se implementa exactamente como la documenta el informe para
# que app e informe den el mismo numero: antes la app usaba una formula aditiva
# propia (calculate_risk v0.1) que no coincidia con lo publicado.
#
# Cada fuente se lleva primero a una escala 0-100 para que el peso dependa solo
# de la formula y no de la escala original de cada variable.

# La app ya usa Critico/Alto/Medio/Bajo en colores, filtros y leyendas; el
# informe nombra los mismos cuatro niveles como Critica/Alta/Media/Baja.
RIESGO_NIVEL_POR_PRIORIDAD = {'Critica': 'Critico', 'Alta': 'Alto', 'Media': 'Medio', 'Baja': 'Bajo'}

SCORE_PESOS = {'simon': 0.45, 'infra': 0.30, 'campo': 0.20, 'brecha': 0.05}

# Maximo observable de risk_score_infra segun el informe (seccion 3.3.2).
INFRA_SCORE_MAX = 6.5

# Puntos por alerta de campo segun severidad. Mantienen la proporcion 2:1 entre
# severidades consecutivas: dos medias ~ una alta, dos bajas ~ una media.
CAMPO_PUNTOS_SEVERIDAD = {'alta': 25, 'media': 12, 'baja': 6}


def riesgo_simon_0_100(promedio, pct_bajos):
    """Desempeno pedagogico en escala 0-100 (mas alto = peor).

    El promedio se invierte porque un promedio bajo es peor, y pesa mas (0.60)
    que el porcentaje de items bajos (0.40) porque resume los 8 items en una
    cifra continua; el porcentaje evita que un item muy bueno tape varios malos.
    """
    promedio = safe_float(promedio, 0.0)
    if promedio <= 0:
        return 0.0
    pct = safe_float(pct_bajos, 0.0)
    if pct <= 1.0:
        pct *= 100  # se guarda como fraccion; el informe lo usa como porcentaje
    promedio_bajo = ((4 - promedio) / 3) * 100
    return round(max(0.0, min(100.0, 0.60 * promedio_bajo + 0.40 * pct)), 1)


def riesgo_infra_0_100(risk_score_infra):
    """Infraestructura en escala 0-100, reescalando sobre su maximo observable."""
    valor = safe_float(risk_score_infra, 0.0)
    if valor <= 0:
        return 0.0
    return round(min(valor / INFRA_SCORE_MAX * 100, 100), 1)


def riesgo_campo_0_100(altas, medias, bajas):
    """Alertas de campo ponderadas por severidad, con tope en 100."""
    total = (safe_int(altas, 0) * CAMPO_PUNTOS_SEVERIDAD['alta']
             + safe_int(medias, 0) * CAMPO_PUNTOS_SEVERIDAD['media']
             + safe_int(bajas, 0) * CAMPO_PUNTOS_SEVERIDAD['baja'])
    return round(min(total, 100), 1)


def brecha_evidencia_0_100(tiene_simon, tiene_infra, tiene_campo):
    """Completitud de la evidencia, NO riesgo.

    Mide cuantas de las tres fuentes faltan. Una brecha alta no dice que la IE
    este bien: dice que no se la puede evaluar con la misma certeza que a las
    demas. Por eso pesa solo 0.05 en el score final.
    """
    faltantes = 3 - sum(1 for tiene in (tiene_simon, tiene_infra, tiene_campo) if tiene)
    return round((faltantes / 3) * 100, 1)


def clasificar_prioridad(score):
    """Cortes del informe v5 (seccion 3.3.4)."""
    valor = safe_float(score, 0.0)
    if valor >= 75:
        return 'Critica'
    if valor >= 60:
        return 'Alta'
    if valor >= 40:
        return 'Media'
    return 'Baja'


def calcular_score_priorizacion(simon_0_100, infra_0_100, campo_0_100, brecha_0_100):
    """Combina los cuatro componentes con los pesos del informe v5."""
    score = (SCORE_PESOS['simon'] * safe_float(simon_0_100, 0.0)
             + SCORE_PESOS['infra'] * safe_float(infra_0_100, 0.0)
             + SCORE_PESOS['campo'] * safe_float(campo_0_100, 0.0)
             + SCORE_PESOS['brecha'] * safe_float(brecha_0_100, 0.0))
    return round(min(max(score, 0.0), 100.0), 1)


def desglose_score(fila):
    """Componentes + score + prioridad de una fila de institucion_resumen.

    Se usa tanto al recalcular (rebuild) como al leer filas viejas que todavia
    no tienen las columnas guardadas, para que la app nunca muestre un score
    vacio mientras no se haya corrido un recalculo.
    """
    simon = fila.get('riesgo_simon_0_100')
    if simon is None:
        simon = riesgo_simon_0_100(
            fila.get('simon_promedio_general_desempeno') or fila.get('simon_promedio_nivel'),
            fila.get('simon_porcentaje_items_bajos'),
        )
    infra = fila.get('riesgo_infra_0_100')
    if infra is None:
        infra = riesgo_infra_0_100(fila.get('risk_score_infra'))
    campo = fila.get('riesgo_campo_0_100')
    if campo is None:
        campo = riesgo_campo_0_100(
            fila.get('campo_alertas_altas'), fila.get('campo_alertas_medias'),
            fila.get('campo_alertas_bajas'),
        )
    brecha = fila.get('brecha_evidencia_0_100')
    if brecha is None:
        brecha = brecha_evidencia_0_100(
            safe_int(fila.get('simon_fichas'), 0) > 0,
            'Censo' in str(fila.get('coverage_category') or ''),
            safe_int(fila.get('campo_documentos'), 0) > 0,
        )
    score = fila.get('score_priorizacion_ie')
    if score is None:
        score = calcular_score_priorizacion(simon, infra, campo, brecha)
    return {
        'riesgo_simon_0_100': round(safe_float(simon, 0.0), 1),
        'riesgo_infra_0_100': round(safe_float(infra, 0.0), 1),
        'riesgo_campo_0_100': round(safe_float(campo, 0.0), 1),
        'brecha_evidencia_0_100': round(safe_float(brecha, 0.0), 1),
        'score_priorizacion_ie': round(safe_float(score, 0.0), 1),
        'prioridad': fila.get('prioridad_v5') or clasificar_prioridad(score),
    }


RISK_MODEL = {
    'name': 'Score de priorizacion Sugka Lab (informe final v5)',
    'type': 'Regla ponderada explicable, 0 a 100',
    'formula': ('score_priorizacion_ie = 0.45 · SIMON + 0.30 · Infraestructura '
                '+ 0.20 · Campo + 0.05 · Brecha de evidencia'),
    'description': (
        'Une las tres fuentes por codigo_modular y resume la prioridad de atencion de cada IE '
        'en una sola escala de 0 a 100. No es un modelo entrenado: el informe descarta la regresion '
        'por falta de muestra y entrega esta regla ponderada, donde cada peso puede revisarse y '
        'discutirse uno por uno. El score prioriza; no reemplaza el criterio del especialista.'
    ),
    'levels': [
        {'level': 'Critico', 'range': '75-100', 'action': 'Visita prioritaria y plan de accion inmediato'},
        {'level': 'Alto', 'range': '60-74', 'action': 'Seguimiento focalizado en la siguiente ronda'},
        {'level': 'Medio', 'range': '40-59', 'action': 'Monitoreo regular con revision mensual'},
        {'level': 'Bajo', 'range': '0-39', 'action': 'Seguimiento ordinario'},
    ],
    'weights': [
        {
            'factor': 'Riesgo SIMON',
            'weight': '0.45',
            'rule': (
                'Unica fuente que observa directamente el desempeno pedagogico en el aula, con un '
                'instrumento estandarizado aplicado por un especialista: la evidencia mas cercana al '
                'problema que se busca priorizar.'
            ),
            'calculo': 'promedio invertido ((4 - promedio) / 3 × 100) × 0.60 + % de items bajos × 0.40',
        },
        {
            'factor': 'Riesgo de infraestructura',
            'weight': '0.30',
            'rule': (
                'Unica fuente con cobertura completa de las 382 IE y con registro estructurado '
                '(Censo/Padron), no extraccion de texto libre: la senal mas objetiva.'
            ),
            'calculo': 'risk_score_infra / 6.5 × 100, con tope en 100',
        },
        {
            'factor': 'Riesgo de Campo',
            'weight': '0.20',
            'rule': (
                'Aporta senales que ni SIMON ni el Censo observan (violencia, ausencia docente, gestion '
                'directiva, acceso), pero su cobertura es parcial y depende de extraccion automatica de '
                'texto. Pesa menos por eso, no porque sus hallazgos importen menos.'
            ),
            'calculo': 'alertas altas × 25 + medias × 12 + bajas × 6, con tope en 100',
        },
        {
            'factor': 'Brecha de evidencia',
            'weight': '0.05',
            'rule': (
                'No mide riesgo, mide completitud. Se mantiene bajo a proposito para que falta de datos '
                'nunca se confunda con alto riesgo, pero presente para que una IE sin evidencia no quede '
                'igual que una ya confirmada como de bajo riesgo en las tres fuentes.'
            ),
            'calculo': 'fuentes faltantes / 3 × 100',
        },
    ],
    'nota_pesos': (
        'Los cuatro pesos son una propuesta inicial sustentada en el razonamiento anterior, no en un '
        'ajuste estadistico: hoy no hay muestra suficiente para ajustarlos con confianza. Estan '
        'documentados componente por componente para que un especialista pueda cuestionarlos y proponer '
        'un ajuste antes de un uso institucional.'
    ),
    'fuente': 'Informe final Sugka Lab v5, seccion 3.3',
}

ALERT_RULES = [
    {
        'code': 'COBERTURA_EVIDENCIA',
        'name': 'Cobertura de evidencia',
        'severity': 'alta',
        'condition': 'coverage_category = Sin evidencia / Solo campo / Solo SIMON',
        'score': 25,
        'description': 'La matriz aumenta la prioridad cuando la IE no tiene evidencia reciente o cuando solo existe una fuente para contrastar el seguimiento.',
        'focus': 'Completar la evidencia faltante antes de cerrar la priorizacion.',
    },
    {
        'code': 'SIMON_REFUERZO',
        'name': 'Desempeno SIMON en refuerzo',
        'severity': 'media',
        'condition': 'simon_promedio_nivel < 3 o respuestas en Nivel I/II',
        'score': 30,
        'description': 'Se prioriza cuando el promedio SIMON queda por debajo de Nivel III o cuando hay concentracion de respuestas en Nivel I y II.',
        'focus': 'Revisar preparacion, ensenanza y compromisos de mejora docente.',
    },
    {
        'code': 'ALERTAS_PEDAGOGICAS',
        'name': 'Alertas pedagogicas acumuladas',
        'severity': 'alta',
        'condition': 'alertas_pedagogicas_campo >= 8 / >= 20',
        'score': 18,
        'description': 'Las respuestas SIMON en Nivel I/II alimentan alertas pedagogicas; a mayor acumulacion, mayor prioridad de acompanamiento.',
        'focus': 'Priorizar acompanamiento pedagogico y compromisos verificables.',
    },
    {
        'code': 'INFRA_CENSO_2025',
        'name': 'Riesgo censal de infraestructura',
        'severity': 'alta',
        'condition': 'risk_score_infra > 0 o risk_flags activos',
        'score': 53,
        'description': 'Usa risk_score_infra del Censo Educativo 2025 y banderas como riesgo estructural, aulas no usadas, alta densidad o falta de senalizacion.',
        'focus': 'Coordinar respuesta con gestion institucional e infraestructura.',
    },
    {
        'code': 'ALERTAS_INFRA_SERVICIOS',
        'name': 'Alertas de infraestructura/servicios',
        'severity': 'media',
        'condition': 'alertas_infraestructura_campo >= 4 / >= 10',
        'score': 8,
        'description': 'Las menciones de conectividad, energia, agua/saneamiento, transporte/acceso y servicios incrementan la prioridad operativa.',
        'focus': 'Atender brechas de servicios que afectan el funcionamiento de la IE.',
    },
]

def risk_level(score):
    if score >= 75:
        return 'Critico'
    if score >= 60:
        return 'Alto'
    if score >= 40:
        return 'Medio'
    return 'Bajo'

def calculate_risk(row):
    """Prioriza una IE con el score del informe v5 (seccion 3.3).

    Antes esta funcion usaba una formula aditiva propia de la app que no
    coincidia con la publicada en el informe: mismos niveles y mismos cortes,
    pero otro calculo. Ahora la app usa la del informe, que es la regla
    ponderada explicable (no hay modelo entrenado: el informe lo descarta
    explicitamente por falta de muestra).

    Se mantienen las llaves `risk_score` / `risk_level` / `risk_reasons` porque
    el resto de la app y el frontend ya las consumen; `risk_level` conserva los
    nombres Critico/Alto/Medio/Bajo que usan los colores y filtros existentes.
    """
    desglose = desglose_score(dict(row))
    score = desglose['score_priorizacion_ie']

    # Razones legibles, ordenadas por cuanto aporta cada componente al total:
    # el especialista debe poder ver POR QUE una IE quedo priorizada.
    aportes = [
        ('SIMON', desglose['riesgo_simon_0_100'], SCORE_PESOS['simon'],
         'desempeno pedagogico observado en aula'),
        ('Infraestructura', desglose['riesgo_infra_0_100'], SCORE_PESOS['infra'],
         'condiciones del local segun el Censo 2025'),
        ('Campo', desglose['riesgo_campo_0_100'], SCORE_PESOS['campo'],
         'alertas registradas en visitas'),
    ]
    aportes.sort(key=lambda a: a[1] * a[2], reverse=True)
    reasons = [
        f'{nombre} {valor}/100 (peso {int(peso * 100)}%): {detalle}'
        for nombre, valor, peso, detalle in aportes
        if valor > 0
    ]
    brecha = desglose['brecha_evidencia_0_100']
    if brecha > 0:
        faltan = int(round(brecha / 100 * 3))
        reasons.append(
            f'Brecha de evidencia {brecha}/100: falta {faltan} de 3 fuentes, '
            'el puntaje se calculo con informacion parcial'
        )
    if not reasons:
        reasons = ['Sin senales de riesgo en las tres fuentes']

    return {
        'risk_score': score,
        'risk_level': RIESGO_NIVEL_POR_PRIORIDAD.get(desglose['prioridad'], 'Bajo'),
        'risk_reasons': reasons[:4],
        # Se exponen los componentes para que la vista pueda mostrar el desglose
        # y la leyenda sin recalcular nada.
        'score_componentes': desglose,
    }


SEMAFORO_PEDAGOGICO = [
    {
        'excel_value': 'bajo rendimiento',
        'label': 'Critico',
        'color': 'Rojo',
        'status_key': 'danger',
        'action': 'Intervencion inmediata y acompanamiento focalizado.',
    },
    {
        'excel_value': 'regular',
        'label': 'Alerta',
        'color': 'Amarillo',
        'status_key': 'warning',
        'action': 'Refuerzo especifico y seguimiento en la siguiente visita.',
    },
    {
        'excel_value': 'bueno',
        'label': 'Estable',
        'color': 'Verde',
        'status_key': 'ok',
        'action': 'Mantener seguimiento ordinario y buenas practicas.',
    },
]

INFRA_FLAG_DESCRIPTIONS = {
    'sin_codigo_local_para_cruce': 'La IE no tenia codigo de local suficiente para cruzar con infraestructura.',
    'sin_registros_edificaciones': 'Tiene codigo local, pero no se encontraron registros en edificaciones.',
    'sin_registros_aulas': 'Tiene codigo local, pero no se encontraron registros de aulas.',
    'edificacion_con_riesgo_estructural': 'Al menos una edificacion presenta riesgo estructural o de colapso.',
    'aulas_registradas_no_en_uso': 'Existen aulas registradas, pero ninguna aparece en uso.',
    'alta_densidad_alumnos_por_aula_en_uso': 'La relacion alumnos por aula en uso es mayor a 35.',
    'conservacion_mala_en_puertas_o_ventanas': 'Alguna aula tiene puertas o ventanas en mal estado.',
    'aulas_en_uso_sin_senalizacion_completa': 'Hay aulas en uso, pero no todas tienen senalizacion de seguridad completa.',
}

INFRA_DEFINITIONS = {
    'alertas_infraestructura_campo': (
        'Numero de alertas de infraestructura o servicios detectadas en informes de campo vinculados a la IE. '
        'Incluye conectividad, energia/luz, agua/saneamiento y transporte/acceso.'
    ),
    'risk_flags': (
        'Etiquetas explicativas del riesgo detectado en infraestructura, aulas, seguridad o disponibilidad de datos.'
    ),
    'risk_score_infra': (
        'Puntaje de priorizacion de riesgo de infraestructura calculado con evidencias censales de edificios/aulas '
        'y alertas de campo. A mayor puntaje, mayor prioridad de atencion.'
    ),
    'formula': [
        '+3.0 si hay edificacion con riesgo estructural.',
        '+2.0 si no hay registros de edificaciones.',
        '+2.0 si no hay registros de aulas.',
        '+1.5 si alumnos por aula en uso > 35.',
        '+1.0 si hay puertas o ventanas en mal estado.',
        '+1.0 si hay aulas en uso sin senalizacion completa.',
        '+min(alertas_infraestructura_campo, 50) / 25, hasta 2 puntos por alertas de campo.',
    ],
    'flags': INFRA_FLAG_DESCRIPTIONS,
}

def clamp_score(value, low=0, high=100):
    return max(low, min(high, int(round(value))))

def simon_level_state(level):
    level = safe_float(level, 0.0)
    if level < 2.5:
        return dict(SEMAFORO_PEDAGOGICO[0])
    if level < 3.5:
        return dict(SEMAFORO_PEDAGOGICO[1])
    return dict(SEMAFORO_PEDAGOGICO[2])

def simon_level_percent(level):
    return clamp_score(safe_float(level, 0.0) / 4 * 100)

def simon_nivel_romano(promedio):
    """'Nivel I'..'Nivel IV' mas cercano a un promedio 1-4, para el pill del KPI."""
    promedio = safe_float(promedio, 0.0)
    if promedio <= 0:
        return None
    numero = max(1, min(4, round(promedio)))
    return f'Nivel {["I", "II", "III", "IV"][numero - 1]}'

def get_simon_question_labels(conn):
    rows = conn.execute('''
        SELECT p.codigo, p.item, s.nombre AS seccion_nombre, s.clave AS seccion_clave
        FROM app_instrumento_pregunta p
        JOIN app_instrumento i ON i.instrumento_id = p.instrumento_id
        LEFT JOIN app_instrumento_seccion s ON s.seccion_id = p.seccion_id
        WHERE i.codigo = 'simon_docente_2026'
        ORDER BY p.orden, p.pregunta_id
    ''').fetchall()
    return {
        row['codigo']: {
            'item': row['item'],
            'seccion_nombre': row['seccion_nombre'],
            'seccion_clave': row['seccion_clave'],
        }
        for row in rows
    }

def simon_values_from_row(row):
    values = []
    for field in SIMON_LEVEL_FIELDS:
        level = safe_int(row.get(field), 0)
        if level > 0:
            values.append(level)
    return values

def _ficha_es_mas_reciente(candidata, actual):
    """True si `candidata` reemplaza a `actual` como la visita mas reciente
    del mismo docente: compara fecha_ejecucion, luego n_visita, luego el id
    de insercion como ultimo desempate (fichas identicas cargadas dos veces)."""
    fecha_a, fecha_b = str(candidata.get('fecha_ejecucion') or ''), str(actual.get('fecha_ejecucion') or '')
    if fecha_a != fecha_b:
        return fecha_a > fecha_b
    visita_a, visita_b = safe_int(candidata.get('n_visita'), 0), safe_int(actual.get('n_visita'), 0)
    if visita_a != visita_b:
        return visita_a > visita_b
    return safe_int(candidata.get('id'), 0) > safe_int(actual.get('id'), 0)

def simon_latest_ficha_por_docente(conn):
    """Una fila de fichas_monitoreo por cada (codigo_modular, docente): si el
    mismo docente tiene mas de una ficha -- una visita de seguimiento real, o
    simplemente una carga duplicada -- se conserva solo la mas reciente.

    Esto alimenta TODAS las metricas agregadas de SIMON (KPIs, graficos,
    ranking de priorizacion): un docente que ya mejoro no debe seguir
    penalizando el promedio ni el riesgo con una evaluacion vieja y superada,
    y una ficha cargada dos veces no debe contarse dos veces."""
    rows = conn.execute('''
        SELECT * FROM fichas_monitoreo WHERE COALESCE(codigo_modular, '') != ''
    ''').fetchall()
    latest = {}
    total_bruto = 0
    for row in rows:
        item = dict(row)
        total_bruto += 1
        item['codigo_modular'] = canonical_codigo_modular(conn, item.get('codigo_modular'))
        key = (item['codigo_modular'], (item.get('docente') or '').strip().upper())
        actual = latest.get(key)
        if actual is None or _ficha_es_mas_reciente(item, actual):
            latest[key] = item
    resultado = list(latest.values())
    resultado_meta = {'total_bruto': total_bruto, 'total_unico': len(resultado), 'duplicados_excluidos': total_bruto - len(resultado)}
    return resultado, resultado_meta

# Regla oficial de "riesgo_intervencion" documentada en el subproyecto de datos SIMON
# (docs/02_criterios_riesgo_intervencion.md), validada 24/24 contra el dataset real
# de fichas SIMON de UGEL Imaza. Se recalcula en vivo a partir de los niveles (a1..b5)
# realmente registrados en cada ficha; no es un valor inventado ni importado como texto.
SIMON_RISK_ORDER = {'Bajo': 0, 'Medio': 1, 'Alto': 2}

def simon_documented_risk(values):
    """values: lista de niveles 1-4 (Nivel I..IV) de una o mas fichas SIMON.
    Implementa exactamente la regla de docs/02_criterios_riesgo_intervencion.md."""
    values = [v for v in values if v and v > 0]
    if not values:
        return None
    total = len(values)
    promedio = sum(values) / total
    nivel_minimo = min(values)
    bajos = sum(1 for v in values if v <= 2)
    porcentaje_bajos = bajos / total

    if nivel_minimo <= 1 or promedio < 2.25 or porcentaje_bajos >= 0.75:
        etiqueta = 'Alto'
    elif promedio < 2.75 or bajos >= 3:
        etiqueta = 'Medio'
    else:
        etiqueta = 'Bajo'

    return {
        'riesgo_intervencion': etiqueta,
        'promedio_general_desempeno': round(promedio, 2),
        'porcentaje_items_bajos': round(porcentaje_bajos, 2),
        'cantidad_items_bajos': bajos,
        'nivel_minimo_obtenido': nivel_minimo,
    }

def worst_case_simon_risk(risk_labels):
    """Agrega varias fichas de una IE con la regla de 'peor caso' entre docentes,
    igual que target_ie_simon.csv del subproyecto SIMON."""
    labels = [label for label in risk_labels if label]
    if not labels:
        return None
    return max(labels, key=lambda label: SIMON_RISK_ORDER.get(label, -1))

FIXED_SIMON_CODES = set(SIMON_CODE_TO_FIELD.keys())

def get_dynamic_responses_for_fichas(conn, ficha_ids):
    """Respuestas de preguntas dinamicas (fuera de las 8 fijas A-01..B-05) registradas
    para un conjunto de fichas, para que se muestren cuando los especialistas las llenan.

    Trae tambien el texto de la pregunta (item) y su seccion, con el mismo
    patron de JOIN que get_simon_question_labels: instrumento_codigo -> codigo
    de app_instrumento -> instrumento_id, porque el codigo de una pregunta
    solo es unico dentro de su instrumento (no globalmente)."""
    ficha_ids = [fid for fid in ficha_ids if fid]
    if not ficha_ids:
        return {}
    placeholders = ', '.join('?' for _ in ficha_ids)
    rows = conn.execute(f'''
        SELECT fri.ficha_id, fri.instrumento_codigo, fri.pregunta_codigo,
               fri.nivel, fri.respuesta_texto, fri.observacion,
               p.item, s.nombre AS seccion_nombre
        FROM ficha_respuesta_instrumento fri
        LEFT JOIN app_instrumento i ON i.codigo = fri.instrumento_codigo
        LEFT JOIN app_instrumento_pregunta p
               ON p.instrumento_id = i.instrumento_id AND p.codigo = fri.pregunta_codigo
        LEFT JOIN app_instrumento_seccion s ON s.seccion_id = p.seccion_id
        WHERE fri.ficha_id IN ({placeholders})
        ORDER BY fri.ficha_id, fri.pregunta_codigo
    ''', ficha_ids).fetchall()
    out = {}
    for row in rows:
        codigo = str(row['pregunta_codigo'] or '').upper()
        if codigo in FIXED_SIMON_CODES:
            continue
        out.setdefault(row['ficha_id'], []).append({
            'pregunta_codigo': row['pregunta_codigo'],
            'instrumento_codigo': row['instrumento_codigo'],
            'nivel': row['nivel'],
            'respuesta_texto': row['respuesta_texto'],
            'observacion': row['observacion'],
            'item': row['item'],
            'seccion_nombre': row['seccion_nombre'],
        })
    return out

def rebuild_simon_operational_summary(conn):
    conn.execute('DELETE FROM app_alerta_priorizada')
    conn.execute('DELETE FROM institucion_resumen')

    ficha_rows = conn.execute('''
        SELECT *
        FROM fichas_monitoreo
        WHERE COALESCE(codigo_modular, '') != ''
    ''').fetchall()
    infra_rows = conn.execute('SELECT * FROM infraestructura_censo_2025').fetchall()
    # Una visita marcada para revision queda fuera del calculo hasta que se
    # resuelva: un dato en duda no debe mover alertas ni priorizacion.
    campo_rows = conn.execute('''
        SELECT * FROM informe_campo_visita
        WHERE estado_fila = 1
          AND COALESCE(revision_estado, 'ok') NOT IN ('pendiente', 'en_revision')
    ''').fetchall()
    # Respuestas del instrumento DINAMICO (UGEL) en Nivel I/II -- alimentan su
    # propia fuente de alertas ('dinamica'), separada de SIMON, para que el
    # selector de fuente del KPI (SIMON / dinamicas / campo / todas) tenga
    # las 4 fuentes con datos reales en vez de solo 3.
    dinamica_rows = conn.execute('''
        SELECT fm.codigo_modular AS raw_code, fri.seccion_clave, fri.pregunta_codigo,
               fri.nivel, fri.observacion
        FROM ficha_respuesta_instrumento fri
        JOIN fichas_monitoreo fm ON fm.id = fri.ficha_id
        WHERE fri.instrumento_tipo = 'dinamica' AND fri.nivel IN (1, 2)
          AND COALESCE(fm.codigo_modular, '') != ''
    ''').fetchall()

    grouped = {}
    for row in ficha_rows:
        code = canonical_codigo_modular(conn, row['codigo_modular'])
        if not code:
            continue
        grouped.setdefault(code, []).append(dict(row))

    infra_by_code = {
        canonical_codigo_modular(conn, row['codigo_modular']): dict(row)
        for row in infra_rows
        if row['codigo_modular']
    }

    campo_by_code = {}
    for row in campo_rows:
        code = canonical_codigo_modular(conn, row['codigo_modular'])
        if not code:
            continue
        campo_by_code.setdefault(code, []).append(dict(row))

    dinamica_by_code = {}
    for row in dinamica_rows:
        code = canonical_codigo_modular(conn, row['raw_code'])
        if not code:
            continue
        bucket = dinamica_by_code.setdefault(code, {})
        seccion = row['seccion_clave'] or 'General'
        entry = bucket.setdefault(seccion, {'count': 0, 'sample': ''})
        entry['count'] += 1
        if not entry['sample'] and row['observacion']:
            entry['sample'] = row['observacion']

    # institucion_resumen.codigo_modular tiene FK a dim_institucion: un codigo
    # que no resuelve a una IE real (typo de OCR, codigo inventado) se excluye
    # aqui en vez de dejar que el INSERT falle -- mismo criterio que ya se usa
    # para informes de campo (resolve_codigo_modular_strict).
    valid_codes = {row['codigo_modular'] for row in conn.execute('SELECT codigo_modular FROM dim_institucion').fetchall()}
    all_codes = sorted((set(grouped) | set(infra_by_code) | set(campo_by_code)) & valid_codes)
    now = datetime.now().isoformat(timespec='seconds')

    for code in all_codes:
        fichas = grouped.get(code, [])
        infra = infra_by_code.get(code, {})
        campo_visitas = campo_by_code.get(code, [])
        all_values = []
        low_total = 0
        prep_low = 0
        teaching_low = 0
        ficha_risk_labels = []
        for ficha in fichas:
            values = simon_values_from_row(ficha)
            all_values.extend(values)
            low_total += sum(1 for value in values if value < 3)
            prep_low += sum(1 for field in ('a1', 'a2', 'a3') if safe_int(ficha.get(field), 0) in (1, 2))
            teaching_low += sum(1 for field in ('b1', 'b2', 'b3', 'b4', 'b5') if safe_int(ficha.get(field), 0) in (1, 2))
            ficha_risk = simon_documented_risk(values)
            if ficha_risk:
                ficha_risk_labels.append(ficha_risk['riesgo_intervencion'])

        # Riesgo de intervencion documentado (regla SIMON), agregado por "peor caso"
        # entre los docentes de la IE -- misma metodologia que target_ie_simon.csv.
        riesgo_intervencion_simon = worst_case_simon_risk(ficha_risk_labels)

        # Senales de informes de campo: conteo de variables concretas observadas
        # (cada una con evidencia textual real, nunca inventada) y una cita
        # representativa por variable para la descripcion de la alerta.
        campo_var_counts = {}
        campo_var_span = {}
        for visita in campo_visitas:
            for var in CAMPO_CONCRETE_VARIABLES:
                if safe_int(visita.get(var), 0) == 1:
                    campo_var_counts[var] = campo_var_counts.get(var, 0) + 1
                    if var not in campo_var_span and visita.get(f'span_{var}'):
                        campo_var_span[var] = visita.get(f'span_{var}')
        campo_alertas_total = sum(campo_var_counts.values())
        campo_infra_count = sum(campo_var_counts.get(v, 0) for v in CAMPO_INFRA_BUCKET)
        campo_pedagogic_count = sum(campo_var_counts.get(v, 0) for v in CAMPO_PEDAGOGIC_BUCKET)

        indicators = len(all_values)
        # El promedio se guarda redondeado a 2 decimales para mostrarlo, pero el
        # score del informe se calcula con el valor exacto: redondear antes
        # desplazaba el resultado (PAKUI daba 62.4 en vez de 62.5).
        promedio_exacto = (sum(all_values) / indicators) if indicators else 0.0
        average = round(promedio_exacto, 2) if indicators else 0.0
        low_share = low_total / indicators if indicators else 0.0
        simon_priority = round(min(22, 5 + (low_share * 11) + max(0, 3 - average) * 5), 2) if indicators else 0.0
        infra_alerts = safe_int(infra.get('alertas_infraestructura_campo'), 0)
        infra_score = safe_float(infra.get('risk_score_infra'), 0.0)
        risk_flags = first_text(infra, 'risk_flags')
        infra_priority = round(min(22, infra_score * 2.7 + min(infra_alerts, 50) / 10), 2) if infra else 0.0
        priority_score = round(max(simon_priority, infra_priority), 2)

        reasons = []
        if indicators and average < 2.5:
            priority_reason = 'Desempeno SIMON por debajo de 2.5; requiere refuerzo pedagogico'
            reasons.append(priority_reason)
        elif indicators and average < 3:
            priority_reason = 'Desempeno SIMON regular; seguimiento pedagogico recomendado'
            reasons.append(priority_reason)
        elif indicators:
            priority_reason = 'Desempeno SIMON estable con seguimiento ordinario'
            reasons.append(priority_reason)
        else:
            priority_reason = ''

        if riesgo_intervencion_simon:
            reasons.append(f'Riesgo de intervencion SIMON (regla documentada): {riesgo_intervencion_simon}')

        if infra:
            reasons.append(f'Riesgo infraestructura censo 2025: {round(infra_score, 2)}')
            if risk_flags:
                readable_flags = [
                    INFRA_FLAG_DESCRIPTIONS.get(flag, flag)
                    for flag in risk_flags.split('|')
                    if flag
                ]
                reasons.extend(readable_flags[:3])

        if campo_visitas:
            reasons.append(
                f'{len(campo_visitas)} visita(s) de campo con {campo_alertas_total} senal(es) '
                f'concretas con evidencia textual'
            )

        sources = []
        if fichas:
            sources.append('SIMON')
        if campo_visitas:
            sources.append('Campo')
        if infra:
            sources.append('Censo 2025')
        if not sources:
            coverage = 'Sin evidencia'
        elif len(sources) == 1:
            coverage = f'Solo {sources[0]}'
        else:
            coverage = ' + '.join(sources)

        monitors = sorted({first_text(f, 'monitor') for f in fichas if first_text(f, 'monitor')})
        docentes = sorted({first_text(f, 'docente') for f in fichas if first_text(f, 'docente')})
        fechas = sorted({first_text(f, 'fecha_ejecucion') for f in fichas if first_text(f, 'fecha_ejecucion')})
        campo_especialistas = sorted({
            first_text(v, 'especialista_detectado') for v in campo_visitas
            if first_text(v, 'especialista_detectado')
        })
        campo_temas = sorted(campo_var_counts.keys())

        # --- Score de priorizacion del informe v5 (seccion 3.3) ---
        # Se cuentan VARIABLES DISTINTAS observadas por severidad (no visitas),
        # tal como lo define el informe: una variable vista en tres visitas
        # sigue siendo una sola senal.
        sev_campo = {'alta': 0, 'media': 0, 'baja': 0}
        for var in campo_var_counts:
            severidad = (CAMPO_VARIABLE_META.get(var, {}) or {}).get('severidad', 'media')
            if severidad in sev_campo:
                sev_campo[severidad] += 1
        v5_simon = riesgo_simon_0_100(promedio_exacto if indicators else 0.0, low_share)
        v5_infra = riesgo_infra_0_100(infra_score)
        v5_campo = riesgo_campo_0_100(sev_campo['alta'], sev_campo['media'], sev_campo['baja'])
        v5_brecha = brecha_evidencia_0_100(bool(fichas), bool(infra), bool(campo_visitas))
        v5_score = calcular_score_priorizacion(v5_simon, v5_infra, v5_campo, v5_brecha)

        conn.execute('''
            INSERT OR REPLACE INTO institucion_resumen (
                codigo_modular, campo_documentos, simon_fichas, simon_indicadores,
                simon_promedio_nivel, total_alertas_campo, alertas_infraestructura_campo,
                alertas_pedagogicas_campo, coverage_category, campo_sin_simon,
                simon_sin_campo, ambas_fuentes, priority_score, priority_reason,
                campo_owners, campo_topics, simon_monitores, simon_docentes, simon_fechas,
                risk_score_infra, risk_flags, riesgo_intervencion_simon,
                simon_promedio_general_desempeno, simon_porcentaje_items_bajos,
                riesgo_simon_0_100, riesgo_infra_0_100, riesgo_campo_0_100,
                brecha_evidencia_0_100, score_priorizacion_ie, prioridad_v5,
                campo_alertas_altas, campo_alertas_medias, campo_alertas_bajas,
                campo_vars_distintas
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, 0, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            code,
            len(campo_visitas),
            len(fichas),
            indicators,
            average,
            low_total + infra_alerts + campo_alertas_total,
            infra_alerts + campo_infra_count,
            low_total + campo_pedagogic_count,
            coverage,
            1 if fichas and not infra else 0,
            1 if fichas and infra else 0,
            priority_score,
            '; '.join(reasons[:5]) or priority_reason,
            '|'.join(campo_especialistas),
            '|'.join(campo_temas),
            '|'.join(monitors),
            '|'.join(docentes[:12]),
            '|'.join(fechas),
            infra_score,
            risk_flags,
            riesgo_intervencion_simon,
            average if indicators else None,
            round(low_share, 2) if indicators else None,
            v5_simon, v5_infra, v5_campo, v5_brecha, v5_score, clasificar_prioridad(v5_score),
            sev_campo['alta'], sev_campo['media'], sev_campo['baja'], len(campo_var_counts),
        ))

        if low_total:
            # Severidad basada en la regla documentada de riesgo_intervencion cuando hay
            # fichas SIMON reales; si no hay suficiente evidencia por ficha, se usa el
            # umbral operativo previo como respaldo.
            if riesgo_intervencion_simon:
                severity = 'alta' if riesgo_intervencion_simon == 'Alto' else 'media'
            else:
                severity = 'alta' if average < 2.25 or low_share >= 0.65 else 'media'
            alert_id = 'simon_' + hashlib.md5(f'{code}:desempeno:{average}:{low_total}'.encode('utf-8')).hexdigest()[:12]
            simon_desc = (
                f'SIMON: {low_total} de {indicators} respuestas estan en Nivel I o II; '
                f'promedio general {average}.'
                + (f' Riesgo de intervencion (regla SIMON documentada): {riesgo_intervencion_simon}.' if riesgo_intervencion_simon else '')
                + ' Requiere refuerzo pedagogico focalizado.'
            )
            conn.execute('''
                INSERT OR REPLACE INTO app_alerta_priorizada (
                    alerta_id, codigo_modular, alerta_codigo, fuente, severidad,
                    score, titulo, descripcion, evidencia_count, estado, created_at
                )
                VALUES (?, ?, 'seguimiento_curricular', 'simon', ?, ?, ?, ?, ?, 'pendiente', ?)
            ''', (
                alert_id,
                code,
                severity,
                round(priority_score, 2),
                'Desempeno SIMON requiere refuerzo',
                simon_desc,
                len(fichas),
                now,
            ))

        if prep_low:
            alert_id = 'simon_' + hashlib.md5(f'{code}:preparacion:{prep_low}'.encode('utf-8')).hexdigest()[:12]
            conn.execute('''
                INSERT OR REPLACE INTO app_alerta_priorizada (
                    alerta_id, codigo_modular, alerta_codigo, fuente, severidad,
                    score, titulo, descripcion, evidencia_count, estado, created_at
                )
                VALUES (?, ?, 'seguimiento_curricular', 'simon', ?, ?, ?, ?, ?, 'pendiente', ?)
            ''', (
                alert_id,
                code,
                'alta' if prep_low >= len(fichas) * 2 else 'media',
                round(min(18, prep_low * 0.9), 2),
                'Preparacion para el aprendizaje en refuerzo',
                f'SIMON Preparacion: {prep_low} respuestas de los items A1-A3 estan en Nivel I o II.',
                prep_low,
                now,
            ))

        if teaching_low:
            alert_id = 'simon_' + hashlib.md5(f'{code}:ensenanza:{teaching_low}'.encode('utf-8')).hexdigest()[:12]
            conn.execute('''
                INSERT OR REPLACE INTO app_alerta_priorizada (
                    alerta_id, codigo_modular, alerta_codigo, fuente, severidad,
                    score, titulo, descripcion, evidencia_count, estado, created_at
                )
                VALUES (?, ?, 'capacitacion_docente', 'simon', ?, ?, ?, ?, ?, 'pendiente', ?)
            ''', (
                alert_id,
                code,
                'alta' if teaching_low >= len(fichas) * 3 else 'media',
                round(min(18, teaching_low * 0.7), 2),
                'Ensenanza para el aprendizaje en refuerzo',
                f'SIMON Ensenanza: {teaching_low} respuestas de los items B1-B5 estan en Nivel I o II.',
                teaching_low,
                now,
            ))

        if infra and infra_score > 0:
            severity = 'alta' if infra_score >= 5 or 'edificacion_con_riesgo_estructural' in risk_flags else 'media'
            alert_id = 'infra_' + hashlib.md5(f'{code}:infra:{infra_score}:{risk_flags}'.encode('utf-8')).hexdigest()[:12]
            flag_labels = [
                INFRA_FLAG_DESCRIPTIONS.get(flag, flag)
                for flag in risk_flags.split('|')
                if flag
            ]
            infra_desc = (
                f'Censo Educativo 2025: risk_score_infra={round(infra_score, 2)}. '
                + (f'Motivos: {"; ".join(flag_labels[:4])}.' if flag_labels else 'Sin banderas especificas registradas.')
            )
            conn.execute('''
                INSERT OR REPLACE INTO app_alerta_priorizada (
                    alerta_id, codigo_modular, alerta_codigo, fuente, severidad,
                    score, titulo, descripcion, evidencia_count, estado, created_at
                )
                VALUES (?, ?, 'agua_saneamiento', 'censo_2025', ?, ?, ?, ?, ?, 'pendiente', ?)
            ''', (
                alert_id,
                code,
                severity,
                round(infra_score, 2),
                'Riesgo de infraestructura censal',
                infra_desc,
                max(1, len([flag for flag in risk_flags.split('|') if flag])),
                now,
            ))

        # Una alerta por cada variable concreta de campo observada, con la
        # cita textual real como evidencia -- nunca se inventa severidad a
        # partir del texto, solo se aplica la prioridad editorial fija por
        # variable (CAMPO_VARIABLE_META).
        for var, count in campo_var_counts.items():
            meta = CAMPO_VARIABLE_META.get(var, {})
            alert_id = 'campo_' + hashlib.md5(f'{code}:{var}:{count}'.encode('utf-8')).hexdigest()[:12]
            span = campo_var_span.get(var, '')
            descripcion = (
                f"{meta.get('label', var)}: observado en {count} visita(s) de campo. "
                f"{meta.get('importancia', '')} "
                + (f'Evidencia: "{span[:280]}"' if span else '')
            )
            conn.execute('''
                INSERT OR REPLACE INTO app_alerta_priorizada (
                    alerta_id, codigo_modular, alerta_codigo, fuente, severidad,
                    score, titulo, descripcion, evidencia_count, estado, created_at
                )
                VALUES (?, ?, 'informe_campo', 'campo', ?, ?, ?, ?, ?, 'pendiente', ?)
            ''', (
                alert_id,
                code,
                meta.get('severidad', 'media'),
                round(min(20, count * 5), 2),
                meta.get('label', var),
                descripcion,
                count,
                now,
            ))

        # Alertas del instrumento dinamico UGEL (opcional): una por seccion con
        # respuestas en Nivel I/II, con una observacion de ejemplo si existe.
        # Nunca amplia all_codes -- una IE solo tiene respuestas dinamicas si
        # ya tiene una ficha SIMON (mismo formulario), asi que ya esta incluida.
        for seccion, info in dinamica_by_code.get(code, {}).items():
            dyn_count = info['count']
            if dyn_count <= 0:
                continue
            dyn_severity = 'alta' if dyn_count >= 3 else 'media'
            dyn_alert_id = 'dinamica_' + hashlib.md5(f'{code}:{seccion}:{dyn_count}'.encode('utf-8')).hexdigest()[:12]
            dyn_desc = (
                f'Instrumento dinamico UGEL: {dyn_count} pregunta(s) de la seccion "{seccion}" en Nivel I o II.'
                + (f' Ejemplo de observacion: "{info["sample"][:200]}"' if info['sample'] else '')
            )
            conn.execute('''
                INSERT OR REPLACE INTO app_alerta_priorizada (
                    alerta_id, codigo_modular, alerta_codigo, fuente, severidad,
                    score, titulo, descripcion, evidencia_count, estado, created_at
                )
                VALUES (?, ?, 'preguntas_dinamicas', 'dinamica', ?, ?, ?, ?, ?, 'pendiente', ?)
            ''', (
                dyn_alert_id, code, dyn_severity, round(min(15, dyn_count * 4), 2),
                f'Instrumento dinámico: {seccion}', dyn_desc, dyn_count, now,
            ))

def unavailable_sheet(kind, message):
    return {
        'source': 'sin_datos',
        'resumen': {
            'estado': {
                'excel_value': 'sin dato',
                'label': 'Sin datos',
                'color': '-',
                'status_key': 'info',
                'action': message,
            },
            'mensaje': message,
        },
        'items': [],
        'radar': [],
    }

def build_real_teacher_sheet(fichas, conn=None):
    dynamic_by_ficha = get_dynamic_responses_for_fichas(conn, [f.get('id') for f in fichas]) if conn else {}
    docentes = []
    for ficha in fichas:
        values = simon_values_from_row(ficha)
        average = round(sum(values) / len(values), 2) if values else 0.0
        semaforo = simon_level_state(average)
        documented_risk = simon_documented_risk(values)
        docentes.append({
            'nombre': first_text(ficha, 'docente') or 'Docente sin nombre',
            'grado': first_text(ficha, 'grado') or first_text(ficha, 'seccion') or 'Sin grado',
            'area': first_text(ficha, 'area') or 'SIMON',
            'fecha': first_text(ficha, 'fecha_ejecucion'),
            'nivel_promedio': average,
            'visita1': {
                'label': 'Visita 1 - Diagnostico',
                'score': simon_level_percent(average),
                'nivel_promedio': average,
                'estado_excel': semaforo['excel_value'],
                'semaforo': semaforo,
            },
            'visita2': None,
            'delta': None,
            'necesita_refuerzo': average < 3,
            # Riesgo de intervencion segun la regla documentada del subproyecto SIMON
            # (docs/02_criterios_riesgo_intervencion.md), calculado en vivo desde los
            # niveles reales de la ficha -- no es un valor adivinado.
            'riesgo_intervencion_simon': documented_risk['riesgo_intervencion'] if documented_risk else None,
            'riesgo_intervencion_detalle': documented_risk,
            # Respuestas a preguntas dinamicas (fuera del instrumento fijo A-01..B-05)
            # que los especialistas hayan llenado para esta ficha.
            'respuestas_dinamicas': dynamic_by_ficha.get(ficha.get('id'), []),
        })

    needs = sum(1 for docente in docentes if docente['necesita_refuerzo'])
    stable = len(docentes) - needs
    critical_initial = sum(1 for docente in docentes if docente['nivel_promedio'] < 2.5)
    return {
        'source': 'simon_real',
        'resumen': {
            'total': len(docentes),
            'necesitan_refuerzo': needs,
            'estables': stable,
            'critico_a_estable': 0,
            'criticos_iniciales': critical_initial,
            'ieap': None,
            'promedio_nivel': round(sum(d['nivel_promedio'] for d in docentes) / len(docentes), 2) if docentes else 0.0,
        },
        'items': docentes,
    }

def infra_status_from_score(score):
    score = safe_float(score, 0.0)
    if score >= 5:
        return dict(SEMAFORO_PEDAGOGICO[0])
    if score >= 3:
        return dict(SEMAFORO_PEDAGOGICO[1])
    return dict(SEMAFORO_PEDAGOGICO[2])

def build_infra_ranking(conn):
    """Ranking de priorizacion por IE del Censo Educativo 2025: a diferencia de
    SIMON, el censo es una sola foto (no hay 'visitas' que deduplicar), asi que
    se ordena directamente por risk_score_infra descendente -- mayor riesgo,
    mayor prioridad de atencion. Tambien agrega la frecuencia real de cada
    bandera de riesgo (risk_flags) para ver que problema es mas comun en las
    382 IE, no solo cuantas estan en rojo."""
    rows = conn.execute('''
        SELECT codigo_modular, nombre_iiee, distrito, risk_score_infra, risk_flags,
               edificaciones_riesgo, aulas, aulas_en_uso, alumnos_por_aula_en_uso
        FROM infraestructura_censo_2025
        ORDER BY risk_score_infra DESC
    ''').fetchall()

    flags_count = {}
    ranking = []
    for row in rows:
        item = dict(row)
        flags = [f for f in str(item.get('risk_flags') or '').split('|') if f]
        for flag in flags:
            flags_count[flag] = flags_count.get(flag, 0) + 1
        estado = infra_status_from_score(item.get('risk_score_infra'))
        ranking.append({
            'codigo_modular': item['codigo_modular'],
            'nombre_iiee': item.get('nombre_iiee') or item['codigo_modular'],
            'distrito': item.get('distrito'),
            'risk_score_infra': round(safe_float(item.get('risk_score_infra'), 0.0), 2),
            'nivel': estado['label'],
            'banderas': len(flags),
            'edificaciones_riesgo': safe_int(item.get('edificaciones_riesgo'), 0),
            'aulas_en_uso': safe_int(item.get('aulas_en_uso'), 0),
            'aulas': safe_int(item.get('aulas'), 0),
        })

    total = len(rows)
    flags_distribution = sorted(
        [
            {
                'flag': flag,
                'descripcion': INFRA_FLAG_DESCRIPTIONS.get(flag, flag),
                'ies': count,
                'porcentaje': round(count / total * 100, 1) if total else 0,
            }
            for flag, count in flags_count.items()
        ],
        key=lambda item: item['ies'],
        reverse=True,
    )

    return {
        'ies': ranking,
        'total_ie': total,
        'criticas': sum(1 for r in ranking if r['nivel'] == 'Critico'),
        'en_alerta': sum(1 for r in ranking if r['nivel'] == 'Alerta'),
        'flags_distribution': flags_distribution,
    }

def build_real_infra_sheet(infra):
    if not infra:
        return unavailable_sheet('infraestructura', 'No hay registro de infraestructura censal 2025 para esta IE.')

    score = safe_float(infra.get('risk_score_infra'), 0.0)
    flags = [flag for flag in first_text(infra, 'risk_flags').split('|') if flag]
    estado = infra_status_from_score(score)
    flags_detail = [
        {
            'flag': flag,
            'descripcion': INFRA_FLAG_DESCRIPTIONS.get(flag, flag),
        }
        for flag in flags
    ]

    critical_flags = {
        'edificacion_con_riesgo_estructural',
        'sin_codigo_local_para_cruce',
        'sin_registros_edificaciones',
        'sin_registros_aulas',
        'aulas_registradas_no_en_uso',
    }
    warning_flags = {
        'alta_densidad_alumnos_por_aula_en_uso',
        'conservacion_mala_en_puertas_o_ventanas',
        'aulas_en_uso_sin_senalizacion_completa',
    }
    brechas_criticas = sum(1 for flag in flags if flag in critical_flags)
    brechas_alerta = sum(1 for flag in flags if flag in warning_flags)

    items = [
        {
            'categoria': 'Edificaciones',
            'valor': safe_int(infra.get('edificaciones'), 0),
            'en_uso': safe_int(infra.get('edificaciones_en_uso'), 0),
            'riesgo': safe_int(infra.get('edificaciones_riesgo'), 0),
            'semaforo': dict(SEMAFORO_PEDAGOGICO[0]) if safe_int(infra.get('edificaciones_riesgo'), 0) else dict(SEMAFORO_PEDAGOGICO[2]),
            'observacion': 'Edificaciones con riesgo estructural' if safe_int(infra.get('edificaciones_riesgo'), 0) else 'Sin riesgo estructural registrado',
        },
        {
            'categoria': 'Aulas',
            'valor': safe_int(infra.get('aulas'), 0),
            'en_uso': safe_int(infra.get('aulas_en_uso'), 0),
            'riesgo': 1 if 'aulas_registradas_no_en_uso' in flags else 0,
            'semaforo': dict(SEMAFORO_PEDAGOGICO[0]) if 'aulas_registradas_no_en_uso' in flags else dict(SEMAFORO_PEDAGOGICO[2]),
            'observacion': 'Aulas registradas no aparecen en uso' if 'aulas_registradas_no_en_uso' in flags else 'Aulas en uso registradas',
        },
        {
            'categoria': 'Densidad',
            'valor': safe_float(infra.get('alumnos_por_aula_en_uso'), 0.0),
            'en_uso': safe_int(infra.get('alumnos_censo'), 0),
            'riesgo': 1 if 'alta_densidad_alumnos_por_aula_en_uso' in flags else 0,
            'semaforo': dict(SEMAFORO_PEDAGOGICO[1]) if 'alta_densidad_alumnos_por_aula_en_uso' in flags else dict(SEMAFORO_PEDAGOGICO[2]),
            'observacion': 'Mas de 35 estudiantes por aula en uso' if 'alta_densidad_alumnos_por_aula_en_uso' in flags else 'Densidad dentro de rango',
        },
        {
            'categoria': 'Alertas de campo',
            'valor': safe_int(infra.get('alertas_infraestructura_campo'), 0),
            'en_uso': None,
            'riesgo': safe_int(infra.get('alertas_infraestructura_campo'), 0),
            'semaforo': infra_status_from_score(min(safe_int(infra.get('alertas_infraestructura_campo'), 0), 50) / 25 * 3),
            'observacion': 'Alertas de infraestructura o servicios asociadas a la IE',
        },
    ]

    return {
        'source': 'censo_2025',
        'resumen': {
            'estado': estado,
            'risk_score_infra': round(score, 2),
            'risk_flags': flags,
            'flags_detalle': flags_detail,
            'alertas_infraestructura_campo': safe_int(infra.get('alertas_infraestructura_campo'), 0),
            'edificaciones': safe_int(infra.get('edificaciones'), 0),
            'edificaciones_en_uso': safe_int(infra.get('edificaciones_en_uso'), 0),
            'edificaciones_riesgo': safe_int(infra.get('edificaciones_riesgo'), 0),
            'aulas': safe_int(infra.get('aulas'), 0),
            'aulas_en_uso': safe_int(infra.get('aulas_en_uso'), 0),
            'alumnos_por_aula_en_uso': safe_float(infra.get('alumnos_por_aula_en_uso'), 0.0),
            'brechas_criticas': brechas_criticas,
            'brechas_en_alerta': brechas_alerta,
            'componentes_estables': max(0, 4 - brechas_criticas - brechas_alerta),
        },
        'items': items,
    }

def build_real_simon_dashboard(conn, scored):
    # Solo la ultima visita por docente -- ver simon_latest_ficha_por_docente:
    # si un docente tiene mas de una ficha, la hoja "Docentes" debe mostrar su
    # estado actual, no una evaluacion vieja ya superada (ni una carga duplicada).
    ficha_rows, _dedup_meta = simon_latest_ficha_por_docente(conn)
    ficha_rows = sorted(ficha_rows, key=lambda item: (item.get('codigo_modular') or '', item.get('fecha_ejecucion') or '', item.get('docente') or ''))
    infra_rows = conn.execute('SELECT * FROM infraestructura_censo_2025').fetchall()
    if not ficha_rows and not infra_rows:
        return {
            'source': 'empty',
            'sheets': ['Docentes', 'Infraestructura', 'Alumnos', 'Progreso comparativo'],
            'selected_codigo': None,
            'semaforo': SEMAFORO_PEDAGOGICO,
            'instituciones': [],
            'definition': {
                'ieap': 'Pendiente: requiere Visita 2 real para comparar Critico a Estable.',
                'rtbg': 'Pendiente: requiere datos reales de brechas de gestion.',
            },
        }

    grouped = {}
    for row in ficha_rows:
        item = dict(row)
        code = canonical_codigo_modular(conn, item.get('codigo_modular'))
        item['codigo_modular'] = code
        grouped.setdefault(code, []).append(item)
    infra_by_code = {
        canonical_codigo_modular(conn, row['codigo_modular']): dict(row)
        for row in infra_rows
        if row['codigo_modular']
    }

    risk_lookup = {str(row.get('codigo_modular')): row for row in scored}
    institutions = []
    all_codes = sorted(set(grouped) | set(infra_by_code))
    for code in all_codes:
        fichas = grouped.get(code, [])
        infra = infra_by_code.get(code, {})
        official = conn.execute('''
            SELECT codigo_modular, nombre_iiee, nivel_modalidad, distrito
            FROM dim_institucion
            WHERE codigo_modular = ?
        ''', (code,)).fetchone()
        base = dict(official) if official else {
            'codigo_modular': code,
            'nombre_iiee': first_text(infra, 'nombre_iiee') or (first_text(fichas[0], 'nombre_ie') if fichas else code),
            'nivel_modalidad': first_text(infra, 'nivel_modalidad') or (first_text(fichas[0], 'nivel_modalidad') if fichas else ''),
            'distrito': first_text(infra, 'distrito'),
        }
        risk = dict(risk_lookup.get(code, {}))
        if not risk:
            summary = conn.execute('SELECT * FROM institucion_resumen WHERE codigo_modular = ?', (code,)).fetchone()
            risk = dict(summary) if summary else {}
            risk.update(calculate_risk(risk))

        docentes = build_real_teacher_sheet(fichas, conn) if fichas else {
            'source': 'sin_datos',
            'resumen': {
                'total': 0,
                'necesitan_refuerzo': 0,
                'estables': 0,
                'critico_a_estable': 0,
                'criticos_iniciales': 0,
                'ieap': None,
                'promedio_nivel': 0.0,
                'mensaje': 'No hay ficha SIMON cargada para esta IE.',
            },
            'items': [],
        }
        promedio_ie = docentes['resumen']['promedio_nivel']
        progreso = {
            'source': 'sin_visita_2',
            'resumen': {
                'promedio_visita1': simon_level_percent(promedio_ie),
                'promedio_visita2': None,
                'mejora': None,
                'dias_brecha_visita1': None,
                'dias_brecha_visita2': None,
                'rtbg': None,
                'mensaje': 'El CSV cargado corresponde a SIMON/Visita 1. Falta Visita 2 real para calcular evolucion.',
            },
            'radar': [],
        }
        base.update(risk)
        base.update({
            'simon_fichas': len(fichas),
            'simon_promedio_nivel': promedio_ie,
            'campo_documentos': safe_int(risk.get('campo_documentos'), 0),
            'total_alertas_campo': safe_int(risk.get('total_alertas_campo'), 0),
            'docentes': docentes,
            'infraestructura': build_real_infra_sheet(infra),
            'alumnos': unavailable_sheet('alumnos', 'El CSV cargado no contiene indicadores reales de alumnos.'),
            'progreso': progreso,
            'kpis': {
                'ieap': None,
                'rtbg': None,
                'docentes_refuerzo': docentes['resumen']['necesitan_refuerzo'],
                'simon_promedio': promedio_ie,
                'risk_score_infra': safe_float(infra.get('risk_score_infra'), 0.0) if infra else 0.0,
            },
        })
        institutions.append(base)

    institutions.sort(
        key=lambda item: (
            safe_float(item.get('risk_score'), 0),
            safe_float(((item.get('infraestructura') or {}).get('resumen') or {}).get('risk_score_infra'), 0),
            item['docentes']['resumen']['necesitan_refuerzo'],
        ),
        reverse=True,
    )
    return {
        'source': 'real_simon_infra',
        'sheets': ['Docentes', 'Infraestructura', 'Alumnos', 'Progreso comparativo'],
        'selected_codigo': institutions[0]['codigo_modular'] if institutions else None,
        'semaforo': SEMAFORO_PEDAGOGICO,
        'instituciones': institutions,
        'definition': {
            'ieap': 'Pendiente: el CSV contiene Visita 1 SIMON; se calculara cuando exista Visita 2 real.',
            'rtbg': 'Pendiente: requiere datos reales de cierre de brechas de gestion.',
            'infraestructura': 'Censo Educativo 2025 cruzado con edificaciones, aulas y alertas de infraestructura.',
        },
    }

def build_simon_indicator_results(conn, rows=None):
    labels = get_simon_question_labels(conn)
    if rows is None:
        rows, _meta = simon_latest_ficha_por_docente(conn)
    if not rows:
        return []
    indicators = []
    for code, field in SIMON_CODE_TO_FIELD.items():
        values = [safe_int(row[field], 0) for row in rows if safe_int(row[field], 0) > 0]
        if not values:
            continue
        good = sum(1 for value in values if value >= 3)
        low = len(values) - good
        average = round(sum(values) / len(values), 2)
        stable_pct = round(good / len(values) * 100)
        state = simon_level_state(average)
        label = labels.get(code, {}).get('item') or code
        indicators.append({
            'area': code,
            'valor': stable_pct,
            'estado': state['label'],
            'lectura': f'{low} de {len(values)} fichas por debajo de Nivel III. Promedio {average}.',
            'pregunta': label,
            'promedio': average,
            'bajo_nivel': low,
            'total': len(values),
        })
    return indicators

def build_simon_ranking(conn, fichas=None, meta=None):
    """Ranking de priorizacion por IE para SIMON: agrupa la ultima visita de
    cada docente (ver simon_latest_ficha_por_docente) por codigo_modular y
    calcula el riesgo de peor caso entre sus docentes -- el mismo criterio
    de agregacion que institucion_resumen, pero calculado en vivo desde
    fichas_monitoreo en vez de depender de un rebuild ya desincronizado.
    Ordenado de mayor a menor prioridad: riesgo Alto primero, y dentro de un
    mismo nivel de riesgo, el promedio general mas bajo primero."""
    if fichas is None or meta is None:
        fichas, meta = simon_latest_ficha_por_docente(conn)
    por_ie = {}
    for ficha in fichas:
        codigo = ficha.get('codigo_modular')
        if not codigo:
            continue
        por_ie.setdefault(codigo, []).append(ficha)

    ranking = []
    for codigo, docentes in por_ie.items():
        official = conn.execute('''
            SELECT nombre_iiee, distrito, nivel_modalidad
            FROM dim_institucion WHERE codigo_modular = ?
        ''', (codigo,)).fetchone()
        promedios = []
        risk_labels = []
        ultima_visita = ''
        for docente in docentes:
            values = simon_values_from_row(docente)
            if values:
                promedios.append(sum(values) / len(values))
            riesgo = simon_documented_risk(values)
            if riesgo:
                risk_labels.append(riesgo['riesgo_intervencion'])
            fecha = str(docente.get('fecha_ejecucion') or '')
            if fecha > ultima_visita:
                ultima_visita = fecha
        riesgo_ie = worst_case_simon_risk(risk_labels)
        ranking.append({
            'codigo_modular': codigo,
            'nombre_iiee': (official['nombre_iiee'] if official else None) or (docentes[0].get('nombre_ie') if docentes else None) or codigo,
            'distrito': official['distrito'] if official else None,
            'docentes_evaluados': len(docentes),
            'promedio_general': round(sum(promedios) / len(promedios), 2) if promedios else 0,
            'riesgo_intervencion': riesgo_ie,
            'ultima_visita': ultima_visita or None,
        })

    ranking.sort(key=lambda item: (-SIMON_RISK_ORDER.get(item['riesgo_intervencion'], -1), item['promedio_general']))
    return {
        'ies': ranking,
        'total_ie': len(ranking),
        'fichas_totales': meta['total_bruto'],
        'fichas_ultima_visita': meta['total_unico'],
        'duplicados_excluidos': meta['duplicados_excluidos'],
    }

def campo_severidad_nivel(variables):
    """Alto/Medio/Bajo a partir de la severidad editorial documentada en
    CAMPO_VARIABLE_META (ver esa constante): Alto si alguna variable marcada
    es de severidad 'alta', Medio si la peor es 'media', Bajo si solo hay
    'baja'. Regla explicita, no inferida del texto."""
    severidades = {CAMPO_VARIABLE_META.get(v, {}).get('severidad', 'media') for v in variables}
    if 'alta' in severidades:
        return 'Alto'
    if 'media' in severidades:
        return 'Medio'
    return 'Bajo' if severidades else None

def build_campo_ranking(conn):
    """Ranking de priorizacion por IE de informes de campo.

    A diferencia de SIMON, la fecha de visita (fecha_visita) solo esta
    disponible en ~1 de cada 3 filas de esta fuente (el resto son informes
    donde la fecha no se pudo extraer con confianza) -- por eso, a diferencia
    del ranking SIMON, aqui NO se aplica "solo la ultima visita": se agregan
    TODOS los hallazgos registrados por IE (de lo contrario se descartaria la
    mayoria de la evidencia real), y se muestra la fecha mas reciente conocida
    solo como dato informativo cuando existe."""
    rows = conn.execute('''
        SELECT v.codigo_modular, v.fecha_visita, v.especialista_detectado, i.distrito,
        ''' + ', '.join(f'v.{var}' for var in CAMPO_CONCRETE_VARIABLES) + '''
        FROM informe_campo_visita v
        LEFT JOIN dim_institucion i ON i.codigo_modular = v.codigo_modular
        WHERE v.estado_fila = 1 AND COALESCE(v.codigo_modular, '') != ''
    ''').fetchall()

    por_ie = {}
    variable_counts = {var: 0 for var in CAMPO_CONCRETE_VARIABLES}
    distrito_variable = {}
    for row in rows:
        item = dict(row)
        codigo = item['codigo_modular']
        distrito = item.get('distrito') or 'Sin distrito'
        entry = por_ie.setdefault(codigo, {'visitas': 0, 'variables': set(), 'ultima_visita': '', 'especialistas': set()})
        entry['visitas'] += 1
        if item.get('especialista_detectado'):
            entry['especialistas'].add(item['especialista_detectado'])
        fecha = str(item.get('fecha_visita') or '')
        if fecha and fecha > entry['ultima_visita']:
            entry['ultima_visita'] = fecha
        for var in CAMPO_CONCRETE_VARIABLES:
            if item.get(var):
                entry['variables'].add(var)
                variable_counts[var] += 1
                distrito_variable.setdefault(distrito, {})
                distrito_variable[distrito][var] = distrito_variable[distrito].get(var, 0) + 1

    official_names = {}
    if por_ie:
        placeholders = ', '.join('?' for _ in por_ie)
        for r in conn.execute(
            f'SELECT codigo_modular, nombre_iiee, distrito FROM dim_institucion WHERE codigo_modular IN ({placeholders})',
            list(por_ie.keys()),
        ).fetchall():
            official_names[r['codigo_modular']] = {'nombre_iiee': r['nombre_iiee'], 'distrito': r['distrito']}

    ranking = []
    for codigo, entry in por_ie.items():
        nivel = campo_severidad_nivel(entry['variables'])
        info = official_names.get(codigo, {})
        ranking.append({
            'codigo_modular': codigo,
            'nombre_iiee': info.get('nombre_iiee') or codigo,
            'distrito': info.get('distrito'),
            'visitas': entry['visitas'],
            'banderas': len(entry['variables']),
            'nivel': nivel,
            'ultima_visita': entry['ultima_visita'] or None,
            'especialistas': len(entry['especialistas']),
        })
    ranking.sort(key=lambda item: (-SIMON_RISK_ORDER.get(item['nivel'], -1), -item['banderas']))

    total_visitas = len(rows)
    con_fecha = sum(1 for r in rows if str(r['fecha_visita'] or ''))
    variables_distribution = sorted(
        [
            {
                'variable': var,
                'label': CAMPO_VARIABLE_META.get(var, {}).get('label', var),
                'severidad': CAMPO_VARIABLE_META.get(var, {}).get('severidad', 'media'),
                'importancia': CAMPO_VARIABLE_META.get(var, {}).get('importancia', ''),
                'count': count,
                'porcentaje': round(count / total_visitas * 100, 1) if total_visitas else 0,
            }
            for var, count in variable_counts.items() if count > 0
        ],
        key=lambda item: item['count'],
        reverse=True,
    )
    infra_total = sum(c for v, c in variable_counts.items() if v in CAMPO_INFRA_BUCKET)
    pedagogico_total = sum(c for v, c in variable_counts.items() if v in CAMPO_PEDAGOGIC_BUCKET)

    # Mapa de calor distrito x variable: solo las variables con al menos una
    # deteccion (mismo orden que variables_distribution) y los distritos
    # reales presentes en los datos -- sin inventar celdas en cero decorativas.
    top_variables = [item['variable'] for item in variables_distribution]
    heatmap = {
        'distritos': sorted(distrito_variable.keys()),
        'variables': [{'variable': v, 'label': CAMPO_VARIABLE_META.get(v, {}).get('label', v)} for v in top_variables],
        'celdas': [
            {'distrito': distrito, 'variable': var, 'count': distrito_variable.get(distrito, {}).get(var, 0)}
            for distrito in sorted(distrito_variable.keys())
            for var in top_variables
        ],
    }

    return {
        'ies': ranking,
        'total_ie': len(ranking),
        'total_visitas': total_visitas,
        'visitas_con_fecha': con_fecha,
        'total_documentos': conn.execute('SELECT COUNT(*) FROM informe_campo_documento').fetchone()[0],
        'riesgo_alto': sum(1 for r in ranking if r['nivel'] == 'Alto'),
        'variables_distribution': variables_distribution,
        'buckets': [
            {'bucket': 'Infraestructura y seguridad', 'count': infra_total},
            {'bucket': 'Pedagogico e institucional', 'count': pedagogico_total},
        ],
        'heatmap': heatmap,
    }

def _condicion_periodo(campo_fecha, anio='', mes=''):
    """Condicion SQL para filtrar por anio, por mes, o por ambos.

    Antes el mes solo se aplicaba junto con un anio: elegir "Julio" sin anio no
    filtraba nada y la vista se veia igual, sin avisar. Ahora cada combinacion
    tiene su comparacion:
      anio + mes -> 'YYYY-MM'   (primeros 7 caracteres)
      solo anio  -> 'YYYY'      (primeros 4)
      solo mes   -> 'MM'        (caracteres 6-7, el mes de cualquier anio)
    Sirve tanto para fechas 'YYYY-MM-DD' como para timestamps completos.
    """
    anio = str(anio or '').strip()
    mes = str(mes or '').strip()
    base = f"COALESCE({campo_fecha}, '')"
    if anio and mes:
        return f"SUBSTR({base}, 1, 7) = ?", [f'{anio}-{mes.zfill(2)}']
    if anio:
        return f"SUBSTR({base}, 1, 4) = ?", [anio]
    if mes:
        return f"SUBSTR({base}, 6, 2) = ?", [mes.zfill(2)]
    return '', []


def build_asistencias_dashboard(conn, anio='', mes='', codigo_ie=''):
    """Seguimiento de la capacitacion docente: cuantos talleres se dictaron,
    cuanta gente asistio y a quienes todavia no se ha capacitado.

    Todo se cuenta contra talleres activos (t.estado_fila = 1): un taller dado
    de baja no debe seguir sumando en los indicadores.
    """
    # Filtros opcionales de periodo (anio o anio+mes) y de institucion. La
    # fecha del taller viene del OCR y puede faltar, por eso se cae a creado_en.
    condiciones = ['COALESCE(a.estado_fila, 1) = 1', 'COALESCE(t.estado_fila, 1) = 1']
    filtros = []
    cond_periodo, params_periodo = _condicion_periodo("NULLIF(t.fecha_inicio, '')", anio, mes)
    if cond_periodo:
        # Si el taller no trae fecha propia (el OCR no siempre la lee) se usa la
        # de registro, para no perder la fila del grafico.
        cond_alt, params_alt = _condicion_periodo('t.creado_en', anio, mes)
        condiciones.append(f"(({cond_periodo}) OR (COALESCE(t.fecha_inicio, '') = '' AND ({cond_alt})))")
        filtros += params_periodo + params_alt
    if codigo_ie:
        condiciones.append('t.codigo_modular_sede = ?')
        filtros.append(codigo_ie)
    base = '''
        FROM taller_asistencia a
        JOIN taller_capacitacion t ON t.id = a.taller_id
        WHERE ''' + ' AND '.join(condiciones) + '''
    '''

    totales = conn.execute(f'''
        SELECT COUNT(*) AS asistencias,
               COUNT(DISTINCT a.profesor_id) AS profesores_distintos,
               COUNT(DISTINCT a.taller_id) AS talleres
        {base}
    ''', filtros).fetchone()

    total_padron = conn.execute(
        'SELECT COUNT(*) FROM profesor WHERE COALESCE(estado_fila, 1) = 1'
    ).fetchone()[0]

    # Por mes: la fecha del taller viene del OCR y puede faltar; en ese caso se
    # cae a la fecha de registro para no perder la fila del grafico.
    por_mes = conn.execute(f'''
        SELECT SUBSTR(COALESCE(NULLIF(t.fecha_inicio, ''), t.creado_en), 1, 7) AS mes,
               COUNT(*) AS asistencias,
               COUNT(DISTINCT a.taller_id) AS talleres
        {base}
        GROUP BY mes
        ORDER BY mes
    ''', filtros).fetchall()

    por_cargo = conn.execute(f'''
        SELECT COALESCE(NULLIF(p.cargo, ''), 'sin dato') AS cargo,
               COUNT(*) AS asistencias,
               COUNT(DISTINCT a.profesor_id) AS profesores
        {base} AND p.id IS NOT NULL
        GROUP BY cargo
        ORDER BY asistencias DESC
    '''.replace('FROM taller_asistencia a', 'FROM taller_asistencia a LEFT JOIN profesor p ON p.id = a.profesor_id'),
        filtros
    ).fetchall()

    top_ies = conn.execute(f'''
        SELECT COALESCE(NULLIF(i.nombre_iiee, ''), 'Sin IE asignada') AS ie,
               i.nivel_modalidad AS nivel,
               COUNT(*) AS asistencias,
               COUNT(DISTINCT a.profesor_id) AS profesores
        {base}
        GROUP BY i.nombre_iiee, i.nivel_modalidad
        ORDER BY asistencias DESC
        LIMIT 10
    '''.replace('FROM taller_asistencia a',
                'FROM taller_asistencia a LEFT JOIN profesor p ON p.id = a.profesor_id '
                'LEFT JOIN dim_institucion i ON i.codigo_modular = COALESCE(a.codigo_modular_ie_detectado, p.codigo_modular_ie)'),
        filtros
    ).fetchall()

    recientes = conn.execute('''
        SELECT t.id, t.nombre_evento, t.compromiso_indicador, t.nivel_educativo,
               t.fecha_inicio, t.fecha_fin, t.creado_en,
               i.nombre_iiee AS sede_nombre, u.nombre AS creado_por_nombre,
               (SELECT COUNT(*) FROM taller_asistencia a
                 WHERE a.taller_id = t.id AND COALESCE(a.estado_fila, 1) = 1) AS asistentes
        FROM taller_capacitacion t
        LEFT JOIN dim_institucion i ON i.codigo_modular = t.codigo_modular_sede
        LEFT JOIN app_user u ON u.user_id = t.creado_por
        WHERE ''' + ' AND '.join(c for c in condiciones if not c.startswith('COALESCE(a.')) + '''
        ORDER BY t.creado_en DESC
        LIMIT 8
    ''', filtros).fetchall()

    capacitados = safe_int(totales['profesores_distintos'] if totales else 0, 0)
    return {
        'total_talleres': safe_int(totales['talleres'] if totales else 0, 0),
        'total_asistencias': safe_int(totales['asistencias'] if totales else 0, 0),
        'profesores_capacitados': capacitados,
        'profesores_padron': total_padron,
        'pct_cobertura': round((capacitados / total_padron) * 100, 1) if total_padron else 0.0,
        'promedio_por_taller': round(
            safe_int(totales['asistencias'] if totales else 0, 0) / totales['talleres'], 1
        ) if totales and totales['talleres'] else 0.0,
        'por_mes': rows_to_list(por_mes),
        'por_cargo': rows_to_list(por_cargo),
        'top_ies': rows_to_list(top_ies),
        'recientes': rows_to_list(recientes),
        'filtros': {'anio': anio, 'mes': mes, 'codigo_modular': codigo_ie},
    }


def _clave_persona(nombre):
    """Clave para agrupar a la misma persona escrita de formas distintas."""
    limpio = _normalize_place_text(nombre)
    return ' '.join(limpio.split())


def build_especialistas_monitoreo(conn, anio='', mes='', codigo_ie=''):
    """Produccion por especialista, contada por el ESPECIALISTA DEL REGISTRO.

    Importante: NO se cuenta por `creado_por` (quien subio el archivo). Un
    administrador puede digitalizar la ficha de otro; el credito es de quien
    hizo el monitoreo. Por eso se agrupa por:
      - fichas SIMON      -> fichas_monitoreo.monitor
      - informes de campo -> informe_campo_visita.especialista_detectado
      - talleres          -> taller_capacitacion.especialista_responsable

    Esos campos son texto libre (no hay FK), asi que se agrupan por nombre
    normalizado y se enlazan con app_user cuando el nombre coincide; si no
    coincide, la persona igual aparece, marcada como no vinculada.
    """
    hay_filtro_periodo = bool(str(anio or '').strip() or str(mes or '').strip())

    personas = {}

    def bucket(nombre_raw):
        clave = _clave_persona(nombre_raw)
        if not clave:
            return None
        if clave not in personas:
            personas[clave] = {
                'nombre': str(nombre_raw).strip(),
                'clave': clave,
                'user_id': '', 'email': '', 'rol_label': '', 'vinculado': False,
                'fichas_simon': 0, 'informes_campo': 0, 'talleres': 0,
                'asistentes_capacitados': 0, 'ultima_actividad': '', 'por_mes': {},
            }
        return personas[clave]

    def suma(persona, campo, fecha, cuantos=1):
        persona[campo] += cuantos
        fecha = str(fecha or '')
        if fecha:
            if fecha > persona['ultima_actividad']:
                persona['ultima_actividad'] = fecha
            mes_k = fecha[:7]
            if len(mes_k) == 7:
                persona['por_mes'][mes_k] = persona['por_mes'].get(mes_k, 0) + cuantos

    def filtro(campo_fecha, where, params):
        cond, valores = _condicion_periodo(campo_fecha, anio, mes)
        if cond:
            where.append(cond)
            params.extend(valores)

    # --- Fichas SIMON ---
    where = ['COALESCE(estado_fila, 1) = 1', "COALESCE(monitor, '') != ''"]
    params = []
    filtro('fecha_ejecucion', where, params)
    if codigo_ie:
        where.append('codigo_modular = ?')
        params.append(codigo_ie)
    consulta = 'SELECT monitor, fecha_ejecucion FROM fichas_monitoreo WHERE ' + ' AND '.join(where)
    for row in conn.execute(consulta, params).fetchall():
        persona = bucket(row['monitor'])
        if persona:
            suma(persona, 'fichas_simon', row['fecha_ejecucion'])

    # --- Informes de campo ---
    where = ['COALESCE(v.estado_fila, 1) = 1', "COALESCE(v.especialista_detectado, '') != ''"]
    params = []
    filtro('v.fecha_visita', where, params)
    if codigo_ie:
        where.append('v.codigo_modular = ?')
        params.append(codigo_ie)
    consulta = ('SELECT v.especialista_detectado, v.fecha_visita, v.creado_en '
                'FROM informe_campo_visita v WHERE ' + ' AND '.join(where))
    for row in conn.execute(consulta, params).fetchall():
        persona = bucket(row['especialista_detectado'])
        if persona:
            suma(persona, 'informes_campo', row['fecha_visita'] or row['creado_en'])

    # --- Talleres dictados ---
    where = ['COALESCE(t.estado_fila, 1) = 1']
    params = []
    filtro("COALESCE(NULLIF(t.fecha_inicio, ''), t.creado_en)", where, params)
    if codigo_ie:
        where.append('t.codigo_modular_sede = ?')
        params.append(codigo_ie)
    consulta = ('SELECT t.especialista_responsable, t.fecha_inicio, t.creado_en, '
                'u.nombre AS creador_nombre, '
                '(SELECT COUNT(*) FROM taller_asistencia a '
                '  WHERE a.taller_id = t.id AND COALESCE(a.estado_fila, 1) = 1) AS asistentes '
                'FROM taller_capacitacion t '
                'LEFT JOIN app_user u ON u.user_id = t.creado_por '
                'WHERE ' + ' AND '.join(where))
    for row in conn.execute(consulta, params).fetchall():
        # Si nadie declaro responsable se usa a quien lo registro, como ultimo
        # recurso; el campo declarado siempre manda.
        persona = bucket(row['especialista_responsable'] or row['creador_nombre'])
        if persona:
            suma(persona, 'talleres', row['fecha_inicio'] or row['creado_en'])
            persona['asistentes_capacitados'] += safe_int(row['asistentes'], 0)

    # --- Enlazar con usuarios de la app ---
    usuarios = conn.execute(
        'SELECT user_id, nombre, email, rol FROM app_user WHERE activo = 1'
    ).fetchall()
    por_clave = {_clave_persona(u['nombre']): u for u in usuarios}
    for clave, persona in personas.items():
        usuario = por_clave.get(clave)
        if usuario:
            persona['user_id'] = usuario['user_id']
            persona['email'] = usuario['email']
            persona['rol_label'] = role_label(usuario['rol'])
            persona['vinculado'] = True

    # Usuarios activos sin ningun registro: aparecen en cero para que se note
    # quien no ha cargado nada (solo cuando no hay filtros).
    if not hay_filtro_periodo and not codigo_ie:
        for usuario in usuarios:
            clave = _clave_persona(usuario['nombre'])
            if clave and clave not in personas:
                personas[clave] = {
                    'nombre': usuario['nombre'], 'clave': clave, 'user_id': usuario['user_id'],
                    'email': usuario['email'], 'rol_label': role_label(usuario['rol']),
                    'vinculado': True, 'fichas_simon': 0, 'informes_campo': 0, 'talleres': 0,
                    'asistentes_capacitados': 0, 'ultima_actividad': '', 'por_mes': {},
                }

    filas = []
    for persona in personas.values():
        persona['total_registros'] = (persona['fichas_simon'] + persona['informes_campo']
                                      + persona['talleres'])
        persona['por_mes'] = [{'mes': m, 'registros': n} for m, n in sorted(persona['por_mes'].items())]
        filas.append(persona)
    filas.sort(key=lambda f: (-f['total_registros'], f['nombre'] or ''))

    activos = [f for f in filas if f['total_registros'] > 0]
    return {
        'especialistas': filas,
        'con_actividad': len(activos),
        'sin_actividad': len(filas) - len(activos),
        'sin_vincular': len([f for f in filas if not f['vinculado'] and f['total_registros'] > 0]),
        'total_fichas': sum(f['fichas_simon'] for f in filas),
        'total_campo': sum(f['informes_campo'] for f in filas),
        'total_talleres': sum(f['talleres'] for f in filas),
        'filtros': {'anio': anio, 'mes': mes, 'codigo_modular': codigo_ie},
    }


def build_priorizacion_resumen(conn):
    """Lectura unificada de las tres fuentes: el score del informe v5 por IE.

    Es lo que responde "que tan priorizado esta el universo completo", en vez de
    mirar cada fuente por separado. Incluye la cobertura de evidencia porque sin
    ella el promedio de score se malinterpreta: casi toda IE tiene solo Censo.
    """
    fila = conn.execute("""
        SELECT COUNT(*) AS total,
               ROUND(AVG(COALESCE(score_priorizacion_ie, 0)), 1) AS score_promedio,
               ROUND(AVG(COALESCE(riesgo_simon_0_100, 0)), 1) AS simon_prom,
               ROUND(AVG(COALESCE(riesgo_infra_0_100, 0)), 1) AS infra_prom,
               ROUND(AVG(COALESCE(riesgo_campo_0_100, 0)), 1) AS campo_prom,
               ROUND(AVG(COALESCE(brecha_evidencia_0_100, 0)), 1) AS brecha_prom,
               SUM(CASE WHEN prioridad_v5 = 'Critica' THEN 1 ELSE 0 END) AS criticas,
               SUM(CASE WHEN prioridad_v5 = 'Alta' THEN 1 ELSE 0 END) AS altas,
               SUM(CASE WHEN prioridad_v5 = 'Media' THEN 1 ELSE 0 END) AS medias,
               SUM(CASE WHEN prioridad_v5 = 'Baja' THEN 1 ELSE 0 END) AS bajas,
               SUM(CASE WHEN COALESCE(brecha_evidencia_0_100, 100) = 0 THEN 1 ELSE 0 END) AS con_tres_fuentes,
               SUM(CASE WHEN COALESCE(simon_fichas, 0) > 0 THEN 1 ELSE 0 END) AS con_simon,
               SUM(CASE WHEN COALESCE(coverage_category, '') LIKE '%Censo%' THEN 1 ELSE 0 END) AS con_infra,
               SUM(CASE WHEN COALESCE(campo_documentos, 0) > 0 THEN 1 ELSE 0 END) AS con_campo
        FROM institucion_resumen
    """).fetchone()
    datos = dict(fila) if fila else {}
    total = safe_int(datos.get('total'), 0)
    prioritarias = safe_int(datos.get('criticas'), 0) + safe_int(datos.get('altas'), 0)
    datos['prioritarias'] = prioritarias
    datos['pct_prioritarias'] = round(prioritarias / total * 100, 1) if total else 0.0
    datos['pct_tres_fuentes'] = round(safe_int(datos.get('con_tres_fuentes'), 0) / total * 100, 1) if total else 0.0
    # Aporte medio de cada fuente al score, ya ponderado: hace visible por que el
    # promedio general es bajo (casi nadie tiene SIMON ni Campo).
    datos['aporte_simon'] = round(safe_float(datos.get('simon_prom'), 0) * SCORE_PESOS['simon'], 1)
    datos['aporte_infra'] = round(safe_float(datos.get('infra_prom'), 0) * SCORE_PESOS['infra'], 1)
    datos['aporte_campo'] = round(safe_float(datos.get('campo_prom'), 0) * SCORE_PESOS['campo'], 1)
    datos['aporte_brecha'] = round(safe_float(datos.get('brecha_prom'), 0) * SCORE_PESOS['brecha'], 1)
    return datos


def build_ie_coverage(conn):
    """Cuantos dias pasaron desde el ultimo contacto real con cada IE, por
    fuente (SIMON / informes de campo). No es lo mismo que el ranking de
    riesgo: una IE puede tener buen desempeno y aun asi llevar meses sin
    visita -- esta vista responde 'a quien no visitamos hace mas tiempo',
    no 'quien esta peor'. Solo incluye IE con al menos un evento real (fichas
    SIMON o informes de campo con fecha), para no listar 382 filas vacias."""
    today = datetime.now().date()

    simon_rows = conn.execute('''
        SELECT codigo_modular, MAX(fecha_ejecucion) AS ultima
        FROM fichas_monitoreo
        WHERE COALESCE(codigo_modular, '') != '' AND COALESCE(fecha_ejecucion, '') != ''
        GROUP BY codigo_modular
    ''').fetchall()
    campo_rows = conn.execute('''
        SELECT codigo_modular, MAX(fecha_visita) AS ultima
        FROM informe_campo_visita
        WHERE estado_fila = 1 AND COALESCE(codigo_modular, '') != '' AND COALESCE(fecha_visita, '') != ''
        GROUP BY codigo_modular
    ''').fetchall()

    def dias_desde(fecha_str):
        try:
            fecha = datetime.strptime(fecha_str[:10], '%Y-%m-%d').date()
            return (today - fecha).days
        except (ValueError, TypeError):
            return None

    por_ie = {}
    for row in simon_rows:
        entry = por_ie.setdefault(row['codigo_modular'], {'simon_fecha': None, 'campo_fecha': None})
        entry['simon_fecha'] = row['ultima']
    for row in campo_rows:
        entry = por_ie.setdefault(row['codigo_modular'], {'simon_fecha': None, 'campo_fecha': None})
        entry['campo_fecha'] = row['ultima']

    official_names = {}
    if por_ie:
        placeholders = ', '.join('?' for _ in por_ie)
        for r in conn.execute(
            f'SELECT codigo_modular, nombre_iiee, distrito FROM dim_institucion WHERE codigo_modular IN ({placeholders})',
            list(por_ie.keys()),
        ).fetchall():
            official_names[r['codigo_modular']] = {'nombre_iiee': r['nombre_iiee'], 'distrito': r['distrito']}

    items = []
    for codigo, entry in por_ie.items():
        info = official_names.get(codigo, {})
        dias_simon = dias_desde(entry['simon_fecha'])
        dias_campo = dias_desde(entry['campo_fecha'])
        items.append({
            'codigo_modular': codigo,
            'nombre_iiee': info.get('nombre_iiee') or codigo,
            'distrito': info.get('distrito'),
            'ultima_visita_simon': entry['simon_fecha'],
            'dias_desde_simon': dias_simon,
            'ultima_visita_campo': entry['campo_fecha'],
            'dias_desde_campo': dias_campo,
            'dias_max': max(d for d in (dias_simon, dias_campo) if d is not None),
        })
    items.sort(key=lambda item: item['dias_max'], reverse=True)

    return {
        'ies': items,
        'total_ie': len(items),
        'sin_visita_simon': sum(1 for i in items if i['dias_desde_simon'] is None),
        'sin_visita_campo': sum(1 for i in items if i['dias_desde_campo'] is None),
    }

def build_institucion_evolucion(conn, codigo):
    """Linea de tiempo combinada SIMON + informes de campo para UNA IE,
    ordenada de mas reciente a mas antigua -- responde 'como va esta IE en
    el tiempo' con eventos reales, sin inventar puntos intermedios."""
    eventos = []
    for row in conn.execute('''
        SELECT fecha_ejecucion, docente, promedio, a1, a2, a3, b1, b2, b3, b4, b5
        FROM fichas_monitoreo WHERE codigo_modular = ?
        ORDER BY fecha_ejecucion
    ''', (codigo,)).fetchall():
        item = dict(row)
        values = simon_values_from_row(item)
        riesgo = simon_documented_risk(values)
        eventos.append({
            'fecha': item.get('fecha_ejecucion') or None,
            'fuente': 'simon',
            'titulo': f"Ficha SIMON · {item.get('docente') or 'Docente sin nombre'}",
            'detalle': f"Promedio {item.get('promedio') or 0} · riesgo {riesgo['riesgo_intervencion']}" if riesgo else f"Promedio {item.get('promedio') or 0}",
            'nivel': riesgo['riesgo_intervencion'] if riesgo else None,
        })

    for row in conn.execute('''
        SELECT v.fecha_visita, v.especialista_detectado, d.nombre_archivo,
        ''' + ', '.join(CAMPO_CONCRETE_VARIABLES) + '''
        FROM informe_campo_visita v
        LEFT JOIN informe_campo_documento d ON d.id = v.documento_id
        WHERE v.codigo_modular = ? AND v.estado_fila = 1
    ''', (codigo,)).fetchall():
        item = dict(row)
        variables = [v for v in CAMPO_CONCRETE_VARIABLES if item.get(v)]
        nivel = campo_severidad_nivel(variables)
        labels = [CAMPO_VARIABLE_META.get(v, {}).get('label', v) for v in variables]
        eventos.append({
            'fecha': item.get('fecha_visita') or None,
            'fuente': 'campo',
            'titulo': f"Informe de campo · {item.get('especialista_detectado') or 'Especialista no identificado'}",
            'detalle': ', '.join(labels) if labels else 'Sin variables de alerta detectadas',
            'nivel': nivel,
            'documento': item.get('nombre_archivo'),
        })

    eventos.sort(key=lambda e: e['fecha'] or '', reverse=True)
    con_fecha = [e for e in eventos if e['fecha']]
    sin_fecha = [e for e in eventos if not e['fecha']]
    return {'eventos': con_fecha + sin_fecha, 'total': len(eventos)}

# ─── AUTENTICACION ────────────────────────────────────────────────────────────

def get_current_user():
    """Usuario autenticado a partir del header 'Authorization: Bearer <token>'.

    Devuelve None si no hay token, el token no existe, expiro, o el usuario
    fue desactivado. No lanza excepciones.
    """
    auth_header = request.headers.get('Authorization', '')
    token = auth_header[7:].strip() if auth_header.lower().startswith('bearer ') else ''
    if not token:
        return None
    conn = get_db()
    row = conn.execute('''
        SELECT u.* FROM app_session s
        JOIN app_user u ON u.user_id = s.user_id
        WHERE s.token = ? AND s.expira_en > CURRENT_TIMESTAMP
    ''', (token,)).fetchone()
    if not row or not row['activo']:
        return None
    return row

def require_auth(fn):
    """Exige una sesion valida (cualquier rol) para acceder al endpoint."""
    @wraps(fn)
    def wrapper(*args, **kwargs):
        user = get_current_user()
        if not user:
            return jsonify({'error': 'Sesión inválida o expirada. Vuelve a iniciar sesión.'}), 401
        g.current_user = user
        return fn(*args, **kwargs)
    return wrapper

def can_edit_row(user, owner_user_id):
    """Administradores y supervisores pueden editar cualquier fila; un
    especialista solo las suyas.

    Filas sin dueno (creado_por/subido_por NULL, tipicamente registros
    historicos importados) solo las puede editar quien edita todo.
    """
    if normalize_role(user['rol']) in ROLES_QUE_EDITAN_TODO:
        return True
    return bool(owner_user_id) and owner_user_id == user['user_id']

def require_admin(fn):
    """Exige una sesion valida con rol administrador.

    Se reserva para la administracion del sistema (usuarios, instituciones,
    profesores, preguntas dinamicas). Para lo que un supervisor tambien debe
    poder hacer, usar require_roles.
    """
    @wraps(fn)
    def wrapper(*args, **kwargs):
        user = get_current_user()
        if not user:
            return jsonify({'error': 'Sesión inválida o expirada. Vuelve a iniciar sesión.'}), 401
        if normalize_role(user['rol']) != 'administrador':
            return jsonify({'error': 'Esta acción requiere permisos de administrador.'}), 403
        g.current_user = user
        return fn(*args, **kwargs)
    return wrapper

def require_roles(*roles_permitidos):
    """Exige una sesion valida cuyo rol este en la lista dada."""
    permitidos = {normalize_role(r) for r in roles_permitidos}

    def decorador(fn):
        @wraps(fn)
        def wrapper(*args, **kwargs):
            user = get_current_user()
            if not user:
                return jsonify({'error': 'Sesión inválida o expirada. Vuelve a iniciar sesión.'}), 401
            if normalize_role(user['rol']) not in permitidos:
                return jsonify({'error': 'No tienes permisos para esta acción.'}), 403
            g.current_user = user
            return fn(*args, **kwargs)
        return wrapper
    return decorador

# ─── HEALTH / API INDEX ───────────────────────────────────────────────────────

@app.route('/', methods=['GET'])
@app.route('/index.html', methods=['GET'])
def frontend_index():
    return send_from_directory(str(ROOT / 'frontend'), 'index.html')


@app.route('/sw.js', methods=['GET'])
def service_worker():
    return send_from_directory(str(ROOT / 'frontend'), 'sw.js')


@app.route('/logo/<path:filename>', methods=['GET'])
def logo_assets(filename):
    return send_from_directory(str(ROOT / 'frontend' / 'logo'), filename)


@app.route('/img/<path:filename>', methods=['GET'])
def img_assets(filename):
    return send_from_directory(str(ROOT / 'frontend' / 'img'), filename)


@app.route('/api', methods=['GET'])
def api_index():
    azure_endpoint, azure_key = get_azure_config()
    return jsonify({
        'service': 'SUGKA LAB API',
        'status': 'ok',
        'frontend_url': request.host_url.rstrip('/'),
        'ocr_status': {
            'azure_configured': bool(azure_endpoint and azure_key),
            'gemini_configured': bool(os.getenv("GEMINI_API_KEY", "")),
            'fallback': 'windows_ocr',
        },
        'endpoints': {
            'dashboard': '/api/dashboard',
            'especialistas': '/api/especialistas',
            'usuarios': '/api/usuarios',
            'instrumentos': '/api/instrumentos',
            'preguntas_dinamicas': '/api/instrumentos/dinamico/preguntas',
            'ficha_vacia_pdf': '/api/ficha-vacia/pdf',
            'instituciones': '/api/instituciones',
            'alertas': '/api/alertas',
            'fichas': '/api/fichas',
            'ocr_upload': '/api/ocr/upload_advanced',
            'sync': '/api/sync',
            'login': '/api/login',
        }
    })

# ─── ESPECIALISTAS / AUTH ─────────────────────────────────────────────────────

@app.route('/api/especialistas', methods=['GET'])
@require_auth
def get_especialistas():
    conn = get_db()
    rows = conn.execute(
        'SELECT especialista_id, nombre, rol_inferido, total_fichas_simon, total_documentos_campo '
        'FROM dim_especialista ORDER BY nombre'
    ).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))

@app.route('/api/usuarios', methods=['GET', 'POST'])
@require_auth
def usuarios():
    conn = get_db()
    try:
        if request.method == 'GET':
            rows = conn.execute('''
                SELECT u.*, COALESCE(u.cargo, e.rol_inferido) AS especialidad
                FROM app_user u
                LEFT JOIN dim_especialista e ON u.especialista_id = e.especialista_id
                ORDER BY
                    CASE u.rol WHEN 'administrador' THEN 0 WHEN 'supervisor' THEN 1 ELSE 2 END,
                    u.nombre
            ''').fetchall()
            return jsonify([public_user(r) for r in rows])

        # Crear usuarios es una accion administrativa.
        if normalize_role(g.current_user['rol']) != 'administrador':
            return jsonify({'error': 'Esta acción requiere permisos de administrador.'}), 403

        data = request.get_json() or {}
        password = str(data.get('password') or '').strip()
        if len(password) < 4:
            return jsonify({'error': 'La contraseña debe tener al menos 4 caracteres.'}), 400

        try:
            user = create_app_user(
                conn,
                nombre=data.get('nombre', ''),
                username=data.get('username') or data.get('email') or data.get('correo') or '',
                email=data.get('email') or data.get('correo') or '',
                password=password,
                rol=data.get('rol', 'especialista'),
                cargo=data.get('cargo') or None,
                activo=data.get('activo', 1),
            )
            conn.commit()
            return jsonify({'success': True, 'user': public_user(user)})
        except sqlite3.IntegrityError:
            conn.rollback()
            return jsonify({'error': 'Ese usuario ya existe.'}), 400
        except ValueError as exc:
            conn.rollback()
            return jsonify({'error': str(exc)}), 400
    finally:
        conn.close()

@app.route('/api/usuarios/<user_id>', methods=['PUT'])
@require_admin
def actualizar_usuario(user_id):
    data = request.get_json() or {}
    conn = get_db()
    try:
        user = conn.execute('SELECT * FROM app_user WHERE user_id = ?', (user_id,)).fetchone()
        if not user:
            return jsonify({'error': 'Usuario no encontrado.'}), 404

        updates = []
        params = []
        for key in ('nombre', 'username', 'email', 'cargo'):
            if key in data:
                updates.append(f'{key} = ?')
                value = data.get(key)
                params.append(str(value).strip().lower() if key in ('username', 'email') else (value or None))
        if 'rol' in data:
            nuevo_rol = normalize_role(data.get('rol'))
            # Con tres roles ya es posible dejar el sistema sin administradores
            # degradando al ultimo a supervisor, no solo desactivandolo: se
            # bloquean ambos caminos con la misma regla.
            if normalize_role(user['rol']) == 'administrador' and nuevo_rol != 'administrador':
                otros_admins = conn.execute(
                    "SELECT COUNT(*) FROM app_user WHERE rol = 'administrador' AND activo = 1 AND user_id != ?",
                    (user_id,),
                ).fetchone()[0]
                if otros_admins < 1:
                    return jsonify({'error': 'Debe quedar al menos un administrador activo.'}), 400
            updates.append('rol = ?')
            params.append(nuevo_rol)
        if 'activo' in data:
            new_active = 1 if data.get('activo') else 0
            if normalize_role(user['rol']) == 'administrador' and new_active == 0:
                active_admins = conn.execute(
                    "SELECT COUNT(*) FROM app_user WHERE rol = 'administrador' AND activo = 1"
                ).fetchone()[0]
                if active_admins <= 1:
                    return jsonify({'error': 'Debe quedar al menos un administrador activo.'}), 400
            updates.append('activo = ?')
            params.append(new_active)

        if not updates:
            return jsonify({'success': True, 'user': public_user(user)})

        updates.append('actualizado_en = CURRENT_TIMESTAMP')
        params.append(user_id)
        conn.execute(f'UPDATE app_user SET {", ".join(updates)} WHERE user_id = ?', params)
        conn.commit()
        updated = conn.execute('SELECT * FROM app_user WHERE user_id = ?', (user_id,)).fetchone()
        return jsonify({'success': True, 'user': public_user(updated)})
    except sqlite3.IntegrityError:
        conn.rollback()
        return jsonify({'error': 'Ese usuario ya existe.'}), 400
    finally:
        conn.close()

@app.route('/api/usuarios/<user_id>/password', methods=['PUT'])
@require_auth
def cambiar_password_usuario(user_id):
    """Un administrador puede cambiar la contraseña de cualquiera sin
    verificarla (reseteo administrativo). Un usuario cambiando SU PROPIA
    contraseña (autoservicio, ver 'Mi cuenta') debe confirmar la actual --
    evita que una sesion robada cambie la contraseña sin conocerla."""
    is_admin = normalize_role(g.current_user['rol']) == 'administrador'
    is_self = str(g.current_user['user_id']) == str(user_id)
    if not is_admin and not is_self:
        return jsonify({'error': 'No tienes permiso para cambiar la contraseña de otro usuario.'}), 403

    data = request.get_json() or {}
    password = str(data.get('password') or '').strip()
    if len(password) < 4:
        return jsonify({'error': 'La contraseña debe tener al menos 4 caracteres.'}), 400

    conn = get_db()
    try:
        user = conn.execute('SELECT * FROM app_user WHERE user_id = ?', (user_id,)).fetchone()
        if not user:
            return jsonify({'error': 'Usuario no encontrado.'}), 404

        if is_self:
            # Un administrador cambiando su PROPIA contraseña también debe confirmar
            # la actual -- el bypass sin verificación es solo para cuando un admin
            # resetea la contraseña de OTRO usuario (is_self=False).
            actual = str(data.get('password_actual') or '').strip()
            if not verify_password(actual, user['password_salt'], user['password_hash']):
                return jsonify({'error': 'La contraseña actual no es correcta.'}), 400

        salt, digest = hash_password(password)
        conn.execute('''
            UPDATE app_user
            SET password_salt = ?, password_hash = ?, actualizado_en = CURRENT_TIMESTAMP
            WHERE user_id = ?
        ''', (salt, digest, user_id))
        conn.commit()
        updated = conn.execute('SELECT * FROM app_user WHERE user_id = ?', (user_id,)).fetchone()
        return jsonify({'success': True, 'user': public_user(updated)})
    finally:
        conn.close()

@app.route('/api/usuarios/<user_id>', methods=['DELETE'])
@require_admin
def eliminar_usuario(user_id):
    """Elimina un usuario -- en realidad un soft-delete (activo=0), nunca un
    DELETE real de la fila, para conservar el historico y la auditoria de
    quien hizo que. app_user ya tenia esta columna 'activo' desde antes, asi
    que no se duplica con un 'estado_fila' nuevo; es el mismo estandar."""
    conn = get_db()
    try:
        user = conn.execute('SELECT * FROM app_user WHERE user_id = ?', (user_id,)).fetchone()
        if not user:
            return jsonify({'error': 'Usuario no encontrado.'}), 404
        if normalize_role(user['rol']) == 'administrador':
            active_admins = conn.execute(
                "SELECT COUNT(*) FROM app_user WHERE rol = 'administrador' AND activo = 1"
            ).fetchone()[0]
            if active_admins <= 1:
                return jsonify({'error': 'Debe quedar al menos un administrador activo.'}), 400
        if user['user_id'] == g.current_user['user_id']:
            return jsonify({'error': 'No puedes eliminar tu propio usuario.'}), 400
        conn.execute(
            "UPDATE app_user SET activo = 0, actualizado_en = CURRENT_TIMESTAMP WHERE user_id = ?",
            (user_id,)
        )
        conn.commit()
        return jsonify({'success': True})
    finally:
        conn.close()

# ─── INSTRUMENTOS / PREGUNTAS ────────────────────────────────────────────────

@app.route('/api/instrumentos', methods=['GET'])
@require_auth
def get_instrumentos():
    conn = get_db()
    try:
        rows = conn.execute('''
            SELECT *
            FROM app_instrumento
            WHERE activo = 1
            ORDER BY CASE tipo WHEN 'fija' THEN 0 ELSE 1 END, nombre
        ''').fetchall()
        instrumentos = []
        for row in rows:
            instrumentos.append(build_instrument_payload(conn, row))
        return jsonify(instrumentos)
    finally:
        conn.close()

@app.route('/api/instrumentos/dinamico/preguntas', methods=['GET', 'POST'])
@require_auth
def preguntas_dinamicas():
    # GET es de lectura (cualquier usuario autenticado): lo necesita el
    # formulario de fichas para mostrar la descripcion de cada pregunta
    # dinamica al editar, igual que las preguntas fijas de SIMON. Crear/editar
    # preguntas si sigue siendo solo de administrador.
    if request.method != 'GET' and normalize_role(g.current_user['rol']) != 'administrador':
        return jsonify({'error': 'Esta acción requiere permisos de administrador.'}), 403
    conn = get_db()
    try:
        payload = dynamic_questions_payload(conn)
        if not payload:
            return jsonify({'error': 'No hay instrumento dinámico activo.'}), 404

        if request.method == 'GET':
            return jsonify(payload)

        data = request.get_json() or {}
        instrumento_id = payload['instrumento']['instrumento_id']
        seccion_id = safe_int(data.get('seccion_id'), 0)
        codigo = normalize_question_code(data.get('codigo'))
        item = str(data.get('item') or '').strip()
        niveles = normalize_levels_payload(data.get('niveles'))
        activo = 1 if data.get('activo', 1) else 0

        if not seccion_id or not codigo or not item:
            return jsonify({'error': 'Sección, código e ítem son obligatorios.'}), 400
        if any(not value for value in niveles.values()):
            return jsonify({'error': 'Completa la descripción de los cuatro niveles.'}), 400

        section = conn.execute('''
            SELECT *
            FROM app_instrumento_seccion
            WHERE instrumento_id = ? AND seccion_id = ?
        ''', (instrumento_id, seccion_id)).fetchone()
        if not section:
            return jsonify({'error': 'La sección seleccionada no pertenece al instrumento dinámico.'}), 400

        exists = conn.execute('''
            SELECT 1
            FROM app_instrumento_pregunta
            WHERE instrumento_id = ? AND UPPER(codigo) = UPPER(?)
            LIMIT 1
        ''', (instrumento_id, codigo)).fetchone()
        if exists:
            return jsonify({'error': 'Ya existe una pregunta dinámica con ese código.'}), 400

        order_row = conn.execute('''
            SELECT COALESCE(MAX(orden), 0) + 1 AS next_order
            FROM app_instrumento_pregunta
            WHERE instrumento_id = ? AND seccion_id = ?
        ''', (instrumento_id, seccion_id)).fetchone()
        next_order = safe_int(order_row['next_order'], 1) if order_row else 1

        cursor = conn.execute('''
            INSERT INTO app_instrumento_pregunta (
                instrumento_id, seccion_id, codigo, item, tipo_respuesta,
                niveles_json, opciones_json, orden, activo
            )
            VALUES (?, ?, ?, ?, 'nivel', ?, ?, ?, ?)
        ''', (
            instrumento_id,
            seccion_id,
            codigo,
            item,
            json_dumps(niveles),
            json_dumps({}),
            next_order,
            activo,
        ))
        sync_instrument_structure_json(conn, instrumento_id)
        conn.commit()
        updated = dynamic_questions_payload(conn)
        question = next(
            (q for q in updated['preguntas'] if q['pregunta_id'] == cursor.lastrowid),
            None,
        )
        return jsonify({'success': True, 'pregunta': question, 'data': updated})
    except sqlite3.IntegrityError:
        conn.rollback()
        return jsonify({'error': 'No se pudo guardar la pregunta dinámica.'}), 400
    finally:
        conn.close()

@app.route('/api/instrumentos/dinamico/preguntas/<int:pregunta_id>', methods=['PUT', 'DELETE'])
@require_admin
def pregunta_dinamica_detalle(pregunta_id):
    conn = get_db()
    try:
        instrument = get_dynamic_instrument_row(conn)
        if not instrument:
            return jsonify({'error': 'No hay instrumento dinámico activo.'}), 404

        question = conn.execute('''
            SELECT *
            FROM app_instrumento_pregunta
            WHERE pregunta_id = ? AND instrumento_id = ?
        ''', (pregunta_id, instrument['instrumento_id'])).fetchone()
        if not question:
            return jsonify({'error': 'Pregunta dinámica no encontrada.'}), 404

        if request.method == 'DELETE':
            conn.execute('DELETE FROM app_instrumento_pregunta WHERE pregunta_id = ?', (pregunta_id,))
            sync_instrument_structure_json(conn, instrument['instrumento_id'])
            conn.commit()
            return jsonify({'success': True, 'deleted_id': pregunta_id, 'data': dynamic_questions_payload(conn)})

        data = request.get_json() or {}
        seccion_id = safe_int(data.get('seccion_id', question['seccion_id']), question['seccion_id'])
        codigo = normalize_question_code(data.get('codigo', question['codigo']))
        item = str(data.get('item', question['item']) or '').strip()
        niveles = normalize_levels_payload(data.get('niveles', json_loads(question['niveles_json'], {})))
        activo = 1 if data.get('activo', question['activo']) else 0

        if not seccion_id or not codigo or not item:
            return jsonify({'error': 'Sección, código e ítem son obligatorios.'}), 400
        if any(not value for value in niveles.values()):
            return jsonify({'error': 'Completa la descripción de los cuatro niveles.'}), 400

        section = conn.execute('''
            SELECT *
            FROM app_instrumento_seccion
            WHERE instrumento_id = ? AND seccion_id = ?
        ''', (instrument['instrumento_id'], seccion_id)).fetchone()
        if not section:
            return jsonify({'error': 'La sección seleccionada no pertenece al instrumento dinámico.'}), 400

        exists = conn.execute('''
            SELECT 1
            FROM app_instrumento_pregunta
            WHERE instrumento_id = ? AND UPPER(codigo) = UPPER(?) AND pregunta_id != ?
            LIMIT 1
        ''', (instrument['instrumento_id'], codigo, pregunta_id)).fetchone()
        if exists:
            return jsonify({'error': 'Ya existe una pregunta dinámica con ese código.'}), 400

        conn.execute('''
            UPDATE app_instrumento_pregunta
            SET seccion_id = ?, codigo = ?, item = ?, niveles_json = ?, activo = ?
            WHERE pregunta_id = ? AND instrumento_id = ?
        ''', (
            seccion_id,
            codigo,
            item,
            json_dumps(niveles),
            activo,
            pregunta_id,
            instrument['instrumento_id'],
        ))
        sync_instrument_structure_json(conn, instrument['instrumento_id'])
        conn.commit()
        updated = dynamic_questions_payload(conn)
        saved = next((q for q in updated['preguntas'] if q['pregunta_id'] == pregunta_id), None)
        return jsonify({'success': True, 'pregunta': saved, 'data': updated})
    except sqlite3.IntegrityError:
        conn.rollback()
        return jsonify({'error': 'No se pudo actualizar la pregunta dinámica.'}), 400
    finally:
        conn.close()

@app.route('/api/instrumentos/dinamico/secciones', methods=['GET', 'POST'])
@require_admin
def secciones_dinamicas():
    conn = get_db()
    try:
        instrument = get_dynamic_instrument_row(conn)
        if not instrument:
            return jsonify({'error': 'No hay instrumento dinámico activo.'}), 404
        instrumento_id = instrument['instrumento_id']

        if request.method == 'GET':
            return jsonify(dynamic_questions_payload(conn)['secciones'])

        data = request.get_json() or {}
        clave = normalize_question_code(data.get('clave'))
        nombre = str(data.get('nombre') or '').strip()
        if not clave or not nombre:
            return jsonify({'error': 'Clave y nombre son obligatorios.'}), 400

        exists = conn.execute('''
            SELECT 1 FROM app_instrumento_seccion
            WHERE instrumento_id = ? AND UPPER(clave) = UPPER(?)
            LIMIT 1
        ''', (instrumento_id, clave)).fetchone()
        if exists:
            return jsonify({'error': 'Ya existe una sección con esa clave.'}), 400

        if data.get('orden') not in (None, ''):
            orden = safe_int(data.get('orden'), 0)
        else:
            order_row = conn.execute('''
                SELECT COALESCE(MAX(orden), 0) + 1 AS next_order
                FROM app_instrumento_seccion WHERE instrumento_id = ?
            ''', (instrumento_id,)).fetchone()
            orden = safe_int(order_row['next_order'], 1) if order_row else 1

        cursor = conn.execute('''
            INSERT INTO app_instrumento_seccion (instrumento_id, clave, nombre, orden, activo)
            VALUES (?, ?, ?, ?, 1)
        ''', (instrumento_id, clave, nombre, orden))
        sync_instrument_structure_json(conn, instrumento_id)
        conn.commit()
        secciones = dynamic_questions_payload(conn)['secciones']
        seccion = next((s for s in secciones if s['seccion_id'] == cursor.lastrowid), None)
        return jsonify({'success': True, 'seccion': seccion, 'secciones': secciones})
    except sqlite3.IntegrityError:
        conn.rollback()
        return jsonify({'error': 'No se pudo guardar la sección.'}), 400
    finally:
        conn.close()

@app.route('/api/instrumentos/dinamico/secciones/<int:seccion_id>', methods=['PUT', 'DELETE'])
@require_admin
def seccion_dinamica_detalle(seccion_id):
    conn = get_db()
    try:
        instrument = get_dynamic_instrument_row(conn)
        if not instrument:
            return jsonify({'error': 'No hay instrumento dinámico activo.'}), 404
        instrumento_id = instrument['instrumento_id']

        section = conn.execute('''
            SELECT * FROM app_instrumento_seccion
            WHERE seccion_id = ? AND instrumento_id = ?
        ''', (seccion_id, instrumento_id)).fetchone()
        if not section:
            return jsonify({'error': 'Sección no encontrada.'}), 404

        if request.method == 'DELETE':
            preguntas_count = conn.execute(
                'SELECT COUNT(*) FROM app_instrumento_pregunta WHERE seccion_id = ?', (seccion_id,)
            ).fetchone()[0]
            if preguntas_count:
                return jsonify({
                    'error': f'Esta sección tiene {preguntas_count} pregunta(s). '
                             'Reasígnalas o elimínalas antes de borrar la sección.'
                }), 400
            conn.execute('DELETE FROM app_instrumento_seccion WHERE seccion_id = ?', (seccion_id,))
            sync_instrument_structure_json(conn, instrumento_id)
            conn.commit()
            return jsonify({'success': True, 'deleted_id': seccion_id, 'secciones': dynamic_questions_payload(conn)['secciones']})

        data = request.get_json() or {}
        clave = normalize_question_code(data.get('clave', section['clave']))
        nombre = str(data.get('nombre', section['nombre']) or '').strip()
        orden = safe_int(data.get('orden', section['orden']), section['orden'])
        activo = 1 if data.get('activo', section['activo']) else 0
        if not clave or not nombre:
            return jsonify({'error': 'Clave y nombre son obligatorios.'}), 400

        exists = conn.execute('''
            SELECT 1 FROM app_instrumento_seccion
            WHERE instrumento_id = ? AND UPPER(clave) = UPPER(?) AND seccion_id != ?
            LIMIT 1
        ''', (instrumento_id, clave, seccion_id)).fetchone()
        if exists:
            return jsonify({'error': 'Ya existe una sección con esa clave.'}), 400

        conn.execute('''
            UPDATE app_instrumento_seccion
            SET clave = ?, nombre = ?, orden = ?, activo = ?
            WHERE seccion_id = ?
        ''', (clave, nombre, orden, activo, seccion_id))
        sync_instrument_structure_json(conn, instrumento_id)
        conn.commit()
        secciones = dynamic_questions_payload(conn)['secciones']
        seccion = next((s for s in secciones if s['seccion_id'] == seccion_id), None)
        return jsonify({'success': True, 'seccion': seccion, 'secciones': secciones})
    except sqlite3.IntegrityError:
        conn.rollback()
        return jsonify({'error': 'No se pudo actualizar la sección.'}), 400
    finally:
        conn.close()

@app.route('/api/instrumentos/dinamico/preguntas/<int:pregunta_id>/comentarios', methods=['GET', 'POST'])
@require_auth
def pregunta_comentarios(pregunta_id):
    conn = get_db()
    try:
        pregunta = conn.execute(
            'SELECT pregunta_id FROM app_instrumento_pregunta WHERE pregunta_id = ?', (pregunta_id,)
        ).fetchone()
        if not pregunta:
            return jsonify({'error': 'Pregunta no encontrada.'}), 404

        if request.method == 'GET':
            rows = conn.execute('''
                SELECT c.*, u.nombre AS autor_nombre
                FROM app_instrumento_pregunta_comentario c
                LEFT JOIN app_user u ON u.user_id = c.autor_id
                WHERE c.pregunta_id = ? AND c.estado_fila = 1
                ORDER BY c.creado_en
            ''', (pregunta_id,)).fetchall()
            return jsonify(rows_to_list(rows))

        data = request.get_json() or {}
        texto = str(data.get('texto') or '').strip()
        if not texto:
            return jsonify({'error': 'El comentario no puede estar vacío.'}), 400
        cur = conn.execute('''
            INSERT INTO app_instrumento_pregunta_comentario (pregunta_id, autor_id, texto)
            VALUES (?, ?, ?)
        ''', (pregunta_id, g.current_user['user_id'], texto))
        conn.commit()
        row = conn.execute('''
            SELECT c.*, u.nombre AS autor_nombre
            FROM app_instrumento_pregunta_comentario c
            LEFT JOIN app_user u ON u.user_id = c.autor_id
            WHERE c.id = ?
        ''', (cur.lastrowid,)).fetchone()
        return jsonify({'success': True, 'comentario': dict(row)})
    finally:
        conn.close()

@app.route('/api/instrumentos/dinamico/preguntas/<int:pregunta_id>/comentarios/<int:comentario_id>', methods=['DELETE'])
@require_auth
def pregunta_comentario_detalle(pregunta_id, comentario_id):
    conn = get_db()
    try:
        comentario = conn.execute('''
            SELECT * FROM app_instrumento_pregunta_comentario
            WHERE id = ? AND pregunta_id = ?
        ''', (comentario_id, pregunta_id)).fetchone()
        if not comentario:
            return jsonify({'error': 'Comentario no encontrado.'}), 404
        if not can_edit_row(g.current_user, comentario['autor_id']):
            return jsonify({'error': 'Solo puedes eliminar tus propios comentarios.'}), 403
        conn.execute(
            'UPDATE app_instrumento_pregunta_comentario SET estado_fila = 0 WHERE id = ?', (comentario_id,)
        )
        conn.commit()
        return jsonify({'success': True})
    finally:
        conn.close()

@app.route('/api/archivos/<int:archivo_id>', methods=['GET'])
@require_auth
def archivo_original(archivo_id):
    """Sirve el archivo original (imagen/PDF) subido para un OCR de ficha o
    de informe de campo, para que se pueda revisar contra lo que quedo
    registrado. Cualquier usuario autenticado puede verlo (misma regla que
    'ver' un registro ajeno), no solo el dueno."""
    conn = get_db()
    try:
        row = conn.execute(
            'SELECT * FROM archivo_subido WHERE id = ? AND COALESCE(estado_fila, 1) = 1',
            (archivo_id,),
        ).fetchone()
        if not row or not row['contenido']:
            return jsonify({'error': 'Archivo no encontrado.'}), 404
        return send_file(
            io.BytesIO(row['contenido']),
            mimetype=row['mimetype'] or 'application/octet-stream',
            as_attachment=False,
            download_name=row['nombre_archivo'] or f'archivo_{archivo_id}',
        )
    finally:
        conn.close()

@app.route('/api/ficha-vacia/pdf', methods=['GET'])
def ficha_vacia_pdf():
    # Publico a proposito: es solo la plantilla vacia (sin datos de ninguna
    # IE) y el frontend la descarga con un <a href> normal, que no puede
    # adjuntar el header Authorization.
    conn = get_db()
    try:
        rows = conn.execute('''
            SELECT *
            FROM app_instrumento
            WHERE activo = 1
            ORDER BY CASE tipo WHEN 'fija' THEN 0 ELSE 1 END, nombre
        ''').fetchall()
        instrumentos = [build_instrument_payload(conn, row) for row in rows]
        pdf_buffer = build_empty_ficha_pdf(instrumentos)
        return send_file(
            pdf_buffer,
            mimetype='application/pdf',
            as_attachment=True,
            download_name='ficha_monitoreo_vacia_sugka.pdf',
        )
    finally:
        conn.close()

SESSION_TTL_HOURS = safe_int(os.getenv('SESSION_TTL_HOURS'), 12) or 12
EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')

def build_login_payload(conn, user):
    """Arma la respuesta de sesion (usuario + token) para login y registro.

    Emite un token de sesion nuevo, lo guarda en app_session y de paso
    limpia sesiones vencidas de cualquier usuario (mantenimiento barato).
    """
    public = public_user(user)
    public['rol_inferido'] = public['rol_label']
    public['rol'] = normalize_role(user['rol'])
    public['especialidad'] = public.get('cargo') or ''

    if user['especialista_id']:
        esp = conn.execute(
            'SELECT * FROM dim_especialista WHERE especialista_id = ?',
            (user['especialista_id'],)
        ).fetchone()
        if esp:
            public['especialista'] = dict(esp)
            if not public['especialidad']:
                public['especialidad'] = esp['rol_inferido']
            public['total_fichas_simon'] = esp['total_fichas_simon']
            public['total_documentos_campo'] = esp['total_documentos_campo']

    token = secrets.token_urlsafe(32)
    expira_en = (datetime.utcnow() + timedelta(hours=SESSION_TTL_HOURS)).strftime('%Y-%m-%d %H:%M:%S')
    conn.execute(
        'INSERT INTO app_session (token, user_id, expira_en) VALUES (?, ?, ?)',
        (token, user['user_id'], expira_en),
    )
    conn.execute('DELETE FROM app_session WHERE expira_en <= CURRENT_TIMESTAMP')
    conn.commit()

    return {'success': True, 'user': public, 'especialista': public, 'token': token}

@app.route('/api/login', methods=['POST'])
def login():
    data = request.get_json() or {}
    username = str(data.get('username') or data.get('usuario') or data.get('email') or '').strip().lower()
    password = str(data.get('password') or data.get('clave') or data.get('pin') or '')

    conn = get_db()
    try:
        user = conn.execute(
            '''
            SELECT * FROM app_user
            WHERE LOWER(username) = LOWER(?) OR LOWER(COALESCE(email, '')) = LOWER(?)
            ''',
            (username, username),
        ).fetchone()

        if not user or not user['activo']:
            return jsonify({'error': 'Usuario no encontrado o inactivo. Contacte con el administrador.'}), 401

        if not verify_password(password, user['password_salt'], user['password_hash']):
            return jsonify({'error': 'Contraseña incorrecta. Si la olvidaste, contacte con el administrador.'}), 401

        return jsonify(build_login_payload(conn, user))
    finally:
        conn.close()

_registro_attempts = {}
REGISTRO_MAX_POR_HORA = 8

def registro_rate_limited(ip):
    """Limite simple en memoria (por IP) para frenar registros masivos.

    Vale para un solo proceso worker (el que usa este servicio en
    render.yaml, --workers 1); con varios workers cada uno llevaria su
    propio contador, lo cual sigue siendo una mitigacion razonable aunque
    no perfecta.
    """
    now = datetime.utcnow().timestamp()
    attempts = [t for t in _registro_attempts.get(ip, []) if now - t < 3600]
    attempts.append(now)
    _registro_attempts[ip] = attempts
    return len(attempts) > REGISTRO_MAX_POR_HORA

@app.route('/api/registro', methods=['POST'])
def registro():
    if registro_rate_limited(request.remote_addr or 'unknown'):
        return jsonify({'error': 'Demasiados intentos de registro. Intenta de nuevo más tarde.'}), 429

    data = request.get_json() or {}
    nombre = str(data.get('nombre') or '').strip()
    email = str(data.get('email') or data.get('correo') or '').strip().lower()
    password = str(data.get('password') or '').strip()

    if not nombre:
        return jsonify({'error': 'El nombre es obligatorio.'}), 400
    if not EMAIL_RE.match(email):
        return jsonify({'error': 'Ingresa un correo electrónico válido.'}), 400
    if len(password) < 6:
        return jsonify({'error': 'La contraseña debe tener al menos 6 caracteres.'}), 400

    conn = get_db()
    try:
        existing = conn.execute(
            'SELECT 1 FROM app_user WHERE LOWER(email) = LOWER(?) OR LOWER(username) = LOWER(?)',
            (email, email),
        ).fetchone()
        if existing:
            return jsonify({'error': 'Ya existe una cuenta registrada con ese correo.'}), 400

        try:
            # rol='especialista' esta fijo a proposito: el auto-registro
            # publico nunca debe poder crear administradores. Para eso
            # sigue existiendo el panel de administración (POST /api/usuarios).
            # activo=0: el auto-registro publico crea la cuenta pendiente de
            # aprobacion en vez de darle acceso inmediato a datos reales de
            # las IE. Un administrador la activa desde Configuracion > Usuarios
            # (el toggle Activar/Desactivar que ya existia para esto).
            user = create_app_user(
                conn,
                nombre=nombre,
                username=email,
                email=email,
                password=password,
                rol='especialista',
                activo=0,
            )
            conn.commit()
        except sqlite3.IntegrityError:
            conn.rollback()
            return jsonify({'error': 'Ya existe una cuenta registrada con ese correo.'}), 400
        except ValueError as exc:
            conn.rollback()
            return jsonify({'error': str(exc)}), 400

        return jsonify({
            'success': True,
            'pending_approval': True,
            'message': 'Cuenta creada. Un administrador debe activarla antes de que puedas ingresar.',
        })
    finally:
        conn.close()

@app.route('/api/logout', methods=['POST'])
@require_auth
def logout():
    auth_header = request.headers.get('Authorization', '')
    token = auth_header[7:].strip() if auth_header.lower().startswith('bearer ') else ''
    if token:
        conn = get_db()
        conn.execute('DELETE FROM app_session WHERE token = ?', (token,))
        conn.commit()
    return jsonify({'success': True})

# ─── INSTITUCIONES ────────────────────────────────────────────────────────────

@app.route('/api/instituciones', methods=['GET'])
@require_auth
def get_instituciones():
    conn = get_db()
    rows = conn.execute('''
        SELECT
            i.codigo_modular,
            i.nombre_iiee,
            i.nivel_modalidad,
            i.distrito,
            i.centro_poblado,
            i.latitud,
            i.longitud,
            i.altitud,
            i.tipo_gestion,
            COALESCE(i.alumnos_censo,  0) AS alumnos_censo,
            COALESCE(i.docentes_censo, 0) AS docentes_censo,
            COALESCE(r.priority_score,      0)              AS priority_score,
            COALESCE(r.total_alertas_campo, 0)              AS total_alertas,
            COALESCE(r.coverage_category,   'Sin evidencia') AS coverage_category,
            COALESCE(r.simon_fichas,        0)              AS simon_fichas,
            COALESCE(r.campo_documentos,    0)              AS campo_documentos,
            COALESCE(r.campo_owners,        '')             AS especialistas,
            COALESCE(r.simon_promedio_nivel,0)              AS simon_promedio,
            r.riesgo_simon_0_100, r.riesgo_infra_0_100, r.riesgo_campo_0_100,
            r.brecha_evidencia_0_100, r.score_priorizacion_ie, r.prioridad_v5,
            r.risk_score_infra, r.simon_promedio_general_desempeno,
            r.simon_porcentaje_items_bajos, r.campo_alertas_altas,
            r.campo_alertas_medias, r.campo_alertas_bajas,
            CASE
                WHEN COALESCE(r.priority_score,0) >= 15 THEN 'alta'
                WHEN COALESCE(r.priority_score,0) >= 8  THEN 'media'
                ELSE 'baja'
            END AS nivel_alerta
        FROM dim_institucion i
        LEFT JOIN institucion_resumen r ON i.codigo_modular = r.codigo_modular
        WHERE i.latitud  IS NOT NULL AND i.latitud  != 0
          AND i.longitud IS NOT NULL AND i.longitud != 0
          AND COALESCE(i.estado_fila, 1) = 1
        ORDER BY COALESCE(r.priority_score,0) DESC
    ''').fetchall()

    # Riesgo por fuente, calculado en vivo con la MISMA regla que el Panel
    # (build_simon_ranking / build_infra_ranking / build_campo_ranking) -- para
    # que el mapa nunca muestre un nivel distinto al que el especialista ya vio
    # en el ranking de esa fuente. None = la IE no tiene datos de esa fuente.
    simon_lookup = {ie['codigo_modular']: ie['riesgo_intervencion'] for ie in build_simon_ranking(conn)['ies']}
    infra_lookup = {ie['codigo_modular']: ie['nivel'] for ie in build_infra_ranking(conn)['ies']}
    campo_lookup = {ie['codigo_modular']: ie['nivel'] for ie in build_campo_ranking(conn)['ies']}

    # nivel_alerta (vista "Todas" del mapa) se recalcula aqui a partir de
    # app_alerta_priorizada -- la MISMA tabla que cuenta Gestion -> Alertas --
    # en vez del priority_score de institucion_resumen (que no incluye Campo
    # ni Dinamica y capea el score a 22, y por eso daba un numero de IEs en
    # alerta alta muy distinto al conteo real de alertas pendientes). Una IE
    # queda en 'alta' si tiene al menos una alerta pendiente de severidad
    # alta/critica, 'media' si tiene alguna de severidad media, y 'baja' si
    # no tiene ninguna alerta pendiente.
    alerta_rows = conn.execute('''
        SELECT codigo_modular,
               MAX(CASE WHEN severidad IN ('alta','critica') THEN 3
                        WHEN severidad = 'media' THEN 2
                        ELSE 1 END) AS sev_rank
        FROM app_alerta_priorizada
        WHERE estado = 'pendiente'
        GROUP BY codigo_modular
    ''').fetchall()
    sev_rank_to_nivel = {3: 'alta', 2: 'media', 1: 'baja'}
    alerta_lookup = {row['codigo_modular']: sev_rank_to_nivel.get(row['sev_rank'], 'baja') for row in alerta_rows}

    result = []
    for row in rows:
        item = dict(row)
        item['riesgo_simon'] = simon_lookup.get(item['codigo_modular'])
        item['riesgo_censo'] = infra_lookup.get(item['codigo_modular'])
        item['riesgo_campo'] = campo_lookup.get(item['codigo_modular'])
        item['nivel_alerta'] = alerta_lookup.get(item['codigo_modular'], 'baja')
        # Score de priorizacion del informe v5: el mapa puede colorear por esto
        # ademas de por alertas pendientes. Se envia el desglose completo para
        # que la ficha de la IE muestre de que esta hecho el puntaje.
        item['score'] = desglose_score(item)
        item['prioridad'] = item['score']['prioridad']
        result.append(item)

    conn.close()
    return jsonify(result)

# CRUD administrativo del listado de instituciones educativas (dim_institucion).
# Distinto del /api/instituciones de arriba, que es de solo lectura y esta
# pensado para el mapa (filtra las que no tienen coordenadas). Este es el
# mantenedor completo: alta, edicion y baja (estado_fila=0, nunca DELETE real)
# para el rol administrador.
DIM_INSTITUCION_EDITABLE_FIELDS = [
    'codigo_institucion', 'codigo_local', 'nombre_iiee', 'nivel_modalidad',
    'tipo_gestion', 'dependencia', 'direccion_iiee', 'departamento',
    'provincia', 'distrito', 'centro_poblado', 'latitud', 'longitud',
    'altitud', 'alumnos_censo', 'docentes_censo', 'secciones_censo',
]

@app.route('/api/admin/instituciones', methods=['GET', 'POST'])
@require_admin
def admin_instituciones():
    conn = get_db()
    try:
        if request.method == 'GET':
            incluir_inactivas = request.args.get('incluir_inactivas') == '1'
            where = '' if incluir_inactivas else 'WHERE COALESCE(estado_fila, 1) = 1'
            rows = conn.execute(f'''
                SELECT * FROM dim_institucion {where}
                ORDER BY COALESCE(estado_fila, 1) DESC, nombre_iiee
            ''').fetchall()
            return jsonify(rows_to_list(rows))

        data = request.get_json() or {}
        codigo = re.sub(r'\D+', '', str(data.get('codigo_modular') or '').strip())
        if not codigo:
            return jsonify({'error': 'El código modular es obligatorio (solo dígitos).'}), 400
        if not str(data.get('nombre_iiee') or '').strip():
            return jsonify({'error': 'El nombre de la IE es obligatorio.'}), 400

        existing = conn.execute('SELECT 1 FROM dim_institucion WHERE codigo_modular = ?', (codigo,)).fetchone()
        if existing:
            return jsonify({'error': f'Ya existe una IE con código modular {codigo}.'}), 400

        columns = ['codigo_modular'] + DIM_INSTITUCION_EDITABLE_FIELDS
        values = [codigo] + [data.get(f) for f in DIM_INSTITUCION_EDITABLE_FIELDS]
        placeholders = ', '.join('?' for _ in columns)
        conn.execute(
            f'INSERT INTO dim_institucion ({", ".join(columns)}, estado_fila) VALUES ({placeholders}, 1)',
            values
        )
        conn.commit()
        row = conn.execute('SELECT * FROM dim_institucion WHERE codigo_modular = ?', (codigo,)).fetchone()
        return jsonify({'success': True, 'institucion': dict(row)})
    except sqlite3.IntegrityError:
        conn.rollback()
        return jsonify({'error': 'No se pudo guardar la institución (código duplicado o dato inválido).'}), 400
    finally:
        conn.close()

@app.route('/api/admin/instituciones/<codigo_modular>', methods=['PUT', 'DELETE'])
@require_admin
def admin_institucion_detalle(codigo_modular):
    conn = get_db()
    try:
        row = conn.execute('SELECT * FROM dim_institucion WHERE codigo_modular = ?', (codigo_modular,)).fetchone()
        if not row:
            return jsonify({'error': 'Institución no encontrada.'}), 404

        if request.method == 'DELETE':
            conn.execute(
                "UPDATE dim_institucion SET estado_fila = 0, actualizado_en = CURRENT_TIMESTAMP WHERE codigo_modular = ?",
                (codigo_modular,)
            )
            conn.commit()
            return jsonify({'success': True})

        data = request.get_json() or {}
        if data.get('restaurar'):
            conn.execute(
                "UPDATE dim_institucion SET estado_fila = 1, actualizado_en = CURRENT_TIMESTAMP WHERE codigo_modular = ?",
                (codigo_modular,)
            )
            conn.commit()
            updated = conn.execute('SELECT * FROM dim_institucion WHERE codigo_modular = ?', (codigo_modular,)).fetchone()
            return jsonify({'success': True, 'institucion': dict(updated)})

        updates = []
        params = []
        for field in DIM_INSTITUCION_EDITABLE_FIELDS:
            if field in data:
                updates.append(f'{field} = ?')
                params.append(data.get(field))
        if not updates:
            return jsonify({'success': True, 'institucion': dict(row)})
        updates.append('actualizado_en = CURRENT_TIMESTAMP')
        params.append(codigo_modular)
        conn.execute(f'UPDATE dim_institucion SET {", ".join(updates)} WHERE codigo_modular = ?', params)
        conn.commit()
        updated = conn.execute('SELECT * FROM dim_institucion WHERE codigo_modular = ?', (codigo_modular,)).fetchone()
        return jsonify({'success': True, 'institucion': dict(updated)})
    except sqlite3.IntegrityError:
        conn.rollback()
        return jsonify({'error': 'No se pudo actualizar la institución.'}), 400
    finally:
        conn.close()

# ─── PROFESORES (core) ────────────────────────────────────────────────────────

PROFESOR_EDITABLE_FIELDS = [
    'dni', 'nombre_completo', 'celular', 'correo', 'cargo',
    'condicion_laboral', 'codigo_modular_ie',
]
PROFESOR_CARGOS = {'docente', 'director', 'otro'}


def normalize_profesor_payload(conn, data, requerir_nombre=True):
    """Valida y normaliza el payload de un profesor. Devuelve (valores, error)."""
    nombre = str(data.get('nombre_completo') or '').strip()
    if requerir_nombre and not nombre:
        return None, 'El nombre completo es obligatorio.'

    dni = re.sub(r'\D+', '', str(data.get('dni') or '')) or None
    if dni and len(dni) > 12:
        return None, 'El DNI no parece válido.'

    cargo = str(data.get('cargo') or 'docente').strip().lower()
    if cargo not in PROFESOR_CARGOS:
        cargo = 'docente'

    codigo_ie = str(data.get('codigo_modular_ie') or '').strip()
    if codigo_ie:
        # Nunca se guarda una IE que no exista en el padron: si no resuelve, se
        # rechaza en vez de dejar un vinculo roto (misma regla que informes de campo).
        resuelto = resolve_codigo_modular_strict(conn, codigo_ie)
        if not resuelto:
            return None, f'La IE "{codigo_ie}" no existe en el padrón.'
        codigo_ie = resuelto
    else:
        codigo_ie = None

    return {
        'dni': dni,
        'nombre_completo': nombre,
        'celular': str(data.get('celular') or '').strip() or None,
        'correo': str(data.get('correo') or '').strip() or None,
        'cargo': cargo,
        'condicion_laboral': str(data.get('condicion_laboral') or '').strip() or None,
        'codigo_modular_ie': codigo_ie,
    }, None


@app.route('/api/profesores', methods=['GET'])
@require_auth
def buscar_profesores():
    """Busqueda de profesores para los selectores del frontend.

    Prioriza coincidencia exacta de DNI y luego ordena por similitud de nombre
    (rapidfuzz). Solo sugiere: nunca decide por el usuario cual es el correcto.
    """
    q = (request.args.get('q') or '').strip()
    codigo_ie = (request.args.get('codigo_modular_ie') or '').strip()
    limit = min(max(safe_int(request.args.get('limit'), 20), 1), 100)

    conn = get_db()
    try:
        where = ['p.estado_fila = 1']
        params = []
        if codigo_ie:
            where.append('p.codigo_modular_ie = ?')
            params.append(codigo_ie)
        rows = conn.execute(f'''
            SELECT p.*, i.nombre_iiee AS ie_nombre, i.centro_poblado AS ie_centro_poblado
            FROM profesor p
            LEFT JOIN dim_institucion i ON i.codigo_modular = p.codigo_modular_ie
            WHERE {' AND '.join(where)}
            ORDER BY p.nombre_completo
        ''', params).fetchall()
    finally:
        conn.close()

    candidatos = rows_to_list(rows)
    if not q:
        return jsonify(candidatos[:limit])

    dni_q = re.sub(r'\D+', '', q)
    exactos = [c for c in candidatos if dni_q and c.get('dni') == dni_q]
    usados = {c['id'] for c in exactos}

    # Primero las coincidencias literales de texto (lo que la persona espera al
    # escribir un apellido), ordenadas por donde empieza el calce; despues las
    # difusas, que son las que rescatan los errores de tipeo. Sin este orden,
    # WRatio castiga que "Chalas" sea mucho mas corto que "CHALAS MONTENEGRO,
    # DORIS" y deja el resultado obvio fuera de las primeras posiciones.
    q_norm = _normalize_place_text(q)
    literales = []
    for c in candidatos:
        if c['id'] in usados:
            continue
        pos = _normalize_place_text(c.get('nombre_completo') or '').find(q_norm)
        if q_norm and pos >= 0:
            literales.append((pos, c))
    literales.sort(key=lambda par: (par[0], par[1].get('nombre_completo') or ''))
    literales = [c for _, c in literales]
    usados |= {c['id'] for c in literales}

    # partial_ratio compara la consulta contra el mejor tramo del nombre, que es
    # lo que hace falta aqui: la consulta suele ser un apellido suelto y mal
    # tipeado frente a un "APELLIDOS, Nombres" completo ("Davilla Atamain" ->
    # "Davila Atamain, Jarnelly Janeth" puntua 97, contra 61 de token_set_ratio).
    restantes = {c['id']: c['nombre_completo'] or '' for c in candidatos if c['id'] not in usados}
    por_id = {c['id']: c for c in candidatos}
    rankeados = rapidfuzz_process.extract(
        q, restantes, scorer=rapidfuzz_fuzz.partial_ratio, limit=limit
    )
    similares = [por_id[cid] for _, score, cid in rankeados if score >= 85]

    return jsonify((exactos + literales + similares)[:limit])


@app.route('/api/admin/profesores', methods=['GET', 'POST'])
@require_admin
def admin_profesores():
    conn = get_db()
    try:
        if request.method == 'GET':
            incluir_inactivos = request.args.get('incluir_inactivos') == '1'
            where = '' if incluir_inactivos else 'WHERE COALESCE(p.estado_fila, 1) = 1'
            rows = conn.execute(f'''
                SELECT p.*, i.nombre_iiee AS ie_nombre, i.centro_poblado AS ie_centro_poblado
                FROM profesor p
                LEFT JOIN dim_institucion i ON i.codigo_modular = p.codigo_modular_ie
                {where}
                ORDER BY COALESCE(p.estado_fila, 1) DESC, p.nombre_completo
            ''').fetchall()
            return jsonify(rows_to_list(rows))

        valores, error = normalize_profesor_payload(conn, request.get_json() or {})
        if error:
            return jsonify({'error': error}), 400

        columnas = list(valores.keys()) + ['fuente', 'creado_por']
        params = list(valores.values()) + ['manual', g.current_user['user_id']]
        placeholders = ', '.join('?' for _ in columnas)
        cur = conn.execute(
            f'INSERT INTO profesor ({", ".join(columnas)}) VALUES ({placeholders})', params
        )
        conn.commit()
        row = conn.execute('SELECT * FROM profesor WHERE id = ?', (cur.lastrowid,)).fetchone()
        return jsonify({'success': True, 'profesor': dict(row)})
    except sqlite3.IntegrityError:
        conn.rollback()
        return jsonify({'error': 'Ya existe un profesor con ese DNI.'}), 400
    finally:
        conn.close()


@app.route('/api/admin/profesores/<int:profesor_id>', methods=['PUT', 'DELETE'])
@require_admin
def admin_profesor_detalle(profesor_id):
    conn = get_db()
    try:
        row = conn.execute('SELECT * FROM profesor WHERE id = ?', (profesor_id,)).fetchone()
        if not row:
            return jsonify({'error': 'Profesor no encontrado.'}), 404

        if request.method == 'DELETE':
            conn.execute(
                'UPDATE profesor SET estado_fila = 0, modificado_por = ?, '
                'actualizado_en = CURRENT_TIMESTAMP WHERE id = ?',
                (g.current_user['user_id'], profesor_id),
            )
            conn.commit()
            return jsonify({'success': True})

        data = request.get_json() or {}
        if data.get('restaurar'):
            conn.execute(
                'UPDATE profesor SET estado_fila = 1, modificado_por = ?, '
                'actualizado_en = CURRENT_TIMESTAMP WHERE id = ?',
                (g.current_user['user_id'], profesor_id),
            )
            conn.commit()
            updated = conn.execute('SELECT * FROM profesor WHERE id = ?', (profesor_id,)).fetchone()
            return jsonify({'success': True, 'profesor': dict(updated)})

        merged = {**dict(row), **{k: v for k, v in data.items() if k in PROFESOR_EDITABLE_FIELDS}}
        valores, error = normalize_profesor_payload(conn, merged)
        if error:
            return jsonify({'error': error}), 400

        updates = [f'{campo} = ?' for campo in valores]
        params = list(valores.values())
        updates.append('modificado_por = ?')
        params.append(g.current_user['user_id'])
        updates.append('actualizado_en = CURRENT_TIMESTAMP')
        params.append(profesor_id)
        conn.execute(f'UPDATE profesor SET {", ".join(updates)} WHERE id = ?', params)
        conn.commit()
        updated = conn.execute('SELECT * FROM profesor WHERE id = ?', (profesor_id,)).fetchone()
        return jsonify({'success': True, 'profesor': dict(updated)})
    except sqlite3.IntegrityError:
        conn.rollback()
        return jsonify({'error': 'Ya existe un profesor con ese DNI.'}), 400
    finally:
        conn.close()

# ─── DASHBOARDS FILTRABLES ────────────────────────────────────────────────────
# Van aparte de /api/dashboard (que es pesado y no acepta filtros) para que las
# vistas de Capacitaciones y Especialistas puedan refiltrarse sin recargar todo.

def _filtros_dashboard():
    return (
        (request.args.get('anio') or '').strip(),
        (request.args.get('mes') or '').strip(),
        (request.args.get('codigo_modular') or '').strip(),
    )


@app.route('/api/dashboard/asistencias', methods=['GET'])
@require_auth
def dashboard_asistencias():
    anio, mes, codigo_ie = _filtros_dashboard()
    conn = get_db()
    try:
        return jsonify(build_asistencias_dashboard(conn, anio, mes, codigo_ie))
    finally:
        conn.close()


@app.route('/api/dashboard/especialistas', methods=['GET'])
@require_auth
def dashboard_especialistas():
    anio, mes, codigo_ie = _filtros_dashboard()
    conn = get_db()
    try:
        return jsonify(build_especialistas_monitoreo(conn, anio, mes, codigo_ie))
    finally:
        conn.close()


@app.route('/api/dashboard/periodos', methods=['GET'])
@require_auth
def dashboard_periodos():
    """Anios con datos reales, para armar el filtro sin inventar opciones."""
    conn = get_db()
    try:
        anios = set()
        for consulta in (
            "SELECT DISTINCT SUBSTR(fecha_ejecucion, 1, 4) FROM fichas_monitoreo WHERE COALESCE(fecha_ejecucion,'') != ''",
            "SELECT DISTINCT SUBSTR(fecha_visita, 1, 4) FROM informe_campo_visita WHERE COALESCE(fecha_visita,'') != ''",
            "SELECT DISTINCT SUBSTR(COALESCE(NULLIF(fecha_inicio,''), creado_en), 1, 4) FROM taller_capacitacion",
        ):
            for fila in conn.execute(consulta).fetchall():
                valor = str(fila[0] or '').strip()
                if len(valor) == 4 and valor.isdigit():
                    anios.add(valor)
        return jsonify({'anios': sorted(anios, reverse=True)})
    finally:
        conn.close()


# ─── CARGA MASIVA (Excel/CSV) ─────────────────────────────────────────────────
# Flujo de dos pasos: /preview parsea y clasifica sin escribir nada, y /commit
# recibe de vuelta las filas ya revisadas por una persona. Nunca se importa a
# ciegas: los duplicados y los errores se muestran antes de tocar la base.

CARGA_MASIVA_EXTENSIONES = {'.xlsx', '.xlsm', '.csv'}

# Cada campo: (clave interna, etiqueta visible, obligatorio, alias aceptados)
CARGA_MASIVA_ESQUEMAS = {
    'profesores': {
        'titulo': 'Profesores y directores',
        'campos': [
            ('nombre_completo', 'Nombre completo', True, ['nombre', 'apellidos y nombres', 'docente', 'director']),
            ('dni', 'DNI', False, ['documento', 'nro documento']),
            ('cargo', 'Cargo', False, ['tipo', 'condicion']),
            ('celular', 'Celular', False, ['telefono', 'numero de celular']),
            ('correo', 'Correo', False, ['email', 'correo electronico']),
            ('condicion_laboral', 'Condición laboral', False, ['condicion laboral', 'situacion laboral']),
            ('codigo_modular_ie', 'Código modular de la IE', False, ['codigo modular', 'ie', 'institucion educativa', 'codigo ie']),
        ],
    },
    'instituciones': {
        'titulo': 'Instituciones educativas',
        'campos': [
            ('codigo_modular', 'Código modular', True, ['codigo', 'cod modular']),
            ('nombre_iiee', 'Nombre de la IE', True, ['nombre', 'nombre ie', 'institucion educativa']),
            ('codigo_local', 'Código local', False, ['cod local', 'codlocal']),
            ('nivel_modalidad', 'Nivel / modalidad', False, ['nivel', 'modalidad']),
            ('tipo_gestion', 'Tipo de gestión', False, ['gestion']),
            ('distrito', 'Distrito', False, []),
            ('centro_poblado', 'Centro poblado', False, ['localidad', 'lugar']),
            ('direccion_iiee', 'Dirección', False, ['direccion']),
            ('latitud', 'Latitud', False, ['lat']),
            ('longitud', 'Longitud', False, ['lon', 'lng']),
        ],
    },
}


def _clave_encabezado(texto):
    return re.sub(r'[^a-z0-9 ]', '', _normalize_place_text(texto).lower()).strip()


def leer_filas_archivo(file_storage):
    """Devuelve (encabezados, filas) de un .xlsx/.csv subido. Sin tocar la BD."""
    nombre = file_storage.filename or ''
    extension = Path(nombre).suffix.lower()
    if extension == '.csv':
        contenido = file_storage.read()
        for encoding in ('utf-8-sig', 'latin-1'):
            try:
                texto = contenido.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError('No se pudo leer el CSV (codificación no reconocida).')
        muestra = texto[:4096]
        try:
            dialecto = csv.Sniffer().sniff(muestra, delimiters=',;\t|')
        except csv.Error:
            dialecto = csv.excel
        lector = csv.reader(io.StringIO(texto), dialecto)
        filas = [fila for fila in lector if any(str(c).strip() for c in fila)]
    else:
        import openpyxl
        libro = openpyxl.load_workbook(file_storage, data_only=True, read_only=True)
        hoja = libro.active
        filas = [
            list(fila) for fila in hoja.iter_rows(values_only=True)
            if fila and any(str(c).strip() for c in fila if c is not None)
        ]
    if not filas:
        raise ValueError('El archivo está vacío.')
    encabezados = [str(c or '').strip() for c in filas[0]]
    return encabezados, filas[1:]


def mapear_columnas(encabezados, esquema):
    """Empareja los encabezados del archivo con los campos del esquema."""
    normalizados = {_clave_encabezado(h): idx for idx, h in enumerate(encabezados) if str(h).strip()}
    mapa = {}
    faltantes = []
    for clave, etiqueta, obligatorio, alias in esquema['campos']:
        posibles = [clave.replace('_', ' '), etiqueta] + list(alias)
        indice = None
        for posible in posibles:
            indice = normalizados.get(_clave_encabezado(posible))
            if indice is not None:
                break
        if indice is None and obligatorio:
            faltantes.append(etiqueta)
        if indice is not None:
            mapa[clave] = indice
    return mapa, faltantes


@app.route('/api/admin/carga-masiva/esquemas', methods=['GET'])
@require_admin
def carga_masiva_esquemas():
    """Columnas esperadas por destino, para mostrarlas antes de subir nada."""
    return jsonify({
        destino: {
            'titulo': cfg['titulo'],
            'campos': [
                {'clave': c, 'etiqueta': e, 'obligatorio': o, 'alias': a}
                for c, e, o, a in cfg['campos']
            ],
        }
        for destino, cfg in CARGA_MASIVA_ESQUEMAS.items()
    })


@app.route('/api/admin/carga-masiva/preview', methods=['POST'])
@require_admin
def carga_masiva_preview():
    destino = (request.form.get('destino') or '').strip()
    if destino not in CARGA_MASIVA_ESQUEMAS:
        return jsonify({'error': 'Destino no válido.'}), 400
    archivo = request.files.get('file')
    if not archivo or not archivo.filename:
        return jsonify({'error': 'No se subió ningún archivo.'}), 400
    if Path(archivo.filename).suffix.lower() not in CARGA_MASIVA_EXTENSIONES:
        return jsonify({'error': 'Solo se aceptan archivos .xlsx o .csv.'}), 400

    esquema = CARGA_MASIVA_ESQUEMAS[destino]
    try:
        encabezados, filas_crudas = leer_filas_archivo(archivo)
    except Exception as exc:
        return jsonify({'error': f'No se pudo leer el archivo: {exc}'}), 400

    mapa, faltantes = mapear_columnas(encabezados, esquema)
    if faltantes:
        return jsonify({
            'error': f'Faltan columnas obligatorias: {", ".join(faltantes)}.',
            'encabezados_detectados': encabezados,
        }), 400

    conn = get_db()
    try:
        filas = []
        vistos_dni = {}
        vistos_codigo = {}
        for numero, cruda in enumerate(filas_crudas, start=2):
            valores = {}
            for clave, indice in mapa.items():
                bruto = cruda[indice] if indice < len(cruda) else None
                valores[clave] = str(bruto).strip() if bruto is not None else ''

            fila = {'fila': numero, 'valores': valores, 'estado': 'nuevo', 'motivo': ''}

            if destino == 'profesores':
                if not valores.get('nombre_completo'):
                    fila.update(estado='error', motivo='Falta el nombre completo.')
                else:
                    dni = re.sub(r'\D+', '', valores.get('dni', ''))
                    valores['dni'] = dni
                    codigo_ie = valores.get('codigo_modular_ie', '')
                    if codigo_ie:
                        resuelto = (resolve_codigo_modular_strict(conn, codigo_ie)
                                    or resolver_ie_por_nombre_y_nivel(conn, codigo_ie)
                                    or resolve_by_nombre_local(conn, codigo_ie))
                        if not resuelto:
                            fila.update(estado='error', motivo=f'La IE "{codigo_ie}" no existe en el padrón.')
                        valores['codigo_modular_ie'] = resuelto or ''
                    if fila['estado'] != 'error' and dni:
                        if dni in vistos_dni:
                            fila.update(estado='error', motivo=f'DNI repetido dentro del archivo (fila {vistos_dni[dni]}).')
                        else:
                            vistos_dni[dni] = numero
                            existente = conn.execute(
                                'SELECT id, nombre_completo FROM profesor WHERE dni = ?', (dni,)
                            ).fetchone()
                            if existente:
                                fila.update(estado='duplicado', motivo=f'Ya existe con ese DNI: {existente["nombre_completo"]}.',
                                            id_existente=existente['id'])
                    elif fila['estado'] != 'error':
                        existente = conn.execute(
                            'SELECT id FROM profesor WHERE dni IS NULL '
                            'AND UPPER(TRIM(nombre_completo)) = UPPER(TRIM(?)) '
                            'AND COALESCE(codigo_modular_ie, "") = ?',
                            (valores['nombre_completo'], valores.get('codigo_modular_ie', '')),
                        ).fetchone()
                        if existente:
                            fila.update(estado='duplicado', motivo='Ya existe una persona con ese nombre en esa IE.',
                                        id_existente=existente['id'])
            else:  # instituciones
                codigo = re.sub(r'\D+', '', valores.get('codigo_modular', ''))
                valores['codigo_modular'] = codigo
                if not codigo:
                    fila.update(estado='error', motivo='Falta el código modular (solo dígitos).')
                elif not valores.get('nombre_iiee'):
                    fila.update(estado='error', motivo='Falta el nombre de la IE.')
                elif codigo in vistos_codigo:
                    fila.update(estado='error', motivo=f'Código modular repetido dentro del archivo (fila {vistos_codigo[codigo]}).')
                else:
                    vistos_codigo[codigo] = numero
                    # Los tres choques que pidio el usuario, en orden de dureza.
                    choque = conn.execute(
                        'SELECT codigo_modular, nombre_iiee FROM dim_institucion WHERE codigo_modular = ?',
                        (codigo,),
                    ).fetchone()
                    motivo = f'Ya existe una IE con el código modular {codigo}.' if choque else ''
                    if not choque and valores.get('codigo_local'):
                        choque = conn.execute(
                            'SELECT codigo_modular, nombre_iiee FROM dim_institucion WHERE codigo_local = ? AND codigo_local != ""',
                            (valores['codigo_local'],),
                        ).fetchone()
                        if choque:
                            motivo = f'Otra IE ya usa el código local {valores["codigo_local"]}: {choque["nombre_iiee"]} ({choque["codigo_modular"]}).'
                    if not choque:
                        choque = conn.execute(
                            'SELECT codigo_modular, nombre_iiee FROM dim_institucion '
                            'WHERE UPPER(TRIM(nombre_iiee)) = UPPER(TRIM(?))',
                            (valores['nombre_iiee'],),
                        ).fetchone()
                        if choque:
                            motivo = f'Ya existe una IE con ese nombre: {choque["nombre_iiee"]} ({choque["codigo_modular"]}).'
                    if choque:
                        fila.update(estado='duplicado', motivo=motivo, id_existente=choque['codigo_modular'])

            filas.append(fila)

        resumen = {
            'nuevos': sum(1 for f in filas if f['estado'] == 'nuevo'),
            'duplicados': sum(1 for f in filas if f['estado'] == 'duplicado'),
            'errores': sum(1 for f in filas if f['estado'] == 'error'),
        }
        return jsonify({
            'success': True,
            'destino': destino,
            'encabezados_detectados': encabezados,
            'columnas_reconocidas': sorted(mapa.keys()),
            'filas': filas,
            'resumen': resumen,
        })
    finally:
        conn.close()


@app.route('/api/admin/carga-masiva/commit', methods=['POST'])
@require_admin
def carga_masiva_commit():
    data = request.get_json() or {}
    destino = (data.get('destino') or '').strip()
    if destino not in CARGA_MASIVA_ESQUEMAS:
        return jsonify({'error': 'Destino no válido.'}), 400
    filas = data.get('filas') or []
    if not filas:
        return jsonify({'error': 'No hay filas para importar.'}), 400
    politica = (data.get('politica_duplicados') or 'omitir').strip()

    conn = get_db()
    user_id = g.current_user['user_id']
    creados = actualizados = omitidos = 0
    try:
        for fila in filas:
            if not isinstance(fila, dict):
                continue
            estado = fila.get('estado')
            valores = fila.get('valores') or {}
            if estado == 'error':
                omitidos += 1
                continue
            if estado == 'duplicado' and politica != 'actualizar':
                omitidos += 1
                continue

            if destino == 'profesores':
                nombre = str(valores.get('nombre_completo') or '').strip()
                if not nombre:
                    omitidos += 1
                    continue
                dni = re.sub(r'\D+', '', str(valores.get('dni') or '')) or None
                cargo = str(valores.get('cargo') or 'docente').strip().lower()
                if cargo not in PROFESOR_CARGOS:
                    cargo = 'director' if 'direct' in cargo else 'docente'
                campos = (
                    dni, nombre,
                    str(valores.get('celular') or '').strip() or None,
                    str(valores.get('correo') or '').strip() or None,
                    cargo,
                    str(valores.get('condicion_laboral') or '').strip() or None,
                    str(valores.get('codigo_modular_ie') or '').strip() or None,
                )
                existente_id = fila.get('id_existente')
                if estado == 'duplicado' and existente_id:
                    conn.execute('''
                        UPDATE profesor SET dni = ?, nombre_completo = ?, celular = COALESCE(?, celular),
                            correo = COALESCE(?, correo), cargo = ?,
                            condicion_laboral = COALESCE(?, condicion_laboral),
                            codigo_modular_ie = COALESCE(?, codigo_modular_ie),
                            estado_fila = 1, modificado_por = ?, actualizado_en = CURRENT_TIMESTAMP
                        WHERE id = ?
                    ''', (*campos, user_id, existente_id))
                    actualizados += 1
                else:
                    conn.execute('''
                        INSERT INTO profesor (dni, nombre_completo, celular, correo, cargo,
                            condicion_laboral, codigo_modular_ie, fuente, creado_por)
                        VALUES (?, ?, ?, ?, ?, ?, ?, 'carga_masiva', ?)
                    ''', (*campos, user_id))
                    creados += 1
            else:
                codigo = re.sub(r'\D+', '', str(valores.get('codigo_modular') or ''))
                if not codigo:
                    omitidos += 1
                    continue
                columnas = ['codigo_local', 'nombre_iiee', 'nivel_modalidad', 'tipo_gestion',
                            'distrito', 'centro_poblado', 'direccion_iiee', 'latitud', 'longitud']
                datos = [str(valores.get(c) or '').strip() or None for c in columnas]
                if estado == 'duplicado':
                    asignaciones = ', '.join(f'{c} = COALESCE(?, {c})' for c in columnas)
                    conn.execute(
                        f'UPDATE dim_institucion SET {asignaciones}, estado_fila = 1, '
                        'actualizado_en = CURRENT_TIMESTAMP WHERE codigo_modular = ?',
                        (*datos, fila.get('id_existente') or codigo),
                    )
                    actualizados += 1
                else:
                    conn.execute(
                        f'INSERT INTO dim_institucion (codigo_modular, {", ".join(columnas)}, estado_fila) '
                        f'VALUES (?, {", ".join("?" for _ in columnas)}, 1)',
                        (codigo, *datos),
                    )
                    creados += 1

        conn.commit()
        return jsonify({'success': True, 'creados': creados, 'actualizados': actualizados, 'omitidos': omitidos})
    except sqlite3.IntegrityError as exc:
        conn.rollback()
        return jsonify({'error': f'No se pudo completar la importación: {exc}'}), 400
    finally:
        conn.close()


# ─── ALERTAS ──────────────────────────────────────────────────────────────────

@app.route('/api/alertas', methods=['GET'])
@require_auth
def get_alertas():
    cm   = request.args.get('codigo_modular')
    limit = min(max(safe_int(request.args.get('limit'), 60), 1), 1000)
    conn = get_db()
    if cm:
        rows = conn.execute('''
            SELECT a.*, c.nombre AS cat_nombre, c.tipo AS cat_tipo
            FROM app_alerta_priorizada a
            LEFT JOIN dim_alerta_categoria c ON a.alerta_codigo = c.alerta_codigo
            WHERE a.codigo_modular = ?
            ORDER BY a.score DESC
        ''', (cm,)).fetchall()
    else:
        rows = conn.execute('''
            SELECT a.*, c.nombre AS cat_nombre, c.tipo AS cat_tipo,
                   i.nombre_iiee, i.distrito
            FROM app_alerta_priorizada a
            LEFT JOIN dim_alerta_categoria c ON a.alerta_codigo = c.alerta_codigo
            LEFT JOIN dim_institucion      i ON a.codigo_modular = i.codigo_modular
            WHERE a.estado = 'pendiente'
            ORDER BY a.score DESC LIMIT ?
        ''', (limit,)).fetchall()
    conn.close()
    return jsonify(rows_to_list(rows))

# ─── DASHBOARD ────────────────────────────────────────────────────────────────

@app.route('/api/dashboard', methods=['GET'])
@require_auth
def get_dashboard():
    conn = get_db()

    total_ies        = conn.execute('SELECT COUNT(*) FROM dim_institucion').fetchone()[0]
    alertas_altas    = conn.execute("SELECT COUNT(DISTINCT codigo_modular) FROM app_alerta_priorizada WHERE severidad='alta' AND estado='pendiente'").fetchone()[0]
    alertas_pend     = conn.execute("SELECT COUNT(*) FROM app_alerta_priorizada WHERE estado='pendiente'").fetchone()[0]
    ies_monitoreadas = conn.execute("SELECT COUNT(*) FROM institucion_resumen WHERE coverage_category != 'Sin evidencia'").fetchone()[0]
    especialistas    = conn.execute('SELECT COUNT(*) FROM dim_especialista').fetchone()[0]

    top = conn.execute('''
        SELECT i.codigo_modular, i.nombre_iiee, i.nivel_modalidad, i.distrito,
               r.priority_score, r.coverage_category, r.total_alertas_campo,
               r.alertas_infraestructura_campo, r.alertas_pedagogicas_campo,
               r.campo_documentos, r.simon_fichas, r.simon_promedio_nivel,
               r.priority_reason, r.risk_score_infra, r.risk_flags
        FROM institucion_resumen r
        JOIN dim_institucion i ON r.codigo_modular = i.codigo_modular
        ORDER BY r.priority_score DESC LIMIT 20
    ''').fetchall()

    por_distrito = conn.execute('''
        SELECT i.distrito,
               COUNT(*) AS total,
               SUM(CASE WHEN COALESCE(r.priority_score,0) >= 15 THEN 1 ELSE 0 END) AS alerta_alta,
               SUM(CASE WHEN COALESCE(r.priority_score,0) >= 8  THEN 1 ELSE 0 END) AS alerta_media
        FROM dim_institucion i
        LEFT JOIN institucion_resumen r ON i.codigo_modular = r.codigo_modular
        WHERE i.distrito IS NOT NULL
        GROUP BY i.distrito ORDER BY total DESC
    ''').fetchall()

    risk_rows = conn.execute('''
        SELECT i.codigo_modular, i.nombre_iiee, i.nivel_modalidad, i.distrito,
               COALESCE(r.priority_score, 0) AS priority_score,
               COALESCE(r.coverage_category, 'Sin evidencia') AS coverage_category,
               COALESCE(r.total_alertas_campo, 0) AS total_alertas_campo,
               COALESCE(r.alertas_infraestructura_campo, 0) AS alertas_infraestructura_campo,
               COALESCE(r.alertas_pedagogicas_campo, 0) AS alertas_pedagogicas_campo,
               COALESCE(r.campo_documentos, 0) AS campo_documentos,
               COALESCE(r.simon_fichas, 0) AS simon_fichas,
               COALESCE(r.simon_promedio_nivel, 0) AS simon_promedio_nivel,
               COALESCE(r.priority_reason, '') AS priority_reason,
               COALESCE(r.risk_score_infra, 0) AS risk_score_infra,
               COALESCE(r.risk_flags, '') AS risk_flags
        FROM dim_institucion i
        LEFT JOIN institucion_resumen r ON i.codigo_modular = r.codigo_modular
    ''').fetchall()

    scored = []
    risk_distribution = {'Critico': 0, 'Alto': 0, 'Medio': 0, 'Bajo': 0}
    for row in risk_rows:
        item = dict(row)
        item.update(calculate_risk(item))
        risk_distribution[item['risk_level']] += 1
        scored.append(item)

    scored.sort(key=lambda r: r['risk_score'], reverse=True)
    top_riesgo = scored[:10]
    risk_focus_count = risk_distribution['Critico'] + risk_distribution['Alto']

    # Estadisticas de SIMON calculadas sobre la ULTIMA VISITA de cada docente
    # (ver simon_latest_ficha_por_docente): un docente que ya mejoro no debe
    # seguir penalizando el promedio general con una evaluacion vieja, y una
    # ficha cargada dos veces no debe contarse dos veces.
    simon_fichas_dedup, simon_dedup_meta = simon_latest_ficha_por_docente(conn)
    actual_fichas = len(simon_fichas_dedup)
    demo_mode = False
    fichas_recolectadas = actual_fichas
    ocr_total = sum(1 for f in simon_fichas_dedup if f.get('source') == 'ocr')
    promedios_fichas = [safe_float(f.get('promedio'), 0.0) for f in simon_fichas_dedup if safe_float(f.get('promedio'), 0.0) > 0]
    promedio_observado = round(sum(promedios_fichas) / len(promedios_fichas), 4) if promedios_fichas else 0.0
    compromisos = sum(1 for f in simon_fichas_dedup if str(f.get('compromisos_monitoreado') or f.get('compromisos') or '').strip())
    infra_stats = conn.execute('''
        SELECT
            COUNT(*) AS total_ie,
            SUM(CASE WHEN risk_score_infra >= 5 THEN 1 ELSE 0 END) AS criticas,
            SUM(CASE WHEN risk_score_infra >= 3 AND risk_score_infra < 5 THEN 1 ELSE 0 END) AS en_alerta,
            AVG(risk_score_infra) AS promedio_score,
            SUM(CASE WHEN edificaciones_riesgo > 0 THEN 1 ELSE 0 END) AS edificaciones_riesgo,
            SUM(CASE WHEN aulas_en_uso = 0 THEN 1 ELSE 0 END) AS sin_aulas_en_uso,
            SUM(alertas_infraestructura_campo) AS alertas_infraestructura
        FROM infraestructura_censo_2025
    ''').fetchone()
    infra_summary = {
        'total_ie': safe_int(infra_stats['total_ie'], 0) if infra_stats else 0,
        'criticas': safe_int(infra_stats['criticas'], 0) if infra_stats else 0,
        'en_alerta': safe_int(infra_stats['en_alerta'], 0) if infra_stats else 0,
        'promedio_score': round(safe_float(infra_stats['promedio_score'], 0.0), 2) if infra_stats else 0.0,
        'edificaciones_riesgo': safe_int(infra_stats['edificaciones_riesgo'], 0) if infra_stats else 0,
        'sin_aulas_en_uso': safe_int(infra_stats['sin_aulas_en_uso'], 0) if infra_stats else 0,
        'alertas_infraestructura': safe_int(infra_stats['alertas_infraestructura'], 0) if infra_stats else 0,
    }
    resultados = build_simon_indicator_results(conn, simon_fichas_dedup)
    docentes_refuerzo_total = sum(1 for f in simon_fichas_dedup if 0 < safe_float(f.get('promedio'), 0.0) < 3)
    item_critico = sorted(
        resultados,
        key=lambda item: (safe_int(item.get('bajo_nivel'), 0), -safe_float(item.get('promedio'), 0)),
        reverse=True,
    )[0] if resultados else {}

    acciones_recomendadas = [
        'Atender primero docentes con promedio SIMON menor a Nivel III.',
        'Priorizar IEs con risk_score_infra alto o riesgo estructural censal.',
        'Programar Visita 2 para calcular mejora real e IEAP.',
        'Revisar los item con mayor concentracion de Nivel I y II.',
    ]

    excel_dashboard = build_real_simon_dashboard(conn, scored)

    simon_risk_rows = conn.execute('''
        SELECT riesgo_intervencion_simon AS nivel, COUNT(*) AS total
        FROM institucion_resumen
        WHERE riesgo_intervencion_simon IS NOT NULL AND riesgo_intervencion_simon != ''
        GROUP BY riesgo_intervencion_simon
    ''').fetchall()
    simon_risk_total_ie = sum(safe_int(row['total'], 0) for row in simon_risk_rows)
    simon_risk_distribution = sorted(
        [
            {
                'nivel': row['nivel'],
                'ies': safe_int(row['total'], 0),
                'porcentaje': round(safe_int(row['total'], 0) / simon_risk_total_ie * 100, 1) if simon_risk_total_ie else 0,
            }
            for row in simon_risk_rows
        ],
        key=lambda item: SIMON_RISK_ORDER.get(item['nivel'], -1),
        reverse=True,
    )
    # Distribucion real de niveles de respuesta (I-IV) y de riesgo_intervencion
    # POR FICHA (no por IE) -- ambas se calculan en vivo desde fichas_monitoreo,
    # sin depender del rebuild de institucion_resumen (ver distribucion_ie arriba).
    nivel_counts = {1: 0, 2: 0, 3: 0, 4: 0}
    ficha_risk_counts = {'Alto': 0, 'Medio': 0, 'Bajo': 0}
    for ficha_row in simon_fichas_dedup:
        valores = simon_values_from_row(ficha_row)
        for valor in valores:
            if valor in nivel_counts:
                nivel_counts[valor] += 1
        ficha_riesgo = simon_documented_risk(valores)
        if ficha_riesgo:
            ficha_risk_counts[ficha_riesgo['riesgo_intervencion']] += 1

    nivel_total = sum(nivel_counts.values())
    distribucion_nivel = [
        {
            'nivel': f'Nivel {romano}',
            'valor': valor,
            'cantidad': nivel_counts[valor],
            'porcentaje': round(nivel_counts[valor] / nivel_total * 100, 1) if nivel_total else 0,
        }
        for valor, romano in ((1, 'I'), (2, 'II'), (3, 'III'), (4, 'IV'))
    ]

    ficha_risk_total = sum(ficha_risk_counts.values())
    distribucion_ficha = [
        {
            'nivel': nivel,
            'fichas': ficha_risk_counts[nivel],
            'porcentaje': round(ficha_risk_counts[nivel] / ficha_risk_total * 100, 1) if ficha_risk_total else 0,
        }
        for nivel in ('Alto', 'Medio', 'Bajo')
    ]

    # Promedio por dimension (AR01 preparacion / AR02 ensenanza), agregado a
    # partir de los promedios por indicador ya calculados en `resultados`.
    promedio_por_codigo = {item['area']: item['promedio'] for item in resultados}
    dimension_valores = {}
    for indicador in SIMON_INDICATOR_DEFINITIONS:
        valor = promedio_por_codigo.get(indicador['codigo_instrumento'])
        if valor is None:
            continue
        dimension_valores.setdefault(indicador['dimension'], []).append(valor)
    distribucion_dimension = [
        {'dimension': dimension, 'promedio': round(sum(valores) / len(valores), 2), 'indicadores': len(valores)}
        for dimension, valores in dimension_valores.items()
    ]

    simon_metodologia = {
        'indicadores': SIMON_INDICATOR_DEFINITIONS,
        'regla_riesgo': SIMON_RISK_RULE,
        'ejemplo': SIMON_RISK_EXAMPLE,
        'distribucion_ie': simon_risk_distribution,
        'distribucion_total_ie': simon_risk_total_ie,
        'distribucion_ficha': distribucion_ficha,
        'distribucion_total_ficha': ficha_risk_total,
        'distribucion_nivel': distribucion_nivel,
        'distribucion_dimension': distribucion_dimension,
    }

    # Ranking de priorizacion por IE -- siempre calculado sobre la ULTIMA
    # VISITA de cada docente (ver simon_latest_ficha_por_docente): no es un
    # filtro opcional, es la unica forma correcta de priorizar (una IE que
    # ya mejoro en su visita mas reciente no debe seguir arriba del ranking
    # por una evaluacion vieja y superada).
    simon_ranking = build_simon_ranking(conn, simon_fichas_dedup, simon_dedup_meta)
    infra_ranking = build_infra_ranking(conn)
    campo_ranking = build_campo_ranking(conn)
    priorizacion = build_priorizacion_resumen(conn)
    asistencias_talleres = build_asistencias_dashboard(conn)
    especialistas_monitoreo = build_especialistas_monitoreo(conn)
    ie_coverage = build_ie_coverage(conn)

    conn.close()
    return jsonify({
        'total_ies':        total_ies,
        'alertas_altas':    alertas_altas,
        'alertas_pendientes': alertas_pend,
        'ies_monitoreadas': ies_monitoreadas,
        'especialistas':    especialistas,
        'pct_monitoreadas': round(ies_monitoreadas / total_ies * 100, 1) if total_ies else 0,
        'top_prioridad':    rows_to_list(top),
        'por_distrito':     rows_to_list(por_distrito),
        'risk_distribution': risk_distribution,
        'risk_focus_count': risk_focus_count,
        'top_riesgo': top_riesgo,
        'risk_model': RISK_MODEL,
        'alert_rules': ALERT_RULES,
        'simon_metodologia': simon_metodologia,
        'simon_ranking': simon_ranking,
        'infra_ranking': infra_ranking,
        'campo_ranking': campo_ranking,
        'priorizacion': priorizacion,
        'asistencias_talleres': asistencias_talleres,
        'especialistas_monitoreo': especialistas_monitoreo,
        'ie_coverage': ie_coverage,
        'infra_definitions': INFRA_DEFINITIONS,
        'infraestructura_censo': infra_summary,
        'excel_dashboard': excel_dashboard,
        'demo_mode': demo_mode,
        'resultados_recolectados': {
            'fichas_recolectadas': fichas_recolectadas,
            'ocr_total': ocr_total,
            'promedio_observado': round(promedio_observado, 2),
            'compromisos_registrados': compromisos,
            'indicadores': resultados,
            'simon': {
                'docentes_refuerzo': docentes_refuerzo_total,
                'promedio_nivel': round(promedio_observado, 2),
                'nivel_actual': simon_nivel_romano(promedio_observado),
                'item_critico_codigo': item_critico.get('area', ''),
                'item_critico_brecha': item_critico.get('bajo_nivel', 0),
                'item_critico_total': item_critico.get('total', 0),
            },
            'infraestructura': infra_summary,
            'acciones_recomendadas': acciones_recomendadas,
        },
    })

import pdfplumber
import google.generativeai as genai
from azure.core.credentials import AzureKeyCredential
from azure.ai.documentintelligence import DocumentIntelligenceClient

AZURE_ENDPOINT, AZURE_KEY = get_azure_config()
GEMINI_KEY = get_gemini_key()

if GEMINI_KEY:
    genai.configure(api_key=GEMINI_KEY)

def extract_with_python(file_bytes):
    text = ""
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
    except Exception as e:
        print(f"Error pdfplumber: {e}")
    return text.strip()

def extract_with_azure(file_bytes, content_type='application/octet-stream'):
    endpoint, key = get_azure_config()
    if not endpoint or not key:
        print("Azure OCR no configurado: define AZURE_DOCUMENT_INTELLIGENCE_ENDPOINT y AZURE_DOCUMENT_INTELLIGENCE_KEY.")
        return ""
    try:
        client = DocumentIntelligenceClient(endpoint=endpoint, credential=AzureKeyCredential(key))
        poller = client.begin_analyze_document(
            "prebuilt-read",
            file_bytes,
            content_type=content_type,
        )
        result = poller.result()
        return result.content or ""
    except Exception as e:
        print(f"Error Azure: {e}")
        return ""

def extract_with_windows_ocr(file_bytes, filename='documento.jpeg'):
    suffix = Path(filename or 'documento.jpeg').suffix or '.jpeg'
    temp_path = None
    script_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp.write(file_bytes)
            temp_path = Path(tmp.name)

        ps_path = str(temp_path).replace("'", "''")
        script = f"""
$imagePath = '{ps_path}'
Add-Type -AssemblyName System.Runtime.WindowsRuntime
$null = [Windows.Storage.StorageFile, Windows.Storage, ContentType=WindowsRuntime]
$null = [Windows.Storage.FileAccessMode, Windows.Storage, ContentType=WindowsRuntime]
$null = [Windows.Graphics.Imaging.BitmapDecoder, Windows.Graphics.Imaging, ContentType=WindowsRuntime]
$null = [Windows.Graphics.Imaging.SoftwareBitmap, Windows.Graphics.Imaging, ContentType=WindowsRuntime]
$null = [Windows.Media.Ocr.OcrEngine, Windows.Foundation, ContentType=WindowsRuntime]
$null = [Windows.Globalization.Language, Windows.Globalization, ContentType=WindowsRuntime]
function AwaitOperation($Operation, $ResultType) {{
  $asTask = ([System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object {{ $_.Name -eq 'AsTask' -and $_.IsGenericMethod -and $_.GetParameters().Count -eq 1 }})[0]
  $task = $asTask.MakeGenericMethod($ResultType).Invoke($null, @($Operation))
  $task.Wait() | Out-Null
  $task.Result
}}
$file = AwaitOperation ([Windows.Storage.StorageFile]::GetFileFromPathAsync($imagePath)) ([Windows.Storage.StorageFile])
$stream = AwaitOperation ($file.OpenAsync([Windows.Storage.FileAccessMode]::Read)) ([Windows.Storage.Streams.IRandomAccessStream])
$decoder = AwaitOperation ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($stream)) ([Windows.Graphics.Imaging.BitmapDecoder])
$bitmap = AwaitOperation ($decoder.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap])
$engine = $null
foreach ($tag in @('es-ES','es-MX')) {{
  $lang = [Windows.Globalization.Language]::new($tag)
  $engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromLanguage($lang)
  if ($null -ne $engine) {{ break }}
}}
if ($null -eq $engine) {{ $engine = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages() }}
if ($null -eq $engine) {{ throw 'Windows OCR no tiene idiomas disponibles.' }}
$result = AwaitOperation ($engine.RecognizeAsync($bitmap)) ([Windows.Media.Ocr.OcrResult])
$result.Text
"""
        with tempfile.NamedTemporaryFile('w', delete=False, suffix='.ps1', encoding='utf-8-sig') as ps_file:
            ps_file.write(script)
            script_path = Path(ps_file.name)

        result = subprocess.run(
            ['powershell', '-NoProfile', '-ExecutionPolicy', 'Bypass', '-File', str(script_path)],
            text=True,
            capture_output=True,
            timeout=45,
            encoding='utf-8',
            errors='replace',
        )
        if result.returncode != 0:
            print(f"Error Windows OCR: {result.stderr.strip()}")
            return ""
        return re.sub(r'\s+', ' ', result.stdout).strip()
    except Exception as e:
        print(f"Error Windows OCR: {e}")
        return ""
    finally:
        if temp_path:
            try:
                temp_path.unlink(missing_ok=True)
            except Exception:
                pass
        if script_path:
            try:
                script_path.unlink(missing_ok=True)
            except Exception:
                pass

def is_pdf_file(filename, mimetype=''):
    lower = (filename or '').lower()
    return lower.endswith('.pdf') or mimetype == 'application/pdf'

def extract_text_from_upload(file_storage):
    file_bytes = file_storage.read()
    filename = file_storage.filename or 'documento'
    mimetype = file_storage.mimetype or 'application/octet-stream'

    if is_pdf_file(filename, mimetype):
        text = extract_with_python(file_bytes)
        if len(text) >= 100:
            return text, 'pdfplumber', filename, file_bytes, mimetype
        print("Usando Azure OCR como fallback para PDF...")
        text = extract_with_azure(file_bytes, mimetype)
        return text, 'azure', filename, file_bytes, mimetype

    text = extract_with_azure(file_bytes, mimetype)
    if len(text) >= 50:
        return text, 'azure', filename, file_bytes, mimetype

    print("Usando Windows OCR como fallback para imagen...")
    local_text = extract_with_windows_ocr(file_bytes, filename)
    if local_text:
        method = 'windows_ocr' if not text else 'azure+windows_ocr'
        return local_text, method, filename, file_bytes, mimetype

    return text, 'azure', filename, file_bytes, mimetype

FORM_LABELS = {
    'REGIÓN', 'REGION', 'UGEL', 'I.E.', 'IE', 'NIVEL/MODALIDAD', 'DIRECTOR(A)',
    'CEL.', 'EMAIL', 'SITUACIÓN LABORAL', 'SITUACION LABORAL', 'DESIGNADO(A)',
    'ENCARGADO(A)', 'NOMBRE', 'DNI', 'GRADO', 'SECCIÓN', 'SECCION',
    'Nº ESTUDIANTES', 'N° ESTUDIANTES', 'AREA', 'ÁREA', 'COMPETENCIA',
    'TÍTULO DE LA SESIÓN', 'TITULO DE LA SESION', 'MONITOR', 'IGED',
}

def compact_lines(text):
    return [line.strip() for line in text.splitlines() if line.strip()]

def is_form_label(line):
    normalized = re.sub(r'\s+', ' ', line.upper().strip(' :'))
    return normalized in FORM_LABELS

def next_line_after(lines, *labels):
    normalized = [line.upper() for line in lines]
    for label in labels:
        label = label.upper()
        for i, line in enumerate(normalized):
            if line == label or line.startswith(label):
                for j in range(i + 1, min(i + 5, len(lines))):
                    candidate = lines[j].strip()
                    if is_form_label(candidate):
                        return ''
                    if candidate and candidate.upper() not in labels:
                        return candidate
    return ''

def section_text(text, start, *ends):
    upper = text.upper()
    start_idx = upper.find(start.upper())
    if start_idx < 0:
        return ''
    end_idx = len(text)
    for end in ends:
        idx = upper.find(end.upper(), start_idx + len(start))
        if idx >= 0:
            end_idx = min(end_idx, idx)
    return text[start_idx:end_idx]

def text_between(text, start, *ends):
    upper = text.upper()
    start_idx = upper.find(start.upper())
    if start_idx < 0:
        return ''
    start_idx += len(start)
    end_idx = len(text)
    for end in ends:
        idx = upper.find(end.upper(), start_idx)
        if idx >= 0:
            end_idx = min(end_idx, idx)
    return re.sub(r'\s+', ' ', text[start_idx:end_idx]).strip(' :.-')

def text_between_raw(text, start, *ends):
    upper = text.upper()
    start_idx = upper.find(start.upper())
    if start_idx < 0:
        return ''
    start_idx += len(start)
    end_idx = len(text)
    for end in ends:
        idx = upper.find(end.upper(), start_idx)
        if idx >= 0:
            end_idx = min(end_idx, idx)
    return text[start_idx:end_idx].strip()

def text_after_all_headings(text, heading, *ends):
    upper = text.upper()
    heading_upper = heading.upper()
    parts = []
    start = 0
    while True:
        idx = upper.find(heading_upper, start)
        if idx < 0:
            break
        chunk_start = idx + len(heading)
        chunk_end = len(text)
        for end in (heading, *ends):
            end_idx = upper.find(end.upper(), chunk_start)
            if end_idx >= 0:
                chunk_end = min(chunk_end, end_idx)
        chunk = re.sub(r'\s+', ' ', text[chunk_start:chunk_end]).strip(' :.-')
        if chunk:
            parts.append(chunk)
        start = chunk_end
    return '\n'.join(parts)

def infer_selected_level(block):
    for line in compact_lines(block):
        if 'NIVEL' not in line.upper():
            continue
        level = normalize_level(line)
        if level and (
            re.search(r'[Xx☒✓✔&\\_]', line)
            or not re.match(r'^\s*[a-dA-D]\.?\s*', line)
        ):
            return level

    selected_patterns = [
        r'(?:X|☒|✓|✔|:SELECTED:|SELECTED)\s*(?:[A-D]\.?\s*)?NIVEL\s*(IV|III|II|I|[1-4])',
        r'(?:[A-D]\.?\s*)?NIVEL\s*(IV|III|II|I|[1-4])\s*(?:X|☒|✓|✔|:SELECTED:|SELECTED)',
    ]
    for pattern in selected_patterns:
        match = re.search(pattern, block, re.IGNORECASE)
        if match:
            return normalize_level(match.group(1))
    return 0

def parse_ficha_from_text(text):
    lines = compact_lines(text)
    ie_section = section_text(
        text,
        'Datos de identificación la IE',
        'Datos de identificación del docente monitoreado',
    )
    docente_section = section_text(
        text,
        'Datos de identificación del docente monitoreado',
        'Datos de identificación del especialista responsable',
    )
    monitor_section = section_text(
        text,
        'Datos de identificación del especialista responsable',
        'A: Preparación',
    )
    ie_lines = compact_lines(ie_section)
    docente_lines = compact_lines(docente_section)
    monitor_lines = compact_lines(monitor_section)

    area_match = re.search(r'(?:ÁREA|AREA)\s+([^\n]+)', docente_section, re.IGNORECASE)
    students_match = re.search(r'(?:N[°º]\s*)?ESTUDIANTES\s*(\d+)', docente_section, re.IGNORECASE)

    parsed = {
        'region': next_line_after(ie_lines, 'REGIÓN', 'REGION') or 'AMAZONAS',
        'ugel': next_line_after(ie_lines, 'UGEL'),
        'nombre_ie': next_line_after(ie_lines, 'I.E.', 'IE'),
        'nivel_modalidad': next_line_after(ie_lines, 'NIVEL/MODALIDAD'),
        'director': next_line_after(ie_lines, 'DIRECTOR(A)'),
        'director_cel': next_line_after(ie_lines, 'CEL.'),
        'docente': next_line_after(docente_lines, 'NOMBRE'),
        'docente_dni': next_line_after(docente_lines, 'DNI'),
        'docente_cel': next_line_after(docente_lines, 'CEL.'),
        'grado': next_line_after(docente_lines, 'GRADO'),
        'seccion': next_line_after(docente_lines, 'SECCIÓN', 'SECCION'),
        'area': area_match.group(1).strip() if area_match else next_line_after(docente_lines, 'ÁREA', 'AREA'),
        'competencia': next_line_after(docente_lines, 'COMPETENCIA'),
        'titulo_sesion': next_line_after(docente_lines, 'TÍTULO DE LA SESIÓN', 'TITULO DE LA SESION'),
        'monitor': next_line_after(monitor_lines, 'MONITOR'),
        'monitor_dni': next_line_after(monitor_lines, 'DNI'),
        'iged': next_line_after(monitor_lines, 'IGED'),
        'observaciones_recomendaciones': text_between(
            text,
            'OBSERVACIONES/ RECOMENDACIONES',
            'COMPROMISOS DEL MONITOREADO',
        ),
        'compromisos_monitoreado': text_after_all_headings(
            text,
            'COMPROMISOS DEL MONITOREADO',
            'Especialista monitoreado',
            'Monitoreado',
        ),
    }

    date_match = re.search(r'\b(\d{1,2})[/-](\d{1,2})[/-](\d{4})\b', text)
    if date_match:
        day, month, year = date_match.groups()
        parsed['fecha_ejecucion'] = f'{year}-{int(month):02d}-{int(day):02d}'

    visit_match = re.search(r'N[°º]?\s*VISITA\s*(\d+)', text, re.IGNORECASE)
    if visit_match:
        parsed['n_visita'] = visit_match.group(1)

    if students_match:
        parsed['nro_estudiantes'] = int(students_match.group(1))

    item_markers = {
        'a1': 'En la planificación curricular del docente se evidencia',
        'a2': 'En la planificación curricular del docente se observa',
        'a3': 'En la programación curricular del docente se observa',
        'b1': 'El docente promueve el interés',
        'b2': 'El docente propone actividades',
        'b3': 'El docente acompaña el proceso',
        'b4': 'El docente se comunica de manera respetuosa',
        'b5': 'El docente establece en su práctica pedagógica',
    }
    marker_values = list(item_markers.items())
    for idx, (key, marker) in enumerate(marker_values):
        end_markers = [m for _, m in marker_values[idx + 1:]] + ['OBSERVACIONES/ RECOMENDACIONES']
        block = text_between_raw(text, marker, *end_markers)
        parsed[key] = infer_selected_level(block)

    parsed['observaciones'] = parsed['observaciones_recomendaciones']
    parsed['compromisos'] = parsed['compromisos_monitoreado']
    return parsed

def process_with_gemini(text):
    if not get_gemini_key():
        print("Gemini no configurado: define GEMINI_API_KEY.")
        return None

    prompt = """
    Eres un asistente experto en extraer datos de fichas de monitoreo docente en Perú.
    Extrae todos los datos disponibles y devuélvelos estrictamente como un objeto JSON válido,
    sin markdown ni texto adicional.

    Usa exactamente estas claves:
    {
      "region": "string",
      "ugel": "string",
      "n_visita": "string",
      "fecha_ejecucion": "YYYY-MM-DD",
      "codigo_modular": "string",
      "nombre_ie": "string",
      "nivel_modalidad": "string",
      "director": "string",
      "director_cel": "string",
      "director_email": "string",
      "director_situacion_laboral": "string",
      "docente": "string",
      "docente_dni": "string",
      "docente_cel": "string",
      "docente_email": "string",
      "docente_situacion_laboral": "string",
      "grado": "string",
      "seccion": "string",
      "nro_estudiantes": "entero",
      "area": "string",
      "competencia": "string",
      "titulo_sesion": "string",
      "monitor": "string",
      "monitor_dni": "string",
      "iged": "string",
      "monitor_email": "string",
      "a1": "entero (1 al 4)",
      "a2": "entero (1 al 4)",
      "a3": "entero (1 al 4)",
      "b1": "entero (1 al 4)",
      "b2": "entero (1 al 4)",
      "b3": "entero (1 al 4)",
      "b4": "entero (1 al 4)",
      "b5": "entero (1 al 4)",
      "a1_observacion": "string",
      "a2_observacion": "string",
      "a3_observacion": "string",
      "b1_observacion": "string",
      "b2_observacion": "string",
      "b3_observacion": "string",
      "b4_observacion": "string",
      "b5_observacion": "string",
      "observaciones_recomendaciones": "string",
      "compromisos_monitoreado": "string"
    }

    Reglas:
    - Si no encuentras un dato, usa "" para textos y 0 para números.
    - Para los ítems A1-A3 y B1-B5, convierte NIVEL I, II, III o IV a 1, 2, 3 o 4.
    - Extrae las observaciones de cada ítem en su clave *_observacion.
    - El bloque "OBSERVACIONES/ RECOMENDACIONES" debe ir completo en observaciones_recomendaciones.
    - El bloque "COMPROMISOS DEL MONITOREADO" debe ir completo en compromisos_monitoreado.
    - No inventes valores. Si el documento es una plantilla vacía, deja esos campos vacíos o en 0.

    TEXTO A ANALIZAR:
    """ + text

    try:
        model = genai.GenerativeModel(GEMINI_DEFAULT_MODEL)
        response = model.generate_content(prompt)
    except Exception as e:
        print(f"Error Gemini request: {e}")
        return None

    raw = response.text.replace('```json', '').replace('```', '').strip()
    try:
        data = json.loads(raw)
        return data
    except Exception as e:
        print(f"Error Gemini JSON: {e}, Raw: {raw}")
        return None

def process_images_with_gemini(files_payload):
    if not get_gemini_key():
        print("Gemini Vision no configurado: define GEMINI_API_KEY.")
        return None

    image_parts = [
        {
            'mime_type': item.get('mimetype') or 'image/jpeg',
            'data': item.get('bytes') or b'',
        }
        for item in files_payload
        if not is_pdf_file(item.get('filename', ''), item.get('mimetype', ''))
    ]
    image_parts = [part for part in image_parts if part['data']]
    if not image_parts:
        return None

    prompt = """
    Eres un asistente experto en leer fotos de fichas de monitoreo docente de Peru.
    Analiza todas las imagenes como paginas de una misma ficha y devuelve solo JSON valido.
    Lee texto impreso y manuscrito. Extrae marcas X o check en NIVEL I, II, III o IV.

    Usa exactamente estas claves:
    {
      "region": "string",
      "ugel": "string",
      "n_visita": "string",
      "fecha_ejecucion": "YYYY-MM-DD",
      "codigo_modular": "string",
      "nombre_ie": "string",
      "nivel_modalidad": "string",
      "director": "string",
      "director_cel": "string",
      "director_email": "string",
      "director_situacion_laboral": "string",
      "docente": "string",
      "docente_dni": "string",
      "docente_cel": "string",
      "docente_email": "string",
      "docente_situacion_laboral": "string",
      "grado": "string",
      "seccion": "string",
      "nro_estudiantes": "entero",
      "area": "string",
      "competencia": "string",
      "titulo_sesion": "string",
      "monitor": "string",
      "monitor_dni": "string",
      "iged": "string",
      "monitor_email": "string",
      "a1": "entero (1 al 4)",
      "a2": "entero (1 al 4)",
      "a3": "entero (1 al 4)",
      "b1": "entero (1 al 4)",
      "b2": "entero (1 al 4)",
      "b3": "entero (1 al 4)",
      "b4": "entero (1 al 4)",
      "b5": "entero (1 al 4)",
      "a1_observacion": "string",
      "a2_observacion": "string",
      "a3_observacion": "string",
      "b1_observacion": "string",
      "b2_observacion": "string",
      "b3_observacion": "string",
      "b4_observacion": "string",
      "b5_observacion": "string",
      "observaciones_recomendaciones": "string",
      "compromisos_monitoreado": "string"
    }

    Reglas:
    - Si no encuentras un dato, usa "" para textos y 0 para numeros.
    - Convierte NIVEL I, II, III, IV a 1, 2, 3, 4.
    - No inventes datos fuera de lo visible.
    - Si hay dos bloques de compromisos, une ambos con salto de linea.
    - Responde estrictamente JSON, sin markdown.
    """

    try:
        model = genai.GenerativeModel(GEMINI_DEFAULT_MODEL)
        response = model.generate_content([prompt, *image_parts])
    except Exception as e:
        print(f"Error Gemini Vision request: {e}")
        return None

    raw = response.text.replace('```json', '').replace('```', '').strip()
    try:
        return json.loads(raw)
    except Exception as e:
        print(f"Error Gemini Vision JSON: {e}, Raw: {raw}")
        return None

# ─── ASISTENCIA A TALLERES: extraccion de listas de participantes ──────────
# La foto es una tabla manuscrita multi-fila, no una ficha de un solo sujeto:
# por eso usa vision (como el OCR de fichas) pero con las reglas estrictas y el
# modo JSON del extractor de informes de campo, que ya esta afinado para no
# inventar filas ni completar celdas ilegibles.

TALLER_SYSTEM_PROMPT = """Eres un asistente que digitaliza listas de asistencia a
talleres y capacitaciones de la UGEL IBIR Imaza a partir de fotos de planillas
llenadas a mano.
Reglas:
1. No inventes datos. Si una celda esta vacia o es ilegible, usa "".
2. Copia la columna de institucion educativa EXACTAMENTE como esta escrita a
   mano, con sus abreviaturas y puntos (ej. "A.A.H", "AAH"). No la expandas ni
   la corrijas: alguien la confirmara despues.
3. Cada fila de la tabla de participantes es un asistente distinto. No agrupes
   ni omitas filas, aunque esten incompletas.
4. Los DNI peruanos tienen 8 digitos; si lees mas o menos, transcribe lo que ves.
5. No interpretes la firma: no es un dato a extraer.
6. Responde exclusivamente JSON valido segun el esquema entregado, sin markdown.
"""

TALLER_JSON_SCHEMA = {
    'documento': {
        'dre': 'string',
        'ugel': 'string',
        'sede': 'string',
        'nombre_evento': 'string (titulo de la lista, ej. LISTA DE PARTICIPANTES EN LA ASISTENCIA TECNICA)',
        'compromiso_indicador': 'string (ej. Compromiso de Desempeno 2026/Indicador 7.3)',
        'nivel_educativo': 'string',
        'nombre_ie_sede': 'string (IE anfitriona del taller)',
        'fecha_inicio': 'YYYY-MM-DD o ""',
        'fecha_fin': 'YYYY-MM-DD o "" si es un solo dia',
    },
    'asistentes': [
        {
            'fila_num': 'entero (numero de fila en la planilla)',
            'nombre_completo': 'string (apellidos y nombres tal como estan escritos)',
            'dni': 'string (solo digitos)',
            'institucion_educativa_texto': 'string (tal como esta escrito, sin corregir)',
            'grado_a_cargo': 'string',
            'area': 'string',
            'celular': 'string',
            'correo': 'string',
        }
    ],
}


def call_taller_gemini_vision(image_parts):
    if not get_gemini_key():
        return {'status': 'missing_env', 'provider': 'gemini_vision', 'payload': {}, 'error': ''}
    prompt = (
        TALLER_SYSTEM_PROMPT
        + '\n\nEsquema JSON de respuesta:\n'
        + json.dumps(TALLER_JSON_SCHEMA, ensure_ascii=False, indent=2)
    )
    try:
        model = genai.GenerativeModel(GEMINI_DEFAULT_MODEL)
        response = model.generate_content(
            [prompt, *image_parts],
            generation_config={'response_mime_type': 'application/json', 'temperature': 0},
        )
        return {'status': 'ok', 'provider': 'gemini_vision', 'payload': parse_json_loose(response.text), 'error': ''}
    except Exception as exc:
        return {'status': 'failed', 'provider': 'gemini_vision', 'payload': {}, 'error': f'{type(exc).__name__}: {exc}'}


def call_taller_gemini_texto(text):
    """Respaldo cuando no hay imagen utilizable (ej. subieron un PDF ya digital)."""
    if not get_gemini_key():
        return {'status': 'missing_env', 'provider': 'gemini_texto', 'payload': {}, 'error': ''}
    prompt = (
        TALLER_SYSTEM_PROMPT
        + '\n\nEsquema JSON de respuesta:\n'
        + json.dumps(TALLER_JSON_SCHEMA, ensure_ascii=False, indent=2)
        + '\n\nTexto de la planilla:\n'
        + truncate_for_llm(text, 18000)
    )
    try:
        model = genai.GenerativeModel(GEMINI_DEFAULT_MODEL)
        response = model.generate_content(
            prompt,
            generation_config={'response_mime_type': 'application/json', 'temperature': 0},
        )
        return {'status': 'ok', 'provider': 'gemini_texto', 'payload': parse_json_loose(response.text), 'error': ''}
    except Exception as exc:
        return {'status': 'failed', 'provider': 'gemini_texto', 'payload': {}, 'error': f'{type(exc).__name__}: {exc}'}


def extract_taller_asistentes(files_payload, text):
    """Devuelve (payload, proveedor, error). No toca la base de datos."""
    image_parts = []
    for item in files_payload:
        if is_pdf_file(item.get('filename', ''), item.get('mimetype', '')):
            continue
        if item.get('bytes'):
            image_parts.append({
                'mime_type': item.get('mimetype') or 'image/jpeg',
                'data': item['bytes'],
            })

    if image_parts:
        resultado = call_taller_gemini_vision(image_parts)
        if resultado['status'] == 'ok' and resultado['payload'].get('asistentes'):
            return resultado['payload'], resultado['provider'], ''

    if len((text or '').strip()) >= 50:
        resultado = call_taller_gemini_texto(text)
        if resultado['status'] == 'ok' and resultado['payload'].get('asistentes'):
            return resultado['payload'], resultado['provider'], ''
        return {}, resultado['provider'], resultado.get('error') or 'sin_asistentes_detectados'

    return {}, 'ninguno', 'sin_imagen_ni_texto_utilizable'


def resolver_ie_por_nombre_y_nivel(conn, nombre_ie, nivel_educativo=''):
    """Resuelve una IE escrita en lenguaje natural ("IE. Alberto Acosta Herrera").

    Quita los prefijos habituales ("IE.", "I.E.", "N°") que impiden el match
    exacto, y cuando el mismo nombre existe en varios niveles (el caso normal:
    un colegio con Inicial, Primaria y Secundaria) desempata con el nivel que
    declara el propio documento. Si sigue habiendo mas de una candidata,
    devuelve '' para que lo confirme una persona.
    """
    nombre = str(nombre_ie or '').strip()
    if not nombre:
        return ''
    nombre = re.sub(r'^\s*(i\.?\s*e\.?|institucion\s+educativa|colegio)\s*[:.\-]?\s*',
                    '', nombre, flags=re.IGNORECASE).strip()
    nombre = re.sub(r'^\s*n[°ºo]\s*', '', nombre, flags=re.IGNORECASE).strip()
    if not nombre:
        return ''

    candidatas = conn.execute(
        'SELECT codigo_modular, nivel_modalidad FROM dim_institucion '
        'WHERE UPPER(TRIM(nombre_iiee)) = UPPER(TRIM(?)) AND COALESCE(estado_fila, 1) = 1',
        (nombre,),
    ).fetchall()
    if not candidatas:
        return ''
    if len(candidatas) == 1:
        return candidatas[0]['codigo_modular']

    nivel = str(nivel_educativo or '').strip().lower()
    if nivel:
        # "Secundaria de menores" -> "secundaria"; "Inicial - Jardin" -> "inicial"
        nivel_base = re.split(r'[\s\-]+', nivel)[0]
        por_nivel = [
            row for row in candidatas
            if (row['nivel_modalidad'] or '').strip().lower().startswith(nivel_base)
        ]
        if len(por_nivel) == 1:
            return por_nivel[0]['codigo_modular']
    return ''


def sugerir_profesores(profesores, nombre, dni):
    """Match exacto por DNI, o candidatos rankeados por nombre.

    Nunca elige por el usuario: devuelve (match_dni, candidatos) para que la
    persona confirme en pantalla, igual que el resto del proyecto nunca acepta
    una coincidencia ambigua de forma automatica.
    """
    dni = re.sub(r'\D+', '', str(dni or ''))
    if dni:
        exacto = next((p for p in profesores if p.get('dni') == dni), None)
        if exacto:
            return exacto, []

    nombre = str(nombre or '').strip()
    if not nombre:
        return None, []

    nombres = {p['id']: p['nombre_completo'] or '' for p in profesores}
    por_id = {p['id']: p for p in profesores}
    candidatos = []
    for _, score, pid in rapidfuzz_process.extract(
        nombre, nombres, scorer=rapidfuzz_fuzz.WRatio, limit=5
    ):
        if score >= 70:
            candidatos.append({**por_id[pid], 'match_score': round(float(score), 1)})
    return None, candidatos


# ─── INFORMES DE CAMPO: extraccion y categorizacion (portado de
#     subproyectos/informes_campo_v3 para que la app sea autocontenida) ───────

CAMPO_SYSTEM_PROMPT = """Eres un analista de datos educativos para UGEL IBIR Imaza.
Tu tarea es convertir un informe de campo heterogeneo en datos estructurados por visita a institucion educativa.
Reglas:
1. No inventes hechos. Usa solo evidencia explicita del texto.
2. Si no hay codigo modular claro, usa cadena vacia "".
3. Si el documento menciona varias IE y no queda claro a cual corresponde un hallazgo, deja codigo_modular="".
4. Diferencia problemas reales de planes, listados, talleres, informes de comision o informes de gestion.
5. Para documentos administrativos sin visita directa a IE, devuelve hallazgos=[].
6. No calcules gravedad, urgencia ni confianza. Solo marca variables concretas con evidencia textual.
7. Cada hallazgo debe tener codigo_modular, variable_detectada, tema, descripcion y evidencia_textual literal (cita exacta del texto).
8. Si no hay evidencia textual, no crees el hallazgo.
9. Devuelve exclusivamente JSON valido que cumpla el esquema proporcionado.
"""

CAMPO_JSON_SCHEMA = {
    'documento': {
        'tipo_documento': 'uno de: informe_visita, informe_monitoreo, biae, informe_biae, acta, informe_gestion, informe_comision, listado, plan, otro',
        'resumen_ejecutivo': 'string, maximo 1200 caracteres',
    },
    'instituciones': [
        {'codigo_modular': '7 digitos o ""', 'nombre_ie': 'string', 'distrito': 'string', 'centro_poblado': 'string'},
    ],
    'hallazgos': [
        {
            'codigo_modular': '7 digitos o ""',
            'nombre_ie': 'nombre o numero local de la IE tal como aparece en el texto (ej. "288" o "288 - Wawaim"), aunque no sepas el codigo oficial',
            'centro_poblado': 'centro poblado de la IE si se menciona, para poder ubicarla aunque no tengas el codigo oficial',
            'variable_detectada': ' | '.join(CAMPO_CONCRETE_VARIABLES + ['otro']),
            'tema': 'string corto',
            'descripcion': 'string, maximo 1000 caracteres',
            'evidencia_textual': 'cita literal del informe, maximo 700 caracteres',
        },
    ],
}

def campo_criteria_block():
    lines = []
    for var in CAMPO_CONCRETE_VARIABLES:
        meta = CAMPO_VARIABLE_META.get(var, {})
        lines.append(f"- {var}: {meta.get('criterio', '')}")
    return '\n'.join(lines)

def build_campo_llm_prompt(text, filename=''):
    return (
        f"Archivo: {filename}\n\n"
        "Extrae informacion estructurada de este informe de campo para construir hallazgos por visita a IE.\n"
        "No asignes puntajes. No estimes gravedad, urgencia ni confianza.\n"
        "Marca solo problemas explicitamente observados y siempre copia la frase textual de evidencia.\n"
        "Si no hay evidencia textual, no crees el hallazgo.\n"
        "Si el hallazgo no se puede asociar a una IE especifica, deja codigo_modular vacio.\n"
        "Muchos informes no traen el codigo modular oficial de 7 digitos, solo el nombre o "
        "numero local de la IE (ej. \"288 - Wawaim\") y su centro poblado. En ese caso deja "
        "codigo_modular vacio pero SIEMPRE completa nombre_ie y centro_poblado del hallazgo "
        "con lo que diga el texto, para poder ubicar la IE despues aunque no tengas el codigo.\n\n"
        "Criterio textual minimo de cada variable:\n"
        f"{campo_criteria_block()}\n\n"
        "Esquema JSON requerido (usa exactamente estas claves):\n"
        f"{json.dumps(CAMPO_JSON_SCHEMA, ensure_ascii=False, indent=2)}\n\n"
        "Texto del informe:\n"
        f"{truncate_for_llm(text, 18000)}"
    )

def truncate_for_llm(text, max_chars):
    text = text or ''
    if len(text) <= max_chars:
        return text
    head = text[: int(max_chars * 0.65)]
    tail = text[-int(max_chars * 0.35):]
    return head + '\n\n[... TEXTO RECORTADO ...]\n\n' + tail

def parse_json_loose(text):
    clean = (text or '').strip()
    if clean.startswith('```'):
        clean = re.sub(r'^```(?:json)?', '', clean, flags=re.IGNORECASE).strip()
        clean = re.sub(r'```$', '', clean).strip()
    try:
        payload = json.loads(clean)
        return payload if isinstance(payload, dict) else {}
    except json.JSONDecodeError:
        start = clean.find('{')
        end = clean.rfind('}')
        if start >= 0 and end > start:
            payload = json.loads(clean[start:end + 1])
            return payload if isinstance(payload, dict) else {}
        raise

def call_campo_gemini(prompt):
    if not get_gemini_key():
        return {'status': 'missing_env', 'provider': 'gemini', 'payload': {}, 'error': ''}
    try:
        model = genai.GenerativeModel(GEMINI_DEFAULT_MODEL)
        response = model.generate_content(
            CAMPO_SYSTEM_PROMPT + '\n\n' + prompt,
            generation_config={'response_mime_type': 'application/json', 'temperature': 0},
        )
        payload = parse_json_loose(response.text)
        return {'status': 'ok', 'provider': 'gemini', 'payload': payload, 'error': ''}
    except Exception as exc:
        return {'status': 'failed', 'provider': 'gemini', 'payload': {}, 'error': f'{type(exc).__name__}: {exc}'}

def call_campo_openai(prompt):
    api_key = get_openai_key()
    if not api_key:
        return {'status': 'missing_env', 'provider': 'openai', 'payload': {}, 'error': ''}
    model = os.getenv('OPENAI_MODEL', 'gpt-4.1-mini')
    request_payload = {
        'model': model,
        'input': [
            {'role': 'system', 'content': CAMPO_SYSTEM_PROMPT},
            {'role': 'user', 'content': prompt + '\n\nResponde solo JSON valido, sin markdown.'},
        ],
        'text': {'format': {'type': 'json_object'}},
    }
    try:
        req = urllib.request.Request(
            'https://api.openai.com/v1/responses',
            data=json.dumps(request_payload).encode('utf-8'),
            headers={
                'Authorization': f'Bearer {api_key}',
                'Content-Type': 'application/json',
            },
            method='POST',
        )
        with urllib.request.urlopen(req, timeout=90) as resp:
            body = json.loads(resp.read().decode('utf-8'))
        text_parts = []
        for output in body.get('output', []) or []:
            for content in output.get('content', []) or []:
                if isinstance(content, dict) and isinstance(content.get('text'), str):
                    text_parts.append(content['text'])
        text = '\n'.join(text_parts).strip() or body.get('output_text', '')
        payload = parse_json_loose(text)
        return {'status': 'ok', 'provider': 'openai', 'payload': payload, 'error': ''}
    except Exception as exc:
        return {'status': 'failed', 'provider': 'openai', 'payload': {}, 'error': f'{type(exc).__name__}: {exc}'}

def extract_informe_campo_hallazgos(text, filename=''):
    """Gemini primero; si falla o no da JSON valido, cae automaticamente a
    OpenAI (si hay OPENAI_API_KEY configurada). Devuelve (payload, proveedor_usado, error)."""
    prompt = build_campo_llm_prompt(text, filename)

    result = call_campo_gemini(prompt)
    if result['status'] == 'ok' and result['payload']:
        return result['payload'], 'gemini', ''

    fallback = call_campo_openai(prompt)
    if fallback['status'] == 'ok' and fallback['payload']:
        return fallback['payload'], 'openai', ''

    error = fallback['error'] or result['error'] or 'sin_proveedor_disponible'
    return {}, 'ninguno', error

def consolidate_campo_hallazgos(payload, official_codes_resolver=None, resolver_by_nombre=None):
    """Agrupa hallazgos por codigo_modular en filas visita-IE, con los pares
    <variable>/span_<variable> materializados, igual que la tabla oficial del
    subproyecto. `official_codes_resolver(codigo)` puede normalizar/validar
    contra dim_institucion; si no resuelve, se intenta `resolver_by_nombre(
    nombre_ie, centro_poblado)` como respaldo (ver resolve_by_nombre_local --
    cubre el caso frecuente de informes que solo traen el numero/nombre local
    de la IE, no el codigo oficial). Si ninguno resuelve, el hallazgo queda
    para revision."""
    hallazgos = payload.get('hallazgos') or []
    instituciones = {inst.get('codigo_modular', ''): inst for inst in (payload.get('instituciones') or [])}

    by_code = {}
    sin_codigo = []
    for h in hallazgos:
        if not isinstance(h, dict):
            continue
        variable = h.get('variable_detectada')
        evidencia = (h.get('evidencia_textual') or '').strip()
        if variable not in CAMPO_CONCRETE_VARIABLES or not evidencia:
            continue  # sin evidencia literal, o variable "otro" -> no se materializa como bandera
        codigo = re.sub(r'\D', '', str(h.get('codigo_modular') or ''))
        if official_codes_resolver:
            codigo = official_codes_resolver(codigo) or ''
        if not codigo and resolver_by_nombre:
            codigo = resolver_by_nombre(h.get('nombre_ie', ''), h.get('centro_poblado', '')) or ''
        if not codigo:
            sin_codigo.append(h)
            continue
        by_code.setdefault(codigo, []).append(h)

    visitas = []
    for codigo, items in by_code.items():
        row = {'codigo_modular': codigo, 'hallazgos': items}
        observadas = []
        for var in CAMPO_CONCRETE_VARIABLES:
            match = next((it for it in items if it.get('variable_detectada') == var), None)
            if match:
                row[var] = 1
                row[f'span_{var}'] = (match.get('evidencia_textual') or '').strip()
                observadas.append(var)
            else:
                row[var] = None
                row[f'span_{var}'] = None
        row['n_variables_observadas'] = len(observadas)
        row['variables_observadas'] = '|'.join(observadas)
        inst = instituciones.get(codigo, {})
        row['nombre_ie_detectado'] = inst.get('nombre_ie', '')
        visitas.append(row)

    return visitas, sin_codigo

@app.route('/api/institucion/<codigo>/evolucion', methods=['GET'])
@require_auth
def institucion_evolucion(codigo):
    conn = get_db()
    try:
        codigo_real = resolve_codigo_modular_strict(conn, codigo) or canonical_codigo_modular(conn, codigo)
        if not codigo_real:
            return jsonify({'error': 'Codigo modular no reconocido.'}), 404
        official = conn.execute(
            'SELECT codigo_modular, nombre_iiee, distrito, nivel_modalidad FROM dim_institucion WHERE codigo_modular = ?',
            (codigo_real,),
        ).fetchone()
        data = build_institucion_evolucion(conn, codigo_real)
        data['institucion'] = dict(official) if official else {'codigo_modular': codigo_real}
        return jsonify(data)
    finally:
        conn.close()

# ─── OCR Y SINCRONIZACIÓN ─────────────────────────────────────────────────────

ALLOWED_OCR_EXTENSIONS = {'.pdf', '.jpg', '.jpeg', '.png', '.webp', '.heic'}
MAX_OCR_FILES = 10

@app.route('/api/ocr/upload_advanced', methods=['POST'])
@require_auth
def ocr_upload_advanced():
    uploaded_files = request.files.getlist('files') or request.files.getlist('file')
    uploaded_files = [file for file in uploaded_files if file and file.filename]
    if not uploaded_files:
        return jsonify({'error': 'No file uploaded'}), 400
    if len(uploaded_files) > MAX_OCR_FILES:
        return jsonify({'error': f'Sube como máximo {MAX_OCR_FILES} archivos a la vez.'}), 400
    for file in uploaded_files:
        ext = Path(file.filename).suffix.lower()
        if ext not in ALLOWED_OCR_EXTENSIONS:
            return jsonify({
                'error': f'Tipo de archivo no permitido: "{file.filename}". '
                         f'Formatos válidos: PDF, JPG, PNG, WEBP, HEIC.'
            }), 400

    extracted_parts = []
    methods = []
    filenames = []
    files_payload = []
    for idx, file in enumerate(uploaded_files, start=1):
        text_part, method, filename, file_bytes, mimetype = extract_text_from_upload(file)
        methods.append(method)
        filenames.append(filename)
        files_payload.append({
            'filename': filename,
            'mimetype': mimetype,
            'bytes': file_bytes,
        })
        if text_part:
            extracted_parts.append(f'--- DOCUMENTO {idx}: {filename} ---\n{text_part}')

    text = '\n\n'.join(extracted_parts).strip()
    method = '+'.join(sorted(set(methods))) if methods else 'unknown'

    raw_extracted = process_images_with_gemini(files_payload)
    parser = 'gemini_vision'
    if not raw_extracted and len(text) >= 50:
        raw_extracted = process_with_gemini(text)
        parser = 'gemini'
    if len(text) < 50 and not raw_extracted:
        return jsonify({'error': 'No se detectó texto en el documento.'}), 400
    if not raw_extracted:
        raw_extracted = parse_ficha_from_text(text)
        parser = 'fallback_reglas'

    extracted_data = normalize_ficha(raw_extracted, source='ocr')
    extracted_data['file_name'] = ' | '.join(filenames)
    extracted_data['extraction_method'] = method
    extracted_data['parser'] = parser
    extracted_data['confidence'] = 0.95 if method == 'pdfplumber' else (0.9 if parser == 'gemini_vision' else (0.78 if parser == 'fallback_reglas' else 0.85))
    extracted_data['raw_text'] = text
    extracted_data['extracted_json'] = json.dumps(raw_extracted, ensure_ascii=False)

    # Resolver el codigo modular contra dim_institucion -- primero por
    # codigo, y si no calza por nombre de IE (igual de estricto que
    # resolve_codigo_modular_strict para informes de campo: nunca se
    # adivina, solo se confirma o se deja para revision manual).
    conn = get_db()
    codigo_resuelto, nombre_bd = resolve_ficha_codigo_modular(
        conn, extracted_data.get('codigo_modular'), extracted_data.get('nombre_ie')
    )
    if codigo_resuelto:
        extracted_data['codigo_modular'] = codigo_resuelto
        ie = conn.execute(
            'SELECT nombre_iiee, nivel_modalidad, distrito FROM dim_institucion WHERE codigo_modular = ?',
            (codigo_resuelto,),
        ).fetchone()
        if ie:
            extracted_data['distrito'] = ie['distrito']
            extracted_data['nivel_modalidad'] = ie['nivel_modalidad']
            if not extracted_data.get('nombre_ie'):
                extracted_data['nombre_ie'] = ie['nombre_iiee']
    extracted_data['codigo_modular_resuelto'] = bool(codigo_resuelto)
    conn.close()

    # Se devuelve cada archivo original en base64 para que, si el especialista
    # confirma la ficha, el frontend lo reenvie y quede guardado en
    # archivo_subido/ficha_archivo -- hoy el archivo se descartaba apenas se
    # le extraia el texto.
    files_meta = [
        {
            'filename': item['filename'],
            'mimetype': item['mimetype'],
            'hash_sha1': hashlib.sha1(item['bytes']).hexdigest(),
            'file_b64': base64.b64encode(item['bytes']).decode('ascii'),
        }
        for item in files_payload if item.get('bytes')
    ]

    return jsonify({
        'success': True,
        'method': method,
        'parser': parser,
        'file_count': len(uploaded_files),
        'files': filenames,
        'files_meta': files_meta,
        'extracted': extracted_data
    })

def resolve_codigo_modular_strict(conn, value):
    """A diferencia de canonical_codigo_modular, NO devuelve un codigo que no
    exista en dim_institucion -- si no se puede resolver, el hallazgo debe
    quedar para revision (regla dura del subproyecto de informes de campo)."""
    code = re.sub(r'\D+', '', str(value or '').strip())
    if not code:
        return ''
    for candidate in (code, code.zfill(7), f'0{code}'):
        exists = conn.execute(
            'SELECT 1 FROM dim_institucion WHERE codigo_modular = ? LIMIT 1',
            (candidate,),
        ).fetchone()
        if exists:
            return candidate
    return ''

def _normalize_place_text(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode('ascii')
    return re.sub(r'[^A-Z0-9]+', '', text.upper())

def resolve_by_nombre_local(conn, nombre_ie, centro_poblado=''):
    """Fallback cuando el informe no trae el codigo modular oficial de 7
    digitos. En el padron de IBIR-IMAZA, dim_institucion.nombre_iiee suele
    guardar el numero local corto de la IE (ej. "288") en vez de un nombre
    propio -- y los especialistas escriben ese mismo numero junto al centro
    poblado en sus informes (ej. "288 - Wawaim"). Si el numero es unico en
    todo el padron Y el centro poblado (cuando el informe trae uno) coincide
    con el de ese registro, se resuelve; si el numero es unico pero el centro
    poblado NO calza (mismo numero local, otro caserio -- pasa en la
    practica), o si hay varios registros con ese numero y ninguno desambigua
    por centro poblado, se deja para revision manual -- mismo principio de
    "nunca adivinar" que el resto de este subproyecto."""
    match = re.match(r'\s*(\d{1,6})\b', str(nombre_ie or ''))
    if not match:
        return ''
    numero_local = match.group(1)
    rows = conn.execute(
        'SELECT codigo_modular, centro_poblado FROM dim_institucion WHERE nombre_iiee = ?',
        (numero_local,),
    ).fetchall()
    if not rows:
        return ''
    cp_norm = _normalize_place_text(centro_poblado)
    if len(rows) == 1:
        row = rows[0]
        db_cp_norm = _normalize_place_text(row['centro_poblado'])
        if cp_norm and db_cp_norm and cp_norm not in db_cp_norm and db_cp_norm not in cp_norm:
            return ''  # mismo numero local, pero el centro poblado no calza -- no adivinar
        return row['codigo_modular']
    if cp_norm:
        filtered = [
            r for r in rows
            if _normalize_place_text(r['centro_poblado'])
            and (cp_norm in _normalize_place_text(r['centro_poblado']) or _normalize_place_text(r['centro_poblado']) in cp_norm)
        ]
        if len(filtered) == 1:
            return filtered[0]['codigo_modular']
    return ''

@app.route('/api/ocr/informe_campo/upload', methods=['POST'])
@require_auth
def ocr_informe_campo_upload():
    uploaded = request.files.get('file') or (request.files.getlist('files') or [None])[0]
    if not uploaded or not uploaded.filename:
        return jsonify({'error': 'No se subio ningun archivo.'}), 400
    ext = Path(uploaded.filename).suffix.lower()
    if ext not in ALLOWED_OCR_EXTENSIONS:
        return jsonify({'error': f'Tipo de archivo no permitido: "{uploaded.filename}".'}), 400

    text, method, filename, file_bytes, mimetype = extract_text_from_upload(uploaded)
    if len(text) < 50:
        return jsonify({'error': 'No se detecto texto suficiente en el documento.'}), 400

    payload, provider, error = extract_informe_campo_hallazgos(text, filename)
    if not payload:
        return jsonify({'error': f'No se pudo categorizar el informe (proveedor: {provider}). {error}'}), 502

    conn = get_db()
    visitas, sin_codigo = consolidate_campo_hallazgos(
        payload,
        official_codes_resolver=lambda c: resolve_codigo_modular_strict(conn, c),
        resolver_by_nombre=lambda nombre, cp: resolve_by_nombre_local(conn, nombre, cp),
    )

    return jsonify({
        'success': True,
        'provider': provider,
        'method': method,
        'file_name': filename,
        'documento': payload.get('documento') or {},
        'instituciones': payload.get('instituciones') or [],
        'visitas': visitas,
        'hallazgos_sin_codigo': sin_codigo,
        'hash_sha1': hashlib.sha1(file_bytes).hexdigest(),
        'file_b64': base64.b64encode(file_bytes).decode('ascii'),
        'mimetype': mimetype,
        'texto_extraido': text,
        'variables_meta': CAMPO_VARIABLE_META,
    })

@app.route('/api/ocr/taller_asistencia/upload', methods=['POST'])
@require_auth
def ocr_taller_asistencia_upload():
    """Extrae la lista de participantes de una foto de planilla de asistencia.

    No escribe nada en la base: devuelve lo detectado mas sugerencias de match
    para que el especialista confirme cada asistente antes de guardar (mismo
    contrato que el OCR de fichas y el de informes de campo).
    """
    uploaded_files = [f for f in request.files.getlist('files') if f and f.filename]
    if not uploaded_files:
        single = request.files.get('file')
        if single and single.filename:
            uploaded_files = [single]
    if not uploaded_files:
        return jsonify({'error': 'No se subió ningún archivo.'}), 400
    if len(uploaded_files) > MAX_OCR_FILES:
        return jsonify({'error': f'Sube como máximo {MAX_OCR_FILES} archivos a la vez.'}), 400
    for archivo in uploaded_files:
        if Path(archivo.filename).suffix.lower() not in ALLOWED_OCR_EXTENSIONS:
            return jsonify({'error': f'Tipo de archivo no permitido: "{archivo.filename}".'}), 400

    partes_texto = []
    files_payload = []
    metodos = []
    for idx, archivo in enumerate(uploaded_files, start=1):
        texto, metodo, filename, file_bytes, mimetype = extract_text_from_upload(archivo)
        metodos.append(metodo)
        files_payload.append({'filename': filename, 'mimetype': mimetype, 'bytes': file_bytes})
        if texto:
            partes_texto.append(f'--- DOCUMENTO {idx}: {filename} ---\n{texto}')
    texto_total = '\n\n'.join(partes_texto).strip()

    payload, provider, error = extract_taller_asistentes(files_payload, texto_total)
    if not payload:
        return jsonify({'error': f'No se pudo leer la lista de participantes (proveedor: {provider}). {error}'}), 502

    conn = get_db()
    try:
        doc = payload.get('documento') or {}
        nombre_sede = str(doc.get('nombre_ie_sede') or '').strip()
        nivel_doc = str(doc.get('nivel_educativo') or '').strip()
        codigo_sede = (
            resolve_codigo_modular_strict(conn, nombre_sede)
            or resolver_ie_por_nombre_y_nivel(conn, nombre_sede, nivel_doc)
            or resolve_ficha_codigo_modular(conn, '', nombre_sede)[0]
            or resolve_by_nombre_local(conn, nombre_sede, doc.get('sede', ''))
        )

        profesores = rows_to_list(conn.execute('''
            SELECT p.id, p.dni, p.nombre_completo, p.cargo, p.celular, p.correo,
                   p.condicion_laboral, p.codigo_modular_ie,
                   i.nombre_iiee AS ie_nombre, i.centro_poblado AS ie_centro_poblado
            FROM profesor p
            LEFT JOIN dim_institucion i ON i.codigo_modular = p.codigo_modular_ie
            WHERE p.estado_fila = 1
        ''').fetchall())

        asistentes = []
        for fila in payload.get('asistentes') or []:
            if not isinstance(fila, dict):
                continue
            nombre = str(fila.get('nombre_completo') or '').strip()
            dni = re.sub(r'\D+', '', str(fila.get('dni') or ''))
            ie_texto = str(fila.get('institucion_educativa_texto') or '').strip()
            match_dni, candidatos = sugerir_profesores(profesores, nombre, dni)

            # La columna de IE viene manuscrita y abreviada ("A.A.H"): se intenta
            # resolver, pero lo normal es que no alcance y lo confirme la persona.
            codigo_ie = (
                resolver_ie_por_nombre_y_nivel(conn, ie_texto, nivel_doc)
                or resolve_ficha_codigo_modular(conn, '', ie_texto)[0]
                or resolve_by_nombre_local(conn, ie_texto)
            )
            if not codigo_ie and match_dni:
                codigo_ie = match_dni.get('codigo_modular_ie') or ''

            asistentes.append({
                'fila_num': fila.get('fila_num'),
                'nombre_detectado': nombre,
                'dni_detectado': dni,
                'institucion_educativa_texto': ie_texto,
                'codigo_modular_ie_sugerido': codigo_ie or '',
                'grado_a_cargo': str(fila.get('grado_a_cargo') or '').strip(),
                'area': str(fila.get('area') or '').strip(),
                'celular_detectado': str(fila.get('celular') or '').strip(),
                'correo_detectado': str(fila.get('correo') or '').strip(),
                'profesor_match': match_dni,
                'candidatos_profesor': candidatos,
            })
    finally:
        conn.close()

    files_meta = [{
        'filename': item['filename'],
        'mimetype': item['mimetype'],
        'hash_sha1': hashlib.sha1(item['bytes']).hexdigest(),
        'file_b64': base64.b64encode(item['bytes']).decode('ascii'),
    } for item in files_payload if item.get('bytes')]

    return jsonify({
        'success': True,
        'provider': provider,
        'method': '+'.join(sorted(set(m for m in metodos if m))),
        'documento': {**doc, 'codigo_modular_sede_sugerido': codigo_sede or ''},
        'asistentes': asistentes,
        'files_meta': files_meta,
    })


@app.route('/api/talleres_capacitacion', methods=['GET', 'POST'])
@require_auth
def talleres_capacitacion():
    conn = get_db()
    try:
        if request.method == 'GET':
            incluir_eliminados = request.args.get('incluir_eliminadas') == '1' or request.args.get('incluir_eliminados') == '1'
            where = '' if incluir_eliminados else 'WHERE COALESCE(t.estado_fila, 1) = 1'
            rows = conn.execute(f'''
                SELECT t.*, i.nombre_iiee AS sede_nombre, i.nivel_modalidad AS sede_nivel,
                       cu.nombre AS creado_por_nombre, cu.email AS creado_por_email,
                       mu.nombre AS modificado_por_nombre,
                       (SELECT COUNT(*) FROM taller_asistencia a
                         WHERE a.taller_id = t.id AND a.estado_fila = 1) AS total_asistentes
                FROM taller_capacitacion t
                LEFT JOIN dim_institucion i ON i.codigo_modular = t.codigo_modular_sede
                LEFT JOIN app_user cu ON cu.user_id = t.creado_por
                LEFT JOIN app_user mu ON mu.user_id = t.modificado_por
                {where}
                ORDER BY t.creado_en DESC
            ''').fetchall()
            return jsonify(rows_to_list(rows))

        data = request.get_json() or {}
        asistentes = data.get('asistentes') or []
        if not asistentes:
            return jsonify({'error': 'No hay asistentes para guardar.'}), 400

        user_id = g.current_user['user_id']
        doc = data.get('documento') or {}
        codigo_sede = resolve_codigo_modular_strict(conn, str(doc.get('codigo_modular_sede') or '')) or None

        cur = conn.execute('''
            INSERT INTO taller_capacitacion (
                nombre_evento, compromiso_indicador, nivel_educativo, dre, ugel,
                codigo_modular_sede, nombre_sede_detectado, fecha_inicio, fecha_fin,
                metodo_extraccion, parser, creado_por
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (
            str(doc.get('nombre_evento') or '').strip(),
            str(doc.get('compromiso_indicador') or '').strip(),
            str(doc.get('nivel_educativo') or '').strip(),
            str(doc.get('dre') or '').strip(),
            str(doc.get('ugel') or '').strip(),
            codigo_sede,
            str(doc.get('nombre_ie_sede') or '').strip(),
            str(doc.get('fecha_inicio') or '').strip(),
            str(doc.get('fecha_fin') or '').strip(),
            str(data.get('method') or '').strip(),
            str(data.get('provider') or '').strip(),
            user_id,
        ))
        taller_id = cur.lastrowid

        for idx, item in enumerate(data.get('files') or []):
            if not isinstance(item, dict) or not item.get('file_b64'):
                continue
            try:
                contenido = base64.b64decode(item['file_b64'])
            except Exception:
                continue
            archivo_cur = conn.execute('''
                INSERT INTO archivo_subido (tipo, nombre_archivo, mimetype, tamano_bytes, contenido, hash_sha1, subido_por)
                VALUES ('taller_asistencia', ?, ?, ?, ?, ?, ?)
            ''', (
                item.get('filename', ''), item.get('mimetype', ''), len(contenido),
                contenido, item.get('hash_sha1', ''), user_id,
            ))
            conn.execute(
                'INSERT INTO taller_archivo (taller_id, archivo_subido_id, orden) VALUES (?, ?, ?)',
                (taller_id, archivo_cur.lastrowid, idx),
            )

        guardados = 0
        profesores_creados = 0
        omitidos = []
        for fila in asistentes:
            if not isinstance(fila, dict):
                continue
            profesor_id = safe_int(fila.get('profesor_id'), 0) or None
            nuevo = fila.get('nuevo_profesor')

            if not profesor_id and isinstance(nuevo, dict):
                valores, error = normalize_profesor_payload(conn, nuevo)
                if error:
                    omitidos.append(f"{fila.get('nombre_detectado', '')}: {error}")
                    continue
                try:
                    columnas = list(valores.keys()) + ['fuente', 'creado_por']
                    params = list(valores.values()) + ['ocr_taller', user_id]
                    placeholders = ', '.join('?' for _ in columnas)
                    profesor_cur = conn.execute(
                        f'INSERT INTO profesor ({", ".join(columnas)}) VALUES ({placeholders})', params
                    )
                    profesor_id = profesor_cur.lastrowid
                    profesores_creados += 1
                except sqlite3.IntegrityError:
                    existente = conn.execute(
                        'SELECT id FROM profesor WHERE dni = ?', (valores['dni'],)
                    ).fetchone()
                    if existente:
                        profesor_id = existente['id']
                    else:
                        omitidos.append(f"{fila.get('nombre_detectado', '')}: no se pudo crear el profesor.")
                        continue

            if not profesor_id:
                omitidos.append(f"{fila.get('nombre_detectado', '')}: sin profesor confirmado.")
                continue

            conn.execute('''
                INSERT INTO taller_asistencia (
                    taller_id, profesor_id, nombre_detectado, dni_detectado, grado_a_cargo,
                    area, celular_detectado, correo_detectado, codigo_modular_ie_detectado,
                    nombre_ie_detectado, match_metodo, match_score, creado_por
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                taller_id,
                profesor_id,
                str(fila.get('nombre_detectado') or '').strip(),
                re.sub(r'\D+', '', str(fila.get('dni_detectado') or '')),
                str(fila.get('grado_a_cargo') or '').strip(),
                str(fila.get('area') or '').strip(),
                str(fila.get('celular_detectado') or '').strip(),
                str(fila.get('correo_detectado') or '').strip(),
                resolve_codigo_modular_strict(conn, str(fila.get('codigo_modular_ie_detectado') or '')) or None,
                str(fila.get('institucion_educativa_texto') or '').strip(),
                str(fila.get('match_metodo') or '').strip(),
                fila.get('match_score'),
                user_id,
            ))
            guardados += 1

        conn.commit()
        return jsonify({
            'success': True,
            'taller_id': taller_id,
            'asistentes_guardados': guardados,
            'profesores_creados': profesores_creados,
            'omitidos': omitidos,
        })
    except sqlite3.IntegrityError as exc:
        conn.rollback()
        return jsonify({'error': f'No se pudo guardar la asistencia: {exc}'}), 400
    finally:
        conn.close()


TALLER_EDITABLE_FIELDS = [
    'nombre_evento', 'compromiso_indicador', 'nivel_educativo', 'dre', 'ugel',
    'fecha_inicio', 'fecha_fin', 'especialista_responsable',
]


@app.route('/api/talleres_capacitacion/<int:taller_id>', methods=['GET', 'PUT', 'DELETE'])
@require_auth
def taller_detalle(taller_id):
    conn = get_db()
    try:
        row = conn.execute('''
            SELECT t.*, i.nombre_iiee AS sede_nombre, i.nivel_modalidad AS sede_nivel,
                   cu.nombre AS creado_por_nombre, cu.email AS creado_por_email,
                   mu.nombre AS modificado_por_nombre
            FROM taller_capacitacion t
            LEFT JOIN dim_institucion i ON i.codigo_modular = t.codigo_modular_sede
            LEFT JOIN app_user cu ON cu.user_id = t.creado_por
            LEFT JOIN app_user mu ON mu.user_id = t.modificado_por
            WHERE t.id = ?
        ''', (taller_id,)).fetchone()
        if not row:
            return jsonify({'error': 'Taller no encontrado.'}), 404
        taller = dict(row)

        if request.method == 'GET':
            asistentes = conn.execute('''
                SELECT a.*, p.nombre_completo AS profesor_nombre, p.dni AS profesor_dni,
                       p.cargo AS profesor_cargo, i.nombre_iiee AS ie_nombre,
                       i.nivel_modalidad AS ie_nivel
                FROM taller_asistencia a
                LEFT JOIN profesor p ON p.id = a.profesor_id
                LEFT JOIN dim_institucion i ON i.codigo_modular = a.codigo_modular_ie_detectado
                WHERE a.taller_id = ? AND COALESCE(a.estado_fila, 1) = 1
                ORDER BY a.id
            ''', (taller_id,)).fetchall()
            archivos = conn.execute('''
                SELECT ar.id, ar.nombre_archivo, ar.mimetype, ta.orden
                FROM taller_archivo ta
                JOIN archivo_subido ar ON ar.id = ta.archivo_subido_id
                WHERE ta.taller_id = ? AND COALESCE(ta.estado_fila, 1) = 1 AND COALESCE(ar.estado_fila, 1) = 1
                ORDER BY ta.orden
            ''', (taller_id,)).fetchall()
            return jsonify({
                'taller': taller,
                'asistentes': rows_to_list(asistentes),
                'archivos': rows_to_list(archivos),
                'puede_editar': can_edit_row(g.current_user, taller.get('creado_por')),
            })

        if not can_edit_row(g.current_user, taller.get('creado_por')):
            return jsonify({'error': 'Solo puedes editar tus propios registros.'}), 403

        if request.method == 'DELETE':
            # Baja logica solo del taller: NO se tocan las asistencias. Si se
            # marcaran tambien, al restaurar volverian los participantes que
            # alguien habia quitado a mano. Los conteos y dashboards filtran por
            # el estado del taller, asi que ocultarlo basta.
            conn.execute(
                'UPDATE taller_capacitacion SET estado_fila = 0, modificado_por = ?, '
                'actualizado_en = CURRENT_TIMESTAMP WHERE id = ?',
                (g.current_user['user_id'], taller_id),
            )
            conn.commit()
            return jsonify({'success': True})

        data = request.get_json() or {}
        if data.get('restaurar'):
            conn.execute(
                'UPDATE taller_capacitacion SET estado_fila = 1, modificado_por = ?, '
                'actualizado_en = CURRENT_TIMESTAMP WHERE id = ?',
                (g.current_user['user_id'], taller_id),
            )
            conn.commit()
            updated = conn.execute('SELECT * FROM taller_capacitacion WHERE id = ?', (taller_id,)).fetchone()
            return jsonify({'success': True, 'taller': dict(updated)})

        updates = []
        params = []
        for campo in TALLER_EDITABLE_FIELDS:
            if campo in data:
                updates.append(f'{campo} = ?')
                params.append(str(data.get(campo) or '').strip())
        if 'codigo_modular_sede' in data:
            codigo = str(data.get('codigo_modular_sede') or '').strip()
            resuelto = resolve_codigo_modular_strict(conn, codigo) if codigo else ''
            if codigo and not resuelto:
                return jsonify({'error': f'La IE sede "{codigo}" no existe en el padrón.'}), 400
            updates.append('codigo_modular_sede = ?')
            params.append(resuelto or None)

        # Bajas individuales de asistentes (quitar a alguien de la lista sin
        # borrar el taller entero).
        for asistencia_id in data.get('asistentes_eliminados') or []:
            conn.execute(
                'UPDATE taller_asistencia SET estado_fila = 0, modificado_por = ?, '
                'actualizado_en = CURRENT_TIMESTAMP WHERE id = ? AND taller_id = ?',
                (g.current_user['user_id'], safe_int(asistencia_id, 0), taller_id),
            )

        if updates:
            updates.append('modificado_por = ?')
            params.append(g.current_user['user_id'])
            updates.append('actualizado_en = CURRENT_TIMESTAMP')
            params.append(taller_id)
            conn.execute(f'UPDATE taller_capacitacion SET {", ".join(updates)} WHERE id = ?', params)

        conn.commit()
        updated = conn.execute('SELECT * FROM taller_capacitacion WHERE id = ?', (taller_id,)).fetchone()
        return jsonify({'success': True, 'taller': dict(updated)})
    finally:
        conn.close()


@app.route('/api/informes_campo', methods=['GET', 'POST'])
@require_auth
def informes_campo():
    conn = get_db()
    try:
        if request.method == 'POST':
            data = request.get_json() or {}
            visitas = data.get('visitas') or []
            if not visitas:
                return jsonify({'error': 'No hay visitas para guardar.'}), 400

            user_id = g.current_user['user_id'] if g.current_user else None
            file_b64 = data.get('file_b64')
            archivo_id = None
            if file_b64:
                try:
                    contenido = base64.b64decode(file_b64)
                except Exception:
                    contenido = b''
                cur = conn.execute('''
                    INSERT INTO archivo_subido (tipo, nombre_archivo, mimetype, tamano_bytes, contenido, hash_sha1, subido_por)
                    VALUES ('informe_campo', ?, ?, ?, ?, ?, ?)
                ''', (
                    data.get('file_name', ''),
                    data.get('mimetype', ''),
                    len(contenido),
                    contenido,
                    data.get('hash_sha1', ''),
                    user_id,
                ))
                archivo_id = cur.lastrowid

            doc_meta = data.get('documento') or {}
            cur = conn.execute('''
                INSERT INTO informe_campo_documento (
                    archivo_subido_id, nombre_archivo, hash_sha1, tipo_documento,
                    especialista_detectado, metodo_extraccion, longitud_texto_documento,
                    texto_extraido, fuente, subido_por
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, 'app_upload', ?)
            ''', (
                archivo_id,
                data.get('file_name', ''),
                data.get('hash_sha1', ''),
                doc_meta.get('tipo_documento', ''),
                data.get('especialista_detectado', g.current_user['nombre'] if g.current_user else ''),
                data.get('method', ''),
                len(data.get('texto_extraido', '') or ''),
                data.get('texto_extraido', ''),
                user_id,
            ))
            documento_id = cur.lastrowid

            guardadas = []
            for visita in visitas:
                codigo = resolve_codigo_modular_strict(conn, visita.get('codigo_modular'))
                if not codigo:
                    continue
                visita_id = hashlib.sha1(
                    f"{documento_id}|{codigo}|{visita.get('fecha_visita', '')}".encode('utf-8')
                ).hexdigest()[:16]
                columns = ['visita_campo_id', 'documento_id', 'codigo_modular', 'nombre_ie_detectado',
                           'fecha_visita', 'especialista_detectado', 'n_variables_observadas',
                           'variables_observadas', 'requiere_revision', 'fuente']
                values = [
                    visita_id, documento_id, codigo, visita.get('nombre_ie_detectado', ''),
                    visita.get('fecha_visita', ''), data.get('especialista_detectado', ''),
                    safe_int(visita.get('n_variables_observadas'), 0),
                    visita.get('variables_observadas', ''), 0, 'app_upload',
                ]
                for var in CAMPO_CONCRETE_VARIABLES:
                    columns.append(var)
                    values.append(visita.get(var))
                    columns.append(f'span_{var}')
                    values.append(visita.get(f'span_{var}'))
                placeholders = ', '.join('?' for _ in columns)
                conn.execute(
                    f'INSERT OR REPLACE INTO informe_campo_visita ({", ".join(columns)}) VALUES ({placeholders})',
                    values,
                )
                for h in visita.get('hallazgos') or []:
                    conn.execute('''
                        INSERT INTO informe_campo_hallazgo (
                            visita_campo_id, documento_id, codigo_modular_hallazgo,
                            variable_detectada, tema, descripcion, evidencia_textual
                        ) VALUES (?, ?, ?, ?, ?, ?, ?)
                    ''', (
                        visita_id, documento_id, h.get('codigo_modular', ''),
                        h.get('variable_detectada', ''), h.get('tema', ''),
                        h.get('descripcion', ''), h.get('evidencia_textual', ''),
                    ))
                guardadas.append(visita_id)

            rebuild_simon_operational_summary(conn)
            conn.commit()
            return jsonify({'success': True, 'visitas_guardadas': len(guardadas), 'visita_ids': guardadas})

        incluir_eliminadas = request.args.get('incluir_eliminadas') == '1'
        q = (request.args.get('q') or '').strip()
        creado_por_email = (request.args.get('creado_por_email') or '').strip()
        limit = safe_int(request.args.get('limit'), 200)
        offset = safe_int(request.args.get('offset'), 0)

        where = [] if incluir_eliminadas else ['COALESCE(v.estado_fila, 1) = 1']
        params = []
        if q:
            where.append('(v.nombre_ie_detectado LIKE ? OR v.codigo_modular LIKE ? OR v.especialista_detectado LIKE ?)')
            like = f'%{q}%'
            params.extend([like, like, like])
        if creado_por_email:
            where.append('cu.email LIKE ?')
            params.append(f'%{creado_por_email}%')
        where_sql = f"WHERE {' AND '.join(where)}" if where else ''

        rows = conn.execute(f'''
            SELECT v.*, d.subido_por AS creado_por, cu.nombre AS creado_por_nombre,
                   cu.email AS creado_por_email, mu.nombre AS modificado_por_nombre
            FROM informe_campo_visita v
            LEFT JOIN informe_campo_documento d ON d.id = v.documento_id
            LEFT JOIN app_user cu ON cu.user_id = d.subido_por
            LEFT JOIN app_user mu ON mu.user_id = v.modificado_por
            {where_sql}
            ORDER BY v.creado_en DESC
            LIMIT ? OFFSET ?
        ''', [*params, limit, offset]).fetchall()
        return jsonify(rows_to_list(rows))
    finally:
        conn.close()

@app.route('/api/informes_campo/<visita_campo_id>', methods=['GET', 'PUT', 'DELETE'])
@require_auth
def informe_campo_detalle(visita_campo_id):
    conn = get_db()
    try:
        row = conn.execute('''
            SELECT v.*, d.subido_por AS creado_por, d.texto_extraido, d.tipo_documento,
                   d.nombre_archivo AS documento_nombre_archivo, d.archivo_subido_id,
                   cu.nombre AS creado_por_nombre, mu.nombre AS modificado_por_nombre
            FROM informe_campo_visita v
            LEFT JOIN informe_campo_documento d ON d.id = v.documento_id
            LEFT JOIN app_user cu ON cu.user_id = d.subido_por
            LEFT JOIN app_user mu ON mu.user_id = v.modificado_por
            WHERE v.visita_campo_id = ?
        ''', (visita_campo_id,)).fetchone()
        if not row:
            return jsonify({'error': 'Informe de campo no encontrado.'}), 404
        visita = dict(row)

        if request.method == 'GET':
            hallazgos = conn.execute('''
                SELECT * FROM informe_campo_hallazgo
                WHERE visita_campo_id = ? AND estado_fila = 1
                ORDER BY id
            ''', (visita_campo_id,)).fetchall()
            return jsonify({
                'visita': visita,
                'hallazgos': rows_to_list(hallazgos),
                'variables_meta': CAMPO_VARIABLE_META,
                'puede_editar': can_edit_row(g.current_user, visita.get('creado_por')),
            })

        if not can_edit_row(g.current_user, visita.get('creado_por')):
            return jsonify({'error': 'Solo puedes editar tus propios registros.'}), 403

        if request.method == 'DELETE':
            conn.execute(
                'UPDATE informe_campo_visita SET estado_fila = 0, modificado_por = ?, '
                'actualizado_en = CURRENT_TIMESTAMP WHERE visita_campo_id = ?',
                (g.current_user['user_id'], visita_campo_id),
            )
            rebuild_simon_operational_summary(conn)
            conn.commit()
            return jsonify({'success': True})

        data = request.get_json() or {}

        if data.get('restaurar'):
            conn.execute(
                'UPDATE informe_campo_visita SET estado_fila = 1, modificado_por = ?, '
                'actualizado_en = CURRENT_TIMESTAMP WHERE visita_campo_id = ?',
                (g.current_user['user_id'], visita_campo_id),
            )
            conn.commit()
            updated = conn.execute(
                'SELECT * FROM informe_campo_visita WHERE visita_campo_id = ?', (visita_campo_id,)
            ).fetchone()
            return jsonify({'success': True, 'visita': dict(updated)})

        updates = []
        params = []

        if 'codigo_modular' in data:
            codigo = resolve_codigo_modular_strict(conn, data.get('codigo_modular'))
            if not codigo:
                return jsonify({'error': 'Código modular no reconocido.'}), 400
            updates.append('codigo_modular = ?')
            params.append(codigo)

        for field in ('nombre_ie_detectado', 'fecha_visita', 'especialista_detectado', 'motivo_revision'):
            if field in data:
                updates.append(f'{field} = ?')
                params.append(data.get(field))

        if 'requiere_revision' in data:
            updates.append('requiere_revision = ?')
            params.append(1 if data.get('requiere_revision') else 0)

        # Flujo de revision. Marcar puede cualquiera (el que nota el problema);
        # resolver le toca al especialista del registro -- o a un administrador
        # o supervisor, que existen justamente para destrabar.
        if 'revision_estado' in data:
            nuevo_estado = str(data.get('revision_estado') or 'ok').strip()
            if nuevo_estado not in REVISION_ESTADOS:
                return jsonify({'error': f'Estado de revisión no válido: {nuevo_estado}.'}), 400
            estado_actual = str(visita.get('revision_estado') or 'ok')
            ahora = datetime.now().isoformat(timespec='seconds')
            yo = g.current_user['user_id']

            if nuevo_estado in ('pendiente', 'en_revision') and estado_actual in ('ok', 'resuelto'):
                # Se abre la revision: queda a cargo del especialista del
                # registro, que es quien puede confirmar lo que observo.
                asignado_nombre = str(visita.get('especialista_detectado') or '').strip()
                asignado_id = ''
                if asignado_nombre:
                    encontrado = conn.execute(
                        'SELECT user_id FROM app_user WHERE activo = 1 AND '
                        'UPPER(TRIM(nombre)) = UPPER(TRIM(?)) LIMIT 1', (asignado_nombre,)
                    ).fetchone()
                    asignado_id = encontrado['user_id'] if encontrado else ''
                updates += ['revision_solicitada_por = ?', 'revision_solicitada_en = ?',
                            'revision_asignado_a = ?', 'revision_asignado_nombre = ?',
                            'revision_resuelto_por = ?', 'revision_resuelto_en = ?']
                params += [yo, ahora, asignado_id, asignado_nombre, None, None]

            if nuevo_estado == 'resuelto':
                asignado = str(visita.get('revision_asignado_a') or '')
                puede_resolver = (
                    normalize_role(g.current_user['rol']) in ROLES_QUE_EDITAN_TODO
                    or not asignado or asignado == yo
                )
                if not puede_resolver:
                    return jsonify({
                        'error': 'Esta revisión le corresponde a '
                                 f'{visita.get("revision_asignado_nombre") or "otro especialista"}.'
                    }), 403
                updates += ['revision_resuelto_por = ?', 'revision_resuelto_en = ?']
                params += [yo, ahora]

            updates.append('revision_estado = ?')
            params.append(nuevo_estado)
            updates.append('requiere_revision = ?')
            params.append(1 if nuevo_estado in ('pendiente', 'en_revision') else 0)

        if 'revision_nota' in data:
            updates.append('revision_nota = ?')
            params.append(str(data.get('revision_nota') or '').strip())

        variables_tocadas = [v for v in CAMPO_CONCRETE_VARIABLES if v in data]
        for var in variables_tocadas:
            updates.append(f'{var} = ?')
            params.append(1 if data.get(var) else None)
            span_key = f'span_{var}'
            if span_key in data:
                updates.append(f'{span_key} = ?')
                params.append(data.get(span_key))

        if variables_tocadas:
            merged_flags = {
                var: (data.get(var) if var in data else visita.get(var))
                for var in CAMPO_CONCRETE_VARIABLES
            }
            observadas = [var for var in CAMPO_CONCRETE_VARIABLES if merged_flags.get(var)]
            updates.append('n_variables_observadas = ?')
            params.append(len(observadas))
            updates.append('variables_observadas = ?')
            params.append(', '.join(observadas))

        if updates:
            updates.append('modificado_por = ?')
            params.append(g.current_user['user_id'])
            updates.append('actualizado_en = CURRENT_TIMESTAMP')
            params.append(visita_campo_id)
            conn.execute(
                f'UPDATE informe_campo_visita SET {", ".join(updates)} WHERE visita_campo_id = ?',
                params,
            )

        if 'hallazgos' in data:
            # Solo se reemplazan los hallazgos vigentes: los que un especialista
            # ya descarto se conservan como registro de la validacion (el GET
            # tampoco los devuelve, asi que no vuelven en el payload).
            conn.execute(
                'DELETE FROM informe_campo_hallazgo '
                'WHERE visita_campo_id = ? AND estado_fila = 1',
                (visita_campo_id,),
            )
            for h in data.get('hallazgos') or []:
                if not isinstance(h, dict):
                    continue
                conn.execute('''
                    INSERT INTO informe_campo_hallazgo (
                        visita_campo_id, documento_id, codigo_modular_hallazgo,
                        variable_detectada, tema, descripcion, evidencia_textual
                    ) VALUES (?, ?, ?, ?, ?, ?, ?)
                ''', (
                    visita_campo_id, visita.get('documento_id'), h.get('codigo_modular', ''),
                    h.get('variable_detectada', ''), h.get('tema', ''),
                    h.get('descripcion', ''), h.get('evidencia_textual', ''),
                ))

        rebuild_simon_operational_summary(conn)
        conn.commit()
        updated = conn.execute(
            'SELECT * FROM informe_campo_visita WHERE visita_campo_id = ?', (visita_campo_id,)
        ).fetchone()
        return jsonify({'success': True, 'visita': dict(updated)})
    finally:
        conn.close()

@app.route('/api/fichas', methods=['GET', 'POST'])
@require_auth
def fichas():
    conn = get_db()
    try:
        if request.method == 'POST':
            data = request.get_json() or {}
            ficha = save_ficha_to_db(conn, data)
            rebuild_simon_operational_summary(conn)
            conn.commit()
            return jsonify({'success': True, 'ficha': ficha})

        incluir_eliminadas = request.args.get('incluir_eliminadas') == '1'
        q = (request.args.get('q') or '').strip()
        creado_por_email = (request.args.get('creado_por_email') or '').strip()
        limit = safe_int(request.args.get('limit'), 100)
        offset = safe_int(request.args.get('offset'), 0)

        where = [] if incluir_eliminadas else ['COALESCE(f.estado_fila, 1) = 1']
        params = []
        if q:
            where.append('(f.nombre_ie LIKE ? OR f.codigo_modular LIKE ? OR f.docente LIKE ?)')
            like = f'%{q}%'
            params.extend([like, like, like])
        if creado_por_email:
            where.append('cu.email LIKE ?')
            params.append(f'%{creado_por_email}%')
        where_sql = f"WHERE {' AND '.join(where)}" if where else ''

        rows = conn.execute(f'''
            SELECT f.*, cu.nombre AS creado_por_nombre, cu.email AS creado_por_email,
                   mu.nombre AS modificado_por_nombre
            FROM fichas_monitoreo f
            LEFT JOIN app_user cu ON cu.user_id = f.creado_por
            LEFT JOIN app_user mu ON mu.user_id = f.modificado_por
            {where_sql}
            ORDER BY f.fecha_sincronizacion DESC, f.id DESC
            LIMIT ? OFFSET ?
        ''', [*params, limit, offset]).fetchall()
        return jsonify(rows_to_list(rows))
    finally:
        conn.close()

@app.route('/api/fichas/<int:ficha_id>', methods=['GET', 'PUT', 'DELETE'])
@require_auth
def ficha_detalle(ficha_id):
    conn = get_db()
    try:
        row = conn.execute('''
            SELECT f.*, cu.nombre AS creado_por_nombre, mu.nombre AS modificado_por_nombre
            FROM fichas_monitoreo f
            LEFT JOIN app_user cu ON cu.user_id = f.creado_por
            LEFT JOIN app_user mu ON mu.user_id = f.modificado_por
            WHERE f.id = ?
        ''', (ficha_id,)).fetchone()
        if not row:
            return jsonify({'error': 'Ficha no encontrada.'}), 404
        ficha = dict(row)

        if request.method == 'GET':
            respuestas = get_dynamic_responses_for_fichas(conn, [ficha_id]).get(ficha_id, [])
            archivos = conn.execute('''
                SELECT a.id, a.nombre_archivo, a.mimetype, a.tamano_bytes, a.hash_sha1, a.creado_en, fa.orden
                FROM ficha_archivo fa
                JOIN archivo_subido a ON a.id = fa.archivo_subido_id
                WHERE fa.ficha_id = ? AND fa.estado_fila = 1 AND a.estado_fila = 1
                ORDER BY fa.orden
            ''', (ficha_id,)).fetchall()
            return jsonify({
                'ficha': ficha,
                'respuestas_dinamicas': respuestas,
                'archivos': rows_to_list(archivos),
                'puede_editar': can_edit_row(g.current_user, ficha.get('creado_por')),
            })

        if not can_edit_row(g.current_user, ficha.get('creado_por')):
            return jsonify({'error': 'Solo puedes editar tus propios registros.'}), 403

        if request.method == 'DELETE':
            conn.execute(
                'UPDATE fichas_monitoreo SET estado_fila = 0, modificado_por = ?, '
                'actualizado_en = CURRENT_TIMESTAMP WHERE id = ?',
                (g.current_user['user_id'], ficha_id),
            )
            rebuild_simon_operational_summary(conn)
            conn.commit()
            return jsonify({'success': True})

        data = request.get_json() or {}

        if data.get('restaurar'):
            conn.execute(
                'UPDATE fichas_monitoreo SET estado_fila = 1, modificado_por = ?, '
                'actualizado_en = CURRENT_TIMESTAMP WHERE id = ?',
                (g.current_user['user_id'], ficha_id),
            )
            conn.commit()
            updated = conn.execute('SELECT * FROM fichas_monitoreo WHERE id = ?', (ficha_id,)).fetchone()
            return jsonify({'success': True, 'ficha': dict(updated)})

        updates = []
        params = []
        merged = dict(ficha)
        for field in FICHA_EDITABLE_FIELDS:
            if field in data:
                updates.append(f'{field} = ?')
                params.append(data.get(field))
                merged[field] = data.get(field)

        if any(field in data for field in SIMON_LEVEL_FIELDS):
            updates.append('promedio = ?')
            params.append(calculate_promedio(merged))

        if updates:
            updates.append('modificado_por = ?')
            params.append(g.current_user['user_id'])
            updates.append('actualizado_en = CURRENT_TIMESTAMP')
            params.append(ficha_id)
            conn.execute(f'UPDATE fichas_monitoreo SET {", ".join(updates)} WHERE id = ?', params)

        if 'respuestas_dinamicas' in data:
            conn.execute(
                "DELETE FROM ficha_respuesta_instrumento WHERE ficha_id = ? AND instrumento_tipo = 'dinamica'",
                (ficha_id,),
            )
            dynamic_instrument = get_dynamic_instrument_row(conn)
            default_instrumento_codigo = dynamic_instrument['codigo'] if dynamic_instrument else 'dinamico'
            for r in data.get('respuestas_dinamicas') or []:
                if not isinstance(r, dict):
                    continue
                pregunta_codigo = first_text(r, 'pregunta_codigo')
                if not pregunta_codigo:
                    continue
                conn.execute('''
                    INSERT INTO ficha_respuesta_instrumento (
                        ficha_id, instrumento_codigo, instrumento_tipo, seccion_clave,
                        pregunta_codigo, nivel, respuesta_texto, observacion, metadata_json
                    ) VALUES (?, ?, 'dinamica', ?, ?, ?, ?, ?, ?)
                ''', (
                    ficha_id,
                    first_text(r, 'instrumento_codigo') or default_instrumento_codigo,
                    first_text(r, 'seccion_clave'),
                    pregunta_codigo,
                    normalize_level(r.get('nivel')),
                    first_text(r, 'respuesta_texto', 'respuesta'),
                    first_text(r, 'observacion'),
                    json_dumps(r),
                ))

        rebuild_simon_operational_summary(conn)
        conn.commit()
        updated = conn.execute('SELECT * FROM fichas_monitoreo WHERE id = ?', (ficha_id,)).fetchone()
        respuestas = get_dynamic_responses_for_fichas(conn, [ficha_id]).get(ficha_id, [])
        return jsonify({'success': True, 'ficha': dict(updated), 'respuestas_dinamicas': respuestas})
    finally:
        conn.close()

@app.route('/api/fichas/<int:ficha_id>/archivos', methods=['GET'])
@require_auth
def ficha_archivos(ficha_id):
    conn = get_db()
    try:
        rows = conn.execute('''
            SELECT a.id, a.nombre_archivo, a.mimetype, a.tamano_bytes, a.hash_sha1, a.creado_en, fa.orden
            FROM ficha_archivo fa
            JOIN archivo_subido a ON a.id = fa.archivo_subido_id
            WHERE fa.ficha_id = ? AND fa.estado_fila = 1 AND a.estado_fila = 1
            ORDER BY fa.orden
        ''', (ficha_id,)).fetchall()
        return jsonify(rows_to_list(rows))
    finally:
        conn.close()

@app.route('/api/sync', methods=['POST'])
@require_auth
def sync_data():
    data = request.get_json() or []
    if not isinstance(data, list):
        data = [data]

    conn = get_db()
    inserted = 0
    try:
        for f in data:
            save_ficha_to_db(conn, f)
            inserted += 1
        rebuild_simon_operational_summary(conn)
        conn.commit()
    except Exception as e:
        conn.rollback()
        print(f"Error al sincronizar: {e}")
        return jsonify({'error': 'No se pudo sincronizar la información. Intenta nuevamente.'}), 500
    finally:
        conn.close()

    return jsonify({
        'success': True,
        'message': f'{inserted} registros guardados en la BD (fichas_monitoreo).',
        'synced_count': inserted
    })


# ─── MAIN ─────────────────────────────────────────────────────────────────────

if __name__ == '__main__':
    # ensure_db_schema() ya se ejecuto al importar el modulo (ver mas arriba).
    port = int(os.getenv('PORT', '8000'))
    debug = os.getenv('FLASK_DEBUG', '0') == '1'
    print('=' * 50)
    print('  SUGKA LAB API Server')
    print(f'  DB : {DB_PATH}')
    print(f'  URL: http://localhost:{port}')
    print('=' * 50)
    app.run(host='0.0.0.0', port=port, debug=debug)
